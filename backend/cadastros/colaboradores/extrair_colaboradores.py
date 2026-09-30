from __future__ import annotations



import getpass

import os

import shutil

import subprocess

import sys

import time

import unicodedata

from collections import Counter

from datetime import date, datetime

from pathlib import Path



import psutil

import win32gui

import win32process

from openpyxl import load_workbook

from pywinauto import Desktop

from pywinauto.controls.uiawrapper import UIAWrapper

from pywinauto.uia_element_info import UIAElementInfo





# ============================================================================

# CONFIGURACAO

# ============================================================================



RAIZ = Path(__file__).resolve().parents[3]



RM_EXE = Path(

    r"C:\TOTVS\RM.NET\RM TEST\RM.exe"

)



RM_ARGUMENTO = "ALIASFILTER=celeo-prd-197741"



PASTA_BRONZE = (

    RAIZ

    / "data"

    / "bronze"

    / "cadastros"

    / "colaboradores"

)



ARQUIVO_TEMPORARIO = (

    PASTA_BRONZE

    / "Colaboradores_RM_novo.xlsx"

)



ARQUIVO_BRONZE = (

    PASTA_BRONZE

    / "Colaboradores_RM.xlsx"

)



COLUNAS_ESPERADAS = [

    "CODCOLIGADA",

    "CODFILIAL",

    "FILIAL",

    "CHAPA",

    "COLABORADOR",

    "CODFUNCAO",

    "FUNCAO",

    "CODSITUACAO",

    "CODSECAO",

    "SECAO",

]



COLUNAS_TEXTO = [

    "CHAPA",

    "CODFUNCAO",

    "CODSECAO",

]



ARQUIVO_TEMPORARIO_FOLHA = (

    PASTA_BRONZE

    / "Colaboradores_Folha_RM_novo.xlsx"

)



ARQUIVO_BRONZE_FOLHA = (

    PASTA_BRONZE

    / "Colaboradores_Folha_RM.xlsx"

)



VISAO_FOLHA = (

    "CRB.99.06 - CODCOLIGADA - CHAPA - COLABORADOR"

)



COLUNAS_ESPERADAS_FOLHA = [

    "CODCOLIGADA",

    "CODPESSOA",

    "CHAPA",

    "NOME",

    "CODSITUACAO",

    "CODSECAO",

    "DATAADMISSAO",

    "DATADEMISSAO",

]



COLUNAS_TEXTO_FOLHA = [

    "CHAPA",

    "CODSITUACAO",

    "CODSECAO",

]



TIMEOUT_LOGIN = 60

TIMEOUT_MAIN = 90

TIMEOUT_NAVEGACAO = 30

TIMEOUT_EXPORTACAO = 60



BM_CLICK = 0x00F5

BM_GETCHECK = 0x00F0





# ============================================================================

# TEXTO / ESPERA

# ============================================================================



def normalizar(valor) -> str:
    texto = str(valor or "").strip()

    texto = unicodedata.normalize(
        "NFKD",
        texto,
    )

    texto = "".join(
        caractere
        for caractere in texto
        if not unicodedata.combining(caractere)
    )

    # Remove marcador de tecla de atalho do Windows.
    # "Avanca&r >" -> "avancar >"
    texto = texto.replace("&", "")

    return texto.lower().strip()





def aguardar(
    funcao,
    timeout: float,
    intervalo: float = 0.25,
):
    inicio = time.time()
    ultimo_erro = None

    while time.time() - inicio < timeout:
        try:
            resultado = funcao()

            if resultado:
                return resultado

        except Exception as erro:
            ultimo_erro = erro

        time.sleep(intervalo)

    if ultimo_erro is not None:
        print(
            "Ultimo erro durante espera:",
            repr(ultimo_erro),
        )

    return None





# ============================================================================

# PROCESSO RM

# ============================================================================



def listar_pids_rm() -> set[int]:
    pids = set()

    for processo in psutil.process_iter(
        ["pid", "name"]
    ):
        try:
            if (
                (processo.info["name"] or "").lower()
                == "rm.exe"
            ):
                pids.add(
                    processo.info["pid"]
                )

        except Exception:
            pass

    return pids





def garantir_rm_fechado():
    pids = listar_pids_rm()

    if pids:
        raise RuntimeError(
            "Ja existe uma instancia do TOTVS RM aberta. "
            "Feche o RM completamente antes de executar."
        )





def abrir_rm() -> int:
    if not RM_EXE.exists():
        raise FileNotFoundError(
            f"RM.exe nao encontrado: {RM_EXE}"
        )

    print("Abrindo TOTVS RM...")

    processo = subprocess.Popen(
        [
            str(RM_EXE),
            RM_ARGUMENTO,
        ]
    )

    print(
        f"Processo iniciado. PID inicial: "
        f"{processo.pid}"
    )

    return processo.pid





# ============================================================================

# UI AUTOMATION

# ============================================================================



def desktop_uia():
    return Desktop(
        backend="uia"
    )





def localizar_janela_uia_auto_id(
    auto_id: str,
):
    for janela in desktop_uia().windows():
        try:
            if (
                janela.element_info.automation_id
                == auto_id
            ):
                return janela

        except Exception:
            pass

    return None





def buscar_descendente_auto_id(
    wrapper,
    auto_id: str,
    tipo: str | None = None,
):
    for controle in wrapper.descendants():
        try:
            info = controle.element_info

            if (
                info.automation_id
                != auto_id
            ):
                continue

            if (
                tipo is not None
                and
                info.control_type
                != tipo
            ):
                continue

            return controle

        except Exception:
            pass

    return None





def buscar_descendente_texto(
    wrapper,
    texto: str,
    tipos: set[str] | None = None,
):
    alvo = normalizar(texto)

    for controle in wrapper.descendants():
        try:
            info = controle.element_info

            if (
                tipos is not None
                and
                info.control_type
                not in tipos
            ):
                continue

            if (
                normalizar(
                    controle.window_text()
                )
                == alvo
            ):
                return controle

        except Exception:
            pass

    return None





def buscar_global_texto(
    texto: str,
    tipos: set[str] | None = None,
):
    alvo = normalizar(texto)

    for janela in desktop_uia().windows():

        controles = [janela]

        try:
            controles += janela.descendants()
        except Exception:
            pass

        for controle in controles:
            try:
                info = controle.element_info

                if (
                    tipos is not None
                    and
                    info.control_type
                    not in tipos
                ):
                    continue

                if (
                    normalizar(
                        controle.window_text()
                    )
                    != alvo
                ):
                    continue

                if not controle.is_visible():
                    continue

                if not controle.is_enabled():
                    continue

                return controle

            except Exception:
                pass

    return None





def acionar_uia(
    controle,
    descricao: str,
):
    erros = []

    try:
        controle.invoke()
        return

    except Exception as erro:
        erros.append(
            repr(erro)
        )

    try:
        controle.select()
        return

    except Exception as erro:
        erros.append(
            repr(erro)
        )

    try:
        controle.iface_invoke.Invoke()
        return

    except Exception as erro:
        erros.append(
            repr(erro)
        )

    raise RuntimeError(
        f"Nao foi possivel acionar "
        f"{descricao}. "
        f"Erros: {erros}"
    )





# ============================================================================

# WIN32

# ============================================================================



def localizar_janela_win32(
    titulo_alvo: str,
    classe_alvo: str | None = None,
):
    alvo = normalizar(
        titulo_alvo
    )

    encontrados = []

    def callback(hwnd, _):
        try:
            titulo = (
                win32gui.GetWindowText(hwnd)
                or ""
            )

            classe = (
                win32gui.GetClassName(hwnd)
                or ""
            )

            if (
                normalizar(titulo)
                != alvo
            ):
                return True

            if (
                classe_alvo is not None
                and
                classe != classe_alvo
            ):
                return True

            encontrados.append(hwnd)

        except Exception:
            pass

        return True

    win32gui.EnumWindows(
        callback,
        None,
    )

    if encontrados:
        return encontrados[0]

    return None





def buscar_descendente_win32_texto(
    hwnd_pai: int,
    texto_alvo: str,
):
    alvo = normalizar(
        texto_alvo
    )

    encontrados = []

    def callback(hwnd, _):
        try:
            titulo = (
                win32gui.GetWindowText(hwnd)
                or ""
            )

            if (
                normalizar(titulo)
                == alvo
            ):
                encontrados.append(hwnd)

        except Exception:
            pass

        return True

    win32gui.EnumChildWindows(
        hwnd_pai,
        callback,
        None,
    )

    if encontrados:
        return encontrados[0]

    return None





def buscar_descendente_win32_id(
    hwnd_pai: int,
    ctrl_id_alvo: int,
    classe_contem: str | None = None,
):
    encontrados = []

    def callback(hwnd, _):
        try:
            ctrl_id = (
                win32gui.GetDlgCtrlID(hwnd)
            )

            classe = (
                win32gui.GetClassName(hwnd)
                or ""
            )

            if ctrl_id != ctrl_id_alvo:
                return True

            if (
                classe_contem
                and
                classe_contem.lower()
                not in classe.lower()
            ):
                return True

            encontrados.append(hwnd)

        except Exception:
            pass

        return True

    win32gui.EnumChildWindows(
        hwnd_pai,
        callback,
        None,
    )

    if encontrados:
        return encontrados[0]

    return None





def localizar_sql_win32(
    hwnd_main: int,
):
    encontrados = []

    def callback(hwnd, _):
        try:
            titulo = (
                win32gui.GetWindowText(hwnd)
                or ""
            )

            if (
                normalizar(titulo)
                == "execucao de sentenca sql"
            ):
                encontrados.append(hwnd)

        except Exception:
            pass

        return True

    win32gui.EnumChildWindows(
        hwnd_main,
        callback,
        None,
    )

    if encontrados:
        return encontrados[0]

    return None





# ============================================================================

# LOGIN

# ============================================================================



def efetuar_login(
    usuario: str,
    senha: str,
):
    print(
        "Aguardando tela de login..."
    )

    login = aguardar(
        lambda: localizar_janela_uia_auto_id(
            "RMSFormLoginUX"
        ),
        TIMEOUT_LOGIN,
    )

    if login is None:
        raise RuntimeError(
            "Tela de login do RM nao apareceu."
        )

    usuario_edit = (
        buscar_descendente_auto_id(
            login,
            "cUser",
            "Edit",
        )
    )

    senha_edit = (
        buscar_descendente_auto_id(
            login,
            "cPassword",
            "Edit",
        )
    )

    entrar = (
        buscar_descendente_auto_id(
            login,
            "btnOk",
            "Button",
        )
    )

    if usuario_edit is None:
        raise RuntimeError(
            "Campo cUser nao encontrado."
        )

    if senha_edit is None:
        raise RuntimeError(
            "Campo cPassword nao encontrado."
        )

    if entrar is None:
        raise RuntimeError(
            "Botao btnOk nao encontrado."
        )

    usuario_edit.set_edit_text(
        usuario
    )

    senha_edit.set_edit_text(
        senha
    )

    print(
        "Efetuando login..."
    )

    acionar_uia(
        entrar,
        "Entrar",
    )





# ============================================================================

# MAINFORM

# ============================================================================



def aguardar_mainform():
    print(
        "Aguardando MainForm..."
    )

    main = aguardar(
        lambda: localizar_janela_uia_auto_id(
            "MainForm"
        ),
        TIMEOUT_MAIN,
    )

    if main is None:
        raise RuntimeError(
            "MainForm do RM nao apareceu."
        )

    print(
        "MainForm encontrado."
    )

    return main





# ============================================================================

# NAVEGACAO

# ============================================================================



def abrir_gestao_pessoas(

    main,

):

    print(

        "Verificando modulo Gestao de Pessoas..."

    )



    pid_rm = (

        main.element_info.process_id

    )



    def localizar_no_pid(

        titulo,

        tipo=None,

    ):

        desktop = Desktop(

            backend="uia"

        )



        for janela in desktop.windows():

            try:

                if (

                    janela.element_info.process_id

                    != pid_rm

                ):

                    continue

            except Exception:

                continue



            controles = [janela]



            try:

                controles += (

                    janela.descendants()

                )

            except Exception:

                pass



            for controle in controles:

                try:

                    info = (

                        controle.element_info

                    )



                    if (

                        info.process_id

                        != pid_rm

                    ):

                        continue



                    if (

                        controle.window_text()

                        != titulo

                    ):

                        continue



                    if (

                        tipo is not None

                        and

                        info.control_type

                        != tipo

                    ):

                        continue



                    if not controle.is_visible():

                        continue



                    if not controle.is_enabled():

                        continue



                    return controle



                except Exception:

                    pass



        return None



    def modulo_gestao_carregado():

        alvos = {

            "planejamento e captacao",

            "gestao organizacional",

            "gestao",

        }



        encontrados = set()



        try:

            controles = (

                main.descendants()

            )

        except Exception:

            controles = []



        for controle in controles:

            try:

                info = (

                    controle.element_info

                )



                if (

                    info.control_type

                    != "TabItem"

                ):

                    continue



                if not controle.is_visible():

                    continue



                titulo = normalizar(

                    controle.window_text()

                )



                if titulo in alvos:

                    encontrados.add(

                        titulo

                    )



            except Exception:

                pass



        return (

            encontrados == alvos

        )



    if modulo_gestao_carregado():

        print(

            "Gestao de Pessoas ja esta ativa."

        )

        return



    print(

        "Gestao de Pessoas ainda nao esta ativa."

    )



    print(

        "Abrindo seletor Application..."

    )



    application = (

        buscar_descendente_texto(

            main,

            "Application",

            {"SplitButton"},

        )

    )



    if application is None:

        raise RuntimeError(

            "SplitButton Application "

            "nao encontrado."

        )



    try:

        main.set_focus()

    except Exception:

        pass


    try:

        application.set_focus()

    except Exception:

        pass


    try:

        application.click_input()

    except Exception as erro:

        raise RuntimeError(

            "Nao foi possivel abrir o seletor Application "
            "com click_input()."

        ) from erro


    time.sleep(0.5)



    rh = aguardar(

        lambda: localizar_no_pid(

            "btnSegmento5RH",

            "MenuItem",

        ),

        TIMEOUT_NAVEGACAO,

    )



    if rh is None:

        raise RuntimeError(

            "btnSegmento5RH nao apareceu "

            "apos abrir Application."

        )



    print(

        "Abrindo segmento RH..."

    )



    expandido = False



    try:

        rh.expand()

        expandido = True

    except Exception:

        pass



    if not expandido:

        try:

            rh.iface_expand_collapse.Expand()

            expandido = True

        except Exception:

            pass



    if not expandido:

        acionar_uia(

            rh,

            "btnSegmento5RH",

        )



    gestao_pessoas = aguardar(

        lambda: localizar_no_pid(

            "btnModuleGESTAODEPESSOAS",

            "Button",

        ),

        TIMEOUT_NAVEGACAO,

    )



    if gestao_pessoas is None:

        raise RuntimeError(

            "btnModuleGESTAODEPESSOAS "

            "nao apareceu."

        )



    print(

        "Abrindo Gestao de Pessoas..."

    )



    acionar_uia(

        gestao_pessoas,

        "btnModuleGESTAODEPESSOAS",

    )



    carregado = aguardar(

        modulo_gestao_carregado,

        TIMEOUT_NAVEGACAO,

    )



    if not carregado:

        raise RuntimeError(

            "Gestao de Pessoas nao foi "

            "confirmada apos a selecao."

        )



    print(

        "Gestao de Pessoas carregada."

    )





def localizar_controle_no_pid(

    main,

    titulo: str,

    tipos: set[str] | None = None,

):

    pid_rm = main.element_info.process_id

    alvo = normalizar(titulo)

    desktop = Desktop(

        backend="uia"

    )

    for janela in desktop.windows():

        try:

            if (

                janela.element_info.process_id

                != pid_rm

            ):

                continue

        except Exception:

            continue

        controles = [janela]

        try:

            controles += janela.descendants()

        except Exception:

            pass

        for controle in controles:

            try:

                info = controle.element_info

                if info.process_id != pid_rm:

                    continue

                if (

                    tipos is not None

                    and

                    info.control_type not in tipos

                ):

                    continue

                if (

                    normalizar(

                        controle.window_text()

                    )

                    != alvo

                ):

                    continue

                if not controle.is_visible():

                    continue

                if not controle.is_enabled():

                    continue

                return controle

            except Exception:

                continue

    return None



def localizar_controle_no_pid_contem(

    main,

    termos: tuple[str, ...],

    tipos: set[str] | None = None,

):

    pid_rm = main.element_info.process_id

    termos_normalizados = tuple(

        normalizar(termo)

        for termo in termos

    )

    desktop = Desktop(

        backend="uia"

    )

    for janela in desktop.windows():

        try:

            if (

                janela.element_info.process_id

                != pid_rm

            ):

                continue

        except Exception:

            continue

        controles = [janela]

        try:

            controles += janela.descendants()

        except Exception:

            pass

        for controle in controles:

            try:

                info = controle.element_info

                if info.process_id != pid_rm:

                    continue

                if (

                    tipos is not None

                    and

                    info.control_type not in tipos

                ):

                    continue

                texto = normalizar(

                    controle.window_text()

                )

                auto_id = normalizar(

                    info.automation_id

                )

                universo = f"{texto} {auto_id}"

                if not all(

                    termo in universo

                    for termo in termos_normalizados

                ):

                    continue

                if not controle.is_visible():

                    continue

                if not controle.is_enabled():

                    continue

                return controle

            except Exception:

                continue

    return None



def modulo_folha_carregado(

    main,

) -> bool:

    alvos = {

        "administracao de pessoal",

        "folha mensal",

        "ferias",

        "rescisao",

        "gestao",

    }

    encontrados = set()

    try:

        controles = main.descendants()

    except Exception:

        controles = []

    for controle in controles:

        try:

            info = controle.element_info

            if info.control_type != "TabItem":

                continue

            if not controle.is_visible():

                continue

            titulo = normalizar(

                controle.window_text()

            )

            if titulo in alvos:

                encontrados.add(titulo)

        except Exception:

            continue

    outros = alvos - {"gestao"}

    return (

        "gestao" in encontrados

        and

        bool(encontrados & outros)

    )



def abrir_folha_pagamento(

    main,

):

    print(

        "Verificando modulo Folha de Pagamento..."

    )

    if modulo_folha_carregado(main):

        print(

            "Folha de Pagamento ja esta ativa."

        )

        return

    print(

        "Trocando para Folha de Pagamento..."

    )

    application = buscar_descendente_texto(

        main,

        "Application",

        {"SplitButton"},

    )

    if application is None:

        raise RuntimeError(

            "SplitButton Application nao encontrado."

        )

    try:

        main.set_focus()

    except Exception:

        pass


    try:

        application.set_focus()

    except Exception:

        pass


    try:

        application.click_input()

    except Exception as erro:

        raise RuntimeError(

            "Nao foi possivel abrir o seletor Application "
            "com click_input()."

        ) from erro


    time.sleep(0.5)

    rh = aguardar(

        lambda: localizar_controle_no_pid(

            main,

            "btnSegmento5RH",

            {"MenuItem"},

        ),

        TIMEOUT_NAVEGACAO,

    )

    if rh is None:

        rh = aguardar(

            lambda: localizar_controle_no_pid_contem(

                main,

                ("segmento", "rh"),

                {"MenuItem", "Button"},

            ),

            5,

        )

    if rh is None:

        raise RuntimeError(

            "Segmento RH nao encontrado."

        )

    print(

        "Abrindo segmento RH..."

    )

    expandido = False

    try:

        rh.expand()

        expandido = True

    except Exception:

        pass

    if not expandido:

        try:

            rh.iface_expand_collapse.Expand()

            expandido = True

        except Exception:

            pass

    if not expandido:

        acionar_uia(

            rh,

            "segmento RH",

        )

    folha = aguardar(

        lambda: localizar_controle_no_pid_contem(

            main,

            ("folha", "pagamento"),

            {"Button", "MenuItem"},

        ),

        TIMEOUT_NAVEGACAO,

    )

    if folha is None:

        raise RuntimeError(

            "Modulo Folha de Pagamento nao encontrado."

        )

    print(

        "Abrindo Folha de Pagamento..."

    )

    acionar_uia(

        folha,

        "Folha de Pagamento",

    )

    carregado = aguardar(

        lambda: modulo_folha_carregado(main),

        TIMEOUT_NAVEGACAO,

    )

    if not carregado:

        raise RuntimeError(

            "Folha de Pagamento nao foi "

            "confirmada apos a selecao."

        )

    print(

        "Folha de Pagamento carregada."

    )


def abrir_visao_folha(

    main,

):

    print(

        "Selecionando aba Gestao..."

    )

    aba = buscar_descendente_texto(

        main,

        "Gestao",

        {"TabItem"},

    )

    if aba is None:

        raise RuntimeError(

            "Aba Gestao nao encontrada na Folha."

        )

    try:

        aba.select()

    except Exception:

        acionar_uia(

            aba,

            "aba Gestao",

        )

    time.sleep(1)

    visoes = buscar_descendente_texto(

        main,

        "Visoes de Dados",

        {"SplitButton"},

    )

    if visoes is None:

        expandir = buscar_descendente_texto(

            main,

            "Expand Collapse",

            {"Button"},

        )

        if expandir is not None:

            print(

                "Expandindo Ribbon..."

            )

            acionar_uia(

                expandir,

                "Expand Collapse",

            )

            time.sleep(1)

        visoes = buscar_descendente_texto(

            main,

            "Visoes de Dados",

            {"SplitButton"},

        )

    if visoes is None:

        raise RuntimeError(

            "Visoes de Dados nao encontrado na Folha."

        )

    print(

        "Abrindo Visoes de Dados..."

    )

    abrir = buscar_descendente_texto(

        visoes,

        "Abrir",

        {"MenuItem"},

    )

    if abrir is None:

        raise RuntimeError(

            "MenuItem Abrir de Visoes de Dados "

            "nao encontrado na Folha."

        )

    acionar_uia(

        abrir,

        "Abrir Visoes de Dados",

    )

    alvo = aguardar(

        lambda: localizar_controle_no_pid(

            main,

            VISAO_FOLHA,

            {"Button", "MenuItem", "ListItem"},

        ),

        TIMEOUT_NAVEGACAO,

    )

    if alvo is None:

        raise RuntimeError(

            "Visao da Folha nao encontrada: "

            f"{VISAO_FOLHA}"

        )

    print(

        "Abrindo visao:"

    )

    print(

        VISAO_FOLHA

    )

    acionar_uia(

        alvo,

        VISAO_FOLHA,

    )


def abrir_pfunc(
    main,
):
    print(
        "Selecionando aba Gestao..."
    )

    aba = (
        buscar_descendente_texto(
            main,
            "Gestao",
            {"TabItem"},
        )
    )

    if aba is None:
        raise RuntimeError(
            "Aba Gestao nao encontrada."
        )

    try:
        aba.select()
    except Exception:
        acionar_uia(
            aba,
            "aba Gestao",
        )

    time.sleep(1)

    visoes = (
        buscar_descendente_texto(
            main,
            "Visoes de Dados",
            {"SplitButton"},
        )
    )

    if visoes is None:
        expandir = (
            buscar_descendente_texto(
                main,
                "Expand Collapse",
                {"Button"},
            )
        )

        if expandir is not None:
            print(
                "Expandindo Ribbon..."
            )

            acionar_uia(
                expandir,
                "Expand Collapse",
            )

            time.sleep(1)

        visoes = (
            buscar_descendente_texto(
                main,
                "Visoes de Dados",
                {"SplitButton"},
            )
        )

    if visoes is None:
        raise RuntimeError(
            "Visoes de Dados nao encontrado."
        )

    print(
        "Abrindo Visoes de Dados..."
    )

    abrir = (
        buscar_descendente_texto(
            visoes,
            "Abrir",
            {"MenuItem"},
        )
    )

    if abrir is None:
        raise RuntimeError(
            "MenuItem Abrir de "
            "Visoes de Dados nao encontrado."
        )

    acionar_uia(
        abrir,
        "Abrir Visoes de Dados",
    )

    pfunc = aguardar(
        lambda: buscar_global_texto(
            "PFUNC - Funcionarios",
            {"Button"},
        ),
        TIMEOUT_NAVEGACAO,
    )

    if pfunc is None:
        raise RuntimeError(
            "PFUNC - Funcionarios "
            "nao encontrado."
        )

    print(
        "Abrindo PFUNC - Funcionarios..."
    )

    acionar_uia(
        pfunc,
        "PFUNC - Funcionarios",
    )





# ============================================================================

# SQL

# ============================================================================



def abrir_sql(
    main,
):
    print(
        "Aguardando Execucao de Sentenca SQL..."
    )

    hwnd_sql = aguardar(
        lambda: localizar_sql_win32(
            main.handle
        ),
        TIMEOUT_NAVEGACAO,
    )

    if hwnd_sql is None:
        raise RuntimeError(
            "Execucao de Sentenca SQL "
            "nao encontrada."
        )

    sql = UIAWrapper(
        UIAElementInfo(
            hwnd_sql
        )
    )

    print(
        "Janela SQL encontrada."
    )

    return sql





def executar_consulta(
    sql,
):
    executar = (
        buscar_descendente_texto(
            sql,
            "buttonExecute",
        )
    )

    if executar is None:
        raise RuntimeError(
            "buttonExecute nao encontrado."
        )

    print(
        "Executando consulta..."
    )

    acionar_uia(
        executar,
        "buttonExecute",
    )

    grid = aguardar(
        lambda: buscar_descendente_auto_id(
            sql,
            "gridResult",
            "Table",
        ),
        TIMEOUT_NAVEGACAO,
    )

    if grid is None:
        raise RuntimeError(
            "gridResult nao encontrado."
        )

    # Espera a grade terminar de popular.
    time.sleep(3)

    print(
        "Consulta executada."
    )





# ============================================================================

# EXPORTACAO XLSX

# ============================================================================



def abrir_assistente_exportacao(
    sql,
):
    exportar = (
        buscar_descendente_texto(
            sql,
            "buttonExport",
        )
    )

    if exportar is None:
        raise RuntimeError(
            "buttonExport nao encontrado."
        )

    print(
        "Abrindo menu Exportar..."
    )

    acionar_uia(
        exportar,
        "buttonExport",
    )

    xlsx = aguardar(
        lambda: buscar_global_texto(
            "Exportar para XLSX",
            {"Button"},
        ),
        TIMEOUT_NAVEGACAO,
    )

    if xlsx is None:
        raise RuntimeError(
            "Exportar para XLSX "
            "nao encontrado."
        )

    print(
        "Selecionando Exportar para XLSX..."
    )

    # Usamos o HWND do botao e PostMessage para evitar
    # bloqueio da chamada UIA quando o TOTVS abre o modal.
    try:
        hwnd_xlsx = xlsx.handle
    except Exception:
        hwnd_xlsx = 0

    if hwnd_xlsx:
        win32gui.PostMessage(
            hwnd_xlsx,
            BM_CLICK,
            0,
            0,
        )
    else:
        # Fallback. Algumas versoes podem nao expor HWND.
        try:
            xlsx.invoke()
        except Exception:
            pass

    hwnd_assistente = aguardar(
        lambda: localizar_janela_win32(
            "Exportacao de Documentos"
        ),
        TIMEOUT_NAVEGACAO,
    )

    if hwnd_assistente is None:
        raise RuntimeError(
            "Exportacao de Documentos "
            "nao apareceu."
        )

    print(
        "Assistente de exportacao aberto."
    )

    return hwnd_assistente





def avancar_assistente(
    hwnd_assistente: int,
):
    checkbox = (
        buscar_descendente_win32_texto(
            hwnd_assistente,
            "Exportacao com reconhecimento de dados",
        )
    )

    if checkbox:
        estado = win32gui.SendMessage(
            checkbox,
            BM_GETCHECK,
            0,
            0,
        )

        print(
            "Reconhecimento de dados:",
            estado,
        )

        # Mantemos desmarcado, conforme o teste validado.
        if estado == 1:
            print(
                "Desmarcando reconhecimento "
                "de dados..."
            )

            win32gui.SendMessage(
                checkbox,
                BM_CLICK,
                0,
                0,
            )

            time.sleep(0.5)

    avancar = (
        buscar_descendente_win32_texto(
            hwnd_assistente,
            "Avancar >",
        )
    )

    if avancar is None:
        raise RuntimeError(
            "Botao Avancar nao encontrado."
        )

    print(
        "Acionando Avancar..."
    )

    win32gui.PostMessage(
        avancar,
        BM_CLICK,
        0,
        0,
    )





# ============================================================================

# SALVAR COMO

# ============================================================================



def aguardar_salvar_como():
    print(
        "Aguardando Salvar como..."
    )

    hwnd_salvar = aguardar(
        lambda: localizar_janela_win32(
            "Salvar como",
            "#32770",
        ),
        TIMEOUT_NAVEGACAO,
    )

    if hwnd_salvar is None:
        raise RuntimeError(
            "Salvar como nao apareceu."
        )

    print(
        "Salvar como encontrado."
    )

    return hwnd_salvar





def preencher_nome_arquivo(

    hwnd_salvar: int,

    caminho: Path,

):

    print(

        "Preenchendo campo Nome "

        "em processo UIA isolado..."

    )



    caminho_texto = str(

        caminho.resolve()

    )



    print(

        "Arquivo temporario:"

    )



    print(

        caminho_texto

    )



    # O dialogo Salvar como pertence a uma

    # janela modal nativa do Windows.

    #

    # Durante a automacao completa do RM,

    # o provedor UIA do processo principal

    # pode deixar de aceitar ElementFromHandle

    # para esse Edit.

    #

    # Um processo Python novo foi validado

    # com sucesso para esse mesmo controle.

    helper = r"""

import sys

import time



import win32gui



from pywinauto.controls.uiawrapper import UIAWrapper

from pywinauto.uia_element_info import UIAElementInfo





hwnd_salvar = int(

    sys.argv[1]

)



caminho = sys.argv[2]



fim = time.time() + 20



ultimo_erro = None

ultimo_hwnd = None





def localizar_nome():

    encontrados = []



    def callback(hwnd, _):

        try:

            ctrl_id = (

                win32gui.GetDlgCtrlID(

                    hwnd

                )

            )



            classe = (

                win32gui.GetClassName(

                    hwnd

                )

                or ""

            )



            if (

                ctrl_id == 1001

                and

                classe.lower() == "edit"

            ):

                encontrados.append(

                    hwnd

                )



        except Exception:

            pass



        return True



    win32gui.EnumChildWindows(

        hwnd_salvar,

        callback,

        None,

    )



    return (

        encontrados[0]

        if encontrados

        else None

    )





while time.time() < fim:



    if not win32gui.IsWindow(

        hwnd_salvar

    ):

        print(

            "Janela Salvar como deixou "

            "de existir.",

            file=sys.stderr,

        )



        raise SystemExit(10)



    hwnd_nome = localizar_nome()



    if hwnd_nome is None:

        time.sleep(0.25)

        continue



    ultimo_hwnd = hwnd_nome



    try:

        campo = UIAWrapper(

            UIAElementInfo(

                hwnd_nome

            )

        )



        try:

            campo.set_edit_text(

                caminho

            )



        except Exception:

            campo.iface_value.SetValue(

                caminho

            )



        try:

            valor = (

                campo.get_value()

            )



        except Exception:

            valor = (

                campo

                .iface_value

                .CurrentValue

            )



        if valor == caminho:



            print(

                "OK"

            )



            print(

                repr(valor)

            )



            raise SystemExit(0)



        ultimo_erro = (

            "Valor lido diferente: "

            + repr(valor)

        )



    except SystemExit:

        raise



    except Exception as erro:

        ultimo_erro = repr(

            erro

        )



    time.sleep(0.3)





print(

    "Falha ao preencher campo Nome. "

    f"Ultimo HWND: {ultimo_hwnd}. "

    f"Ultimo erro: {ultimo_erro}",

    file=sys.stderr,

)



raise SystemExit(20)

"""



    try:

        resultado = subprocess.run(

            [

                sys.executable,

                "-c",

                helper,

                str(hwnd_salvar),

                caminho_texto,

            ],

            capture_output=True,

            text=True,

            timeout=30,

        )



    except subprocess.TimeoutExpired as erro:

        raise RuntimeError(

            "Subprocesso responsavel pelo "

            "campo Nome excedeu o tempo limite."

        ) from erro



    if resultado.returncode != 0:



        detalhe = (

            resultado.stderr.strip()

            or

            resultado.stdout.strip()

            or

            "sem detalhe"

        )



        raise RuntimeError(

            "Nao foi possivel preencher "

            "o campo Nome no processo "

            "UIA isolado. "

            f"Detalhe: {detalhe}"

        )



    print(

        "Campo Nome preenchido e "

        "confirmado com sucesso."

    )



    if resultado.stdout.strip():

        print(

            resultado.stdout.strip()

        )





def salvar_xlsx(
    hwnd_salvar: int,
):
    print(
        "Aguardando botao Salvar..."
    )

    def localizar_botao_salvar():
        hwnd_atual = (
            localizar_janela_win32(
                "Salvar como",
                "#32770",
            )
        )

        if hwnd_atual is None:
            return None

        return (
            buscar_descendente_win32_id(
                hwnd_atual,
                1,
                "button",
            )
        )

    botao_salvar = aguardar(
        localizar_botao_salvar,
        15,
        0.2,
    )

    if botao_salvar is None:
        raise RuntimeError(
            "Botao Salvar "
            "(CTRL_ID=1) "
            "nao apareceu em 15 segundos."
        )

    print(
        f"Botao Salvar encontrado. "
        f"HWND: {botao_salvar}"
    )

    print(
        "Acionando Salvar..."
    )

    win32gui.PostMessage(
        botao_salvar,
        BM_CLICK,
        0,
        0,
    )





def responder_nao_abrir():
    hwnd_info = aguardar(
        lambda: localizar_janela_win32(
            "Informacao"
        ),
        10,
    )

    if hwnd_info is None:
        print(
            "Aviso: confirmacao para abrir "
            "o arquivo nao apareceu."
        )
        return

    botao_nao = (
        buscar_descendente_win32_texto(
            hwnd_info,
            "Nao",
        )
    )

    if botao_nao is None:
        print(
            "Aviso: botao Nao nao encontrado."
        )
        return

    print(
        "Respondendo Nao para abrir "
        "o arquivo gerado..."
    )

    win32gui.PostMessage(
        botao_nao,
        BM_CLICK,
        0,
        0,
    )





# ============================================================================

# ARQUIVO

# ============================================================================



def esperar_arquivo_estavel(
    caminho: Path,
    timeout: float,
):
    inicio = time.time()
    tamanho_anterior = None
    repeticoes = 0

    while time.time() - inicio < timeout:

        if caminho.exists():
            tamanho = (
                caminho.stat().st_size
            )

            if (
                tamanho > 0
                and
                tamanho == tamanho_anterior
            ):
                repeticoes += 1
            else:
                repeticoes = 0

            tamanho_anterior = tamanho

            if repeticoes >= 2:
                return True

        time.sleep(1)

    return False





# ============================================================================

# CONTROLE ENTRE AS DUAS EXTRACOES

# ============================================================================



def fechar_janela_sql(

    sql,

):

    print(

        "Fechando janela da consulta "

        "antes de trocar de modulo..."

    )

    hwnd = getattr(

        sql,

        "handle",

        None,

    )

    if hwnd is None:

        return

    try:

        if not win32gui.IsWindow(hwnd):

            return

    except Exception:

        return

    WM_CLOSE = 0x0010

    win32gui.PostMessage(

        hwnd,

        WM_CLOSE,

        0,

        0,

    )

    inicio = time.time()

    while time.time() - inicio < 15:

        try:

            if not win32gui.IsWindow(hwnd):

                print(

                    "Janela SQL fechada."

                )

                time.sleep(1)

                return

        except Exception:

            return

        time.sleep(0.25)

    raise RuntimeError(

        "Janela SQL permaneceu aberta "

        "apos tentativa de fechamento."

    )



def exportar_consulta_atual(

    sql,

    caminho_temporario: Path,

    rotulo: str,

):

    executar_consulta(

        sql

    )

    hwnd_assistente = (

        abrir_assistente_exportacao(

            sql

        )

    )

    avancar_assistente(

        hwnd_assistente

    )

    hwnd_salvar = (

        aguardar_salvar_como()

    )

    preencher_nome_arquivo(

        hwnd_salvar,

        caminho_temporario,

    )

    salvar_xlsx(

        hwnd_salvar

    )

    responder_nao_abrir()

    print(

        f"Aguardando arquivo exportado - {rotulo}..."

    )

    if not esperar_arquivo_estavel(

        caminho_temporario,

        TIMEOUT_EXPORTACAO,

    ):

        raise RuntimeError(

            f"Arquivo temporario de {rotulo} "

            "nao apareceu ou nao estabilizou."

        )

    print(

        f"Arquivo exportado - {rotulo}:"

    )

    print(

        caminho_temporario

    )

    print(

        "Tamanho:",

        caminho_temporario.stat().st_size,

        "bytes",

    )


# ============================================================================

# VALIDACAO

# ============================================================================



def validar_exportacao(
    caminho: Path,
) -> int:
    print()
    print("=" * 90)
    print("VALIDANDO EXPORTACAO")
    print("=" * 90)

    wb = load_workbook(
        caminho,
        read_only=True,
        data_only=True,
    )

    try:
        ws = wb.active

        primeira_linha = next(
            ws.iter_rows(
                min_row=1,
                max_row=1,
                values_only=True,
            )
        )

        cabecalho = [
            str(valor).strip()
            if valor is not None
            else ""
            for valor in primeira_linha
        ]

        print(
            "Cabecalho:"
        )
        print(
            cabecalho
        )

        if cabecalho != COLUNAS_ESPERADAS:
            raise RuntimeError(
                "Cabecalho diferente "
                "do contrato esperado."
            )

        indices = {
            coluna: cabecalho.index(coluna)
            for coluna in COLUNAS_ESPERADAS
        }

        total = 0
        chaves = set()
        duplicadas = []
        problemas_tipo = []

        tipos = {
            coluna: Counter()
            for coluna in COLUNAS_TEXTO
        }

        exemplos = {
            coluna: []
            for coluna in COLUNAS_TEXTO
        }

        for numero_linha, linha in enumerate(
            ws.iter_rows(
                min_row=2,
                values_only=True,
            ),
            start=2,
        ):
            if not any(
                valor is not None
                and
                str(valor).strip() != ""
                for valor in linha
            ):
                continue

            total += 1

            codcoligada = linha[
                indices["CODCOLIGADA"]
            ]

            codfilial = linha[
                indices["CODFILIAL"]
            ]

            chapa = linha[
                indices["CHAPA"]
            ]

            if (
                codcoligada is None
                or
                codfilial is None
                or
                chapa is None
            ):
                raise RuntimeError(
                    "Chave RM incompleta "
                    f"na linha {numero_linha}."
                )

            chave = (
                str(codcoligada).strip(),
                str(codfilial).strip(),
                str(chapa).strip(),
            )

            if chave in chaves:
                duplicadas.append(
                    (
                        numero_linha,
                        chave,
                    )
                )
            else:
                chaves.add(chave)

            for coluna in COLUNAS_TEXTO:
                valor = linha[
                    indices[coluna]
                ]

                if valor is None:
                    continue

                tipo = type(valor).__name__

                tipos[coluna][tipo] += 1

                if (
                    len(exemplos[coluna])
                    < 5
                ):
                    exemplos[coluna].append(
                        repr(valor)
                    )

                if not isinstance(
                    valor,
                    str,
                ):
                    problemas_tipo.append(
                        (
                            numero_linha,
                            coluna,
                            valor,
                            tipo,
                        )
                    )

        if total <= 0:
            raise RuntimeError(
                "Exportacao sem registros."
            )

        if duplicadas:
            raise RuntimeError(
                "CODCOLIGADA + CODFILIAL + "
                "CHAPA possui duplicidade. "
                f"Exemplos: {duplicadas[:5]}"
            )

        if problemas_tipo:
            raise RuntimeError(
                "Campos textuais foram "
                "exportados com outro tipo. "
                f"Exemplos: "
                f"{problemas_tipo[:5]}"
            )

        print()
        print(
            f"Registros: {total}"
        )

        print(
            f"Chaves unicas: "
            f"{len(chaves)}"
        )

        print(
            "Cabecalho: OK"
        )

        print(
            "Unicidade: OK"
        )

        print(
            "Tipos textuais: OK"
        )

        for coluna in COLUNAS_TEXTO:
            print()
            print(coluna)

            print(
                "  Tipos:",
                dict(tipos[coluna]),
            )

            print(
                "  Exemplos:",
                exemplos[coluna],
            )

        return total

    finally:
        wb.close()







def valor_data_valido(

    valor,

) -> bool:

    if valor is None:

        return True

    if isinstance(

        valor,

        (datetime, date),

    ):

        return True

    texto = str(valor).strip()

    if not texto:

        return True

    formatos = [

        "%d/%m/%Y",

        "%Y-%m-%d",

        "%Y-%m-%d %H:%M:%S",

    ]

    for formato in formatos:

        try:

            datetime.strptime(

                texto,

                formato,

            )

            return True

        except ValueError:

            pass

    return False



def validar_exportacao_folha(

    caminho: Path,

) -> int:

    print()

    print("=" * 90)

    print(

        "VALIDANDO EXPORTACAO - FOLHA DE PAGAMENTO"

    )

    print("=" * 90)

    wb = load_workbook(

        caminho,

        read_only=True,

        data_only=True,

    )

    try:

        ws = wb.active

        primeira_linha = next(

            ws.iter_rows(

                min_row=1,

                max_row=1,

                values_only=True,

            )

        )

        cabecalho = [

            str(valor).strip()

            if valor is not None

            else ""

            for valor in primeira_linha

        ]

        print(

            "Cabecalho:"

        )

        print(

            cabecalho

        )

        if cabecalho != COLUNAS_ESPERADAS_FOLHA:

            raise RuntimeError(

                "Cabecalho da Folha diferente "

                "do contrato esperado."

            )

        indices = {

            coluna: cabecalho.index(coluna)

            for coluna in COLUNAS_ESPERADAS_FOLHA

        }

        total = 0

        chaves = set()

        duplicadas = []

        problemas_tipo = []

        problemas_data = []

        tipos = {

            coluna: Counter()

            for coluna in COLUNAS_TEXTO_FOLHA

        }

        exemplos = {

            coluna: []

            for coluna in COLUNAS_TEXTO_FOLHA

        }

        for numero_linha, linha in enumerate(

            ws.iter_rows(

                min_row=2,

                values_only=True,

            ),

            start=2,

        ):

            if not any(

                valor is not None

                and

                str(valor).strip() != ""

                for valor in linha

            ):

                continue

            total += 1

            codcoligada = linha[

                indices["CODCOLIGADA"]

            ]

            chapa = linha[

                indices["CHAPA"]

            ]

            nome = linha[

                indices["NOME"]

            ]

            if (

                codcoligada is None

                or

                chapa is None

            ):

                raise RuntimeError(

                    "Chave CODCOLIGADA + CHAPA "

                    "incompleta na Folha "

                    f"na linha {numero_linha}."

                )

            if (

                nome is None

                or

                not str(nome).strip()

            ):

                raise RuntimeError(

                    "NOME vazio na Folha "

                    f"na linha {numero_linha}."

                )

            chave = (

                str(codcoligada).strip(),

                str(chapa).strip(),

            )

            if chave in chaves:

                duplicadas.append(

                    (

                        numero_linha,

                        chave,

                    )

                )

            else:

                chaves.add(chave)

            for coluna in COLUNAS_TEXTO_FOLHA:

                valor = linha[

                    indices[coluna]

                ]

                if valor is None:

                    continue

                tipo = type(valor).__name__

                tipos[coluna][tipo] += 1

                if len(exemplos[coluna]) < 5:

                    exemplos[coluna].append(

                        repr(valor)

                    )

                if not isinstance(

                    valor,

                    str,

                ):

                    problemas_tipo.append(

                        (

                            numero_linha,

                            coluna,

                            valor,

                            tipo,

                        )

                    )

            for coluna in [

                "DATAADMISSAO",

                "DATADEMISSAO",

            ]:

                valor = linha[

                    indices[coluna]

                ]

                if not valor_data_valido(valor):

                    problemas_data.append(

                        (

                            numero_linha,

                            coluna,

                            valor,

                        )

                    )

        if total <= 0:

            raise RuntimeError(

                "Exportacao da Folha sem registros."

            )

        if duplicadas:

            raise RuntimeError(

                "CODCOLIGADA + CHAPA possui "

                "duplicidade na Folha. "

                f"Exemplos: {duplicadas[:5]}"

            )

        if problemas_tipo:

            raise RuntimeError(

                "Campos textuais da Folha foram "

                "exportados com outro tipo. "

                f"Exemplos: {problemas_tipo[:5]}"

            )

        if problemas_data:

            raise RuntimeError(

                "Datas invalidas na Folha. "

                f"Exemplos: {problemas_data[:5]}"

            )

        print()

        print(

            f"Registros: {total}"

        )

        print(

            "Chaves CODCOLIGADA + CHAPA unicas: "

            f"{len(chaves)}"

        )

        print(

            "Cabecalho: OK"

        )

        print(

            "Unicidade: OK"

        )

        print(

            "Tipos textuais: OK"

        )

        print(

            "Datas: OK"

        )

        for coluna in COLUNAS_TEXTO_FOLHA:

            print()

            print(coluna)

            print(

                "  Tipos:",

                dict(tipos[coluna]),

            )

            print(

                "  Exemplos:",

                exemplos[coluna],

            )

        return total

    finally:

        wb.close()


# ============================================================================

# PROMOCAO CONJUNTA PARA BRONZE

# ============================================================================



def promover_bronzes():

    print()

    print("=" * 90)

    print(

        "PROMOVENDO AS DUAS FONTES "

        "PARA O BRONZE OFICIAL"

    )

    print("=" * 90)

    pares = [

        (

            ARQUIVO_TEMPORARIO,

            ARQUIVO_BRONZE,

        ),

        (

            ARQUIVO_TEMPORARIO_FOLHA,

            ARQUIVO_BRONZE_FOLHA,

        ),

    ]

    backups = {}

    promovidos = []

    try:

        for _, destino in pares:

            backup = destino.with_name(

                destino.name + ".backup"

            )

            if backup.exists():

                backup.unlink()

            if destino.exists():

                shutil.copy2(

                    destino,

                    backup,

                )

                backups[destino] = backup

        for origem, destino in pares:

            os.replace(

                origem,

                destino,

            )

            promovidos.append(destino)

    except PermissionError as erro:

        for destino in promovidos:

            backup = backups.get(destino)

            try:

                if backup is not None and backup.exists():

                    shutil.copy2(

                        backup,

                        destino,

                    )

                elif destino.exists():

                    destino.unlink()

            except Exception:

                pass

        raise RuntimeError(

            "Nao foi possivel promover as duas "

            "fontes para o Bronze. Verifique se "

            "algum arquivo esta aberto no Excel."

        ) from erro

    except Exception:

        for destino in promovidos:

            backup = backups.get(destino)

            try:

                if backup is not None and backup.exists():

                    shutil.copy2(

                        backup,

                        destino,

                    )

                elif destino.exists():

                    destino.unlink()

            except Exception:

                pass

        raise

    finally:

        for backup in backups.values():

            try:

                if backup.exists():

                    backup.unlink()

            except Exception:

                pass

    print(

        "Bronze Gestao de Pessoas atualizado:"

    )

    print(

        ARQUIVO_BRONZE

    )

    print()

    print(

        "Bronze Folha de Pagamento atualizado:"

    )

    print(

        ARQUIVO_BRONZE_FOLHA

    )



# ============================================================================

# ENCERRAMENTO DO RM

# ============================================================================



def fechar_rm(
    main_rm,
):
    from pywinauto.controls.hwndwrapper import HwndWrapper

    print()
    print(
        "Fechando TOTVS RM..."
    )

    try:
        hwnd_main = main_rm.handle
    except Exception:
        print(
            "Aviso: nao foi possivel obter "
            "o HWND do MainForm."
        )
        return False

    try:
        pid_rm = (
            main_rm
            .element_info
            .process_id
        )
    except Exception:
        pid_rm = None

    if not win32gui.IsWindow(
        hwnd_main
    ):
        print(
            "RM ja estava fechado."
        )
        return True

    WM_CLOSE = 0x0010

    win32gui.PostMessage(
        hwnd_main,
        WM_CLOSE,
        0,
        0,
    )

    print(
        "Aguardando confirmacao de saida..."
    )

    def localizar_confirmacao():
        encontrados = []

        def callback(hwnd, _):
            try:
                if not win32gui.IsWindowVisible(
                    hwnd
                ):
                    return True

                titulo = (
                    win32gui.GetWindowText(hwnd)
                    or ""
                ).strip()

                if titulo != "RM":
                    return True

                if pid_rm is not None:
                    _, pid = (
                        win32process.GetWindowThreadProcessId(
                            hwnd
                        )
                    )

                    if pid != pid_rm:
                        return True

                encontrados.append(
                    hwnd
                )

            except Exception:
                pass

            return True

        win32gui.EnumWindows(
            callback,
            None,
        )

        return (
            encontrados[0]
            if encontrados
            else None
        )

    hwnd_confirmacao = aguardar(
        localizar_confirmacao,
        10,
        0.2,
    )

    if hwnd_confirmacao is None:
        print(
            "Aviso: confirmacao de saida "
            "nao foi encontrada."
        )
        return False

    print(
        "Confirmacao de saida encontrada."
    )

    def localizar_ok():
        encontrados = []

        def callback(hwnd, _):
            try:
                titulo = (
                    win32gui.GetWindowText(hwnd)
                    or ""
                ).strip()

                if (
                    titulo == "OK"
                    and
                    win32gui.IsWindowVisible(hwnd)
                    and
                    win32gui.IsWindowEnabled(hwnd)
                ):
                    encontrados.append(
                        hwnd
                    )

            except Exception:
                pass

            return True

        win32gui.EnumChildWindows(
            hwnd_confirmacao,
            callback,
            None,
        )

        return (
            encontrados[0]
            if encontrados
            else None
        )

    hwnd_ok = aguardar(
        localizar_ok,
        5,
        0.2,
    )

    if hwnd_ok is None:
        print(
            "Aviso: controle OK da confirmacao "
            "nao foi encontrado."
        )
        return False

    print(
        "Confirmando encerramento..."
    )

    try:
        HwndWrapper(
            hwnd_ok
        ).click()

    except Exception as erro:
        print(
            "Aviso: falha ao acionar OK:"
        )
        print(
            repr(erro)
        )
        return False

    inicio_espera = time.time()

    while (
        time.time() - inicio_espera
        < 20
    ):
        if not win32gui.IsWindow(
            hwnd_main
        ):
            print(
                "TOTVS RM fechado."
            )
            return True

        time.sleep(
            0.5
        )

    print(
        "Aviso: o Bronze foi atualizado, "
        "mas o RM nao fechou automaticamente."
    )

    return False





# ============================================================================

# MAIN

# ============================================================================



def main():

    print()

    print("=" * 90)

    print(

        "EXTRACAO RM - CADASTRO COMPLETO DE COLABORADORES"

    )

    print("=" * 90)

    usuario = input(

        "Usuario/e-mail RM: "

    ).strip()

    if not usuario:

        raise RuntimeError(

            "Usuario nao informado."

        )

    senha = getpass.getpass(

        "Senha RM: "

    )

    if not senha:

        raise RuntimeError(

            "Senha nao informada."

        )

    garantir_rm_fechado()

    PASTA_BRONZE.mkdir(

        parents=True,

        exist_ok=True,

    )

    for temporario in [

        ARQUIVO_TEMPORARIO,

        ARQUIVO_TEMPORARIO_FOLHA,

    ]:

        if temporario.exists():

            print(

                "Removendo temporario de "

                "execucao anterior:"

            )

            print(

                temporario

            )

            temporario.unlink()

    main_rm = None

    try:

        abrir_rm()

        efetuar_login(

            usuario,

            senha,

        )

        senha = None

        main_rm = aguardar_mainform()

        print()

        print("=" * 90)

        print(

            "FONTE 1/2 - GESTAO DE PESSOAS"

        )

        print("=" * 90)

        abrir_gestao_pessoas(

            main_rm

        )

        abrir_pfunc(

            main_rm

        )

        sql_gestao = abrir_sql(

            main_rm

        )

        exportar_consulta_atual(

            sql_gestao,

            ARQUIVO_TEMPORARIO,

            "Gestao de Pessoas",

        )

        total_gestao = validar_exportacao(

            ARQUIVO_TEMPORARIO

        )

        fechar_janela_sql(

            sql_gestao

        )

        print()

        print("=" * 90)

        print(

            "FONTE 2/2 - FOLHA DE PAGAMENTO"

        )

        print("=" * 90)

        abrir_folha_pagamento(

            main_rm

        )

        abrir_visao_folha(

            main_rm

        )

        sql_folha = abrir_sql(

            main_rm

        )

        exportar_consulta_atual(

            sql_folha,

            ARQUIVO_TEMPORARIO_FOLHA,

            "Folha de Pagamento",

        )

        total_folha = validar_exportacao_folha(

            ARQUIVO_TEMPORARIO_FOLHA

        )

        fechar_janela_sql(

            sql_folha

        )

        promover_bronzes()

        print()

        print("=" * 90)

        print(

            "EXTRACAO COMPLETA CONCLUIDA COM SUCESSO"

        )

        print("=" * 90)

        print(

            "Gestao de Pessoas:",

            total_gestao,

            "registros",

        )

        print(

            "Folha de Pagamento:",

            total_folha,

            "registros",

        )

        print()

        print(

            "Bronze Gestao:"

        )

        print(

            ARQUIVO_BRONZE

        )

        print()

        print(

            "Bronze Folha:"

        )

        print(

            ARQUIVO_BRONZE_FOLHA

        )

        print()

        print(

            "Nenhuma regra de Prata foi aplicada."

        )

    finally:

        if main_rm is not None:

            print()

            print(

                "Fechando TOTVS RM apos "

                "as duas extracoes..."

            )

            rm_fechado = fechar_rm(

                main_rm

            )

            print(

                "RM encerrado automaticamente:",

                "SIM" if rm_fechado else "NAO",

            )



if __name__ == "__main__":

    try:

        main()



    except KeyboardInterrupt:

        print()

        print(

            "Execucao cancelada pelo usuario."

        )



        raise SystemExit(130)



    except Exception as erro:

        print()

        print("=" * 90)

        print("EXTRACAO INTERROMPIDA")

        print("=" * 90)



        print(

            type(erro).__name__ + ":",

            erro,

        )



        print()

        print(

            "Os Bronzes oficiais so sao promovidos "

            "depois que as duas extracoes passam "

            "pelos respectivos Quality Gates."

        )



        print("=" * 90)



        raise
