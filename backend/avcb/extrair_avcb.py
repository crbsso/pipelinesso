import base64
import os
import re
import zipfile

from pathlib import Path
from urllib.parse import (
    urlparse,
    parse_qs,
    unquote,
    quote,
)

from dotenv import load_dotenv
from openpyxl import load_workbook

from playwright.sync_api import (
    sync_playwright,
    TimeoutError as PlaywrightTimeoutError,
)


# =========================================================
# CONFIGURAÇÕES DO PROJETO
# =========================================================

PASTA_PROJETO = Path(__file__).resolve().parents[2]


# =========================================================
# CONFIGURAÇÕES DO AVCB
# =========================================================

NOME_ABA_AVCB = "Oficial"

load_dotenv(
    PASTA_PROJETO / ".env"
)

URL_AVCB = os.getenv(
    "AVCB_SHAREPOINT_URL"
)

if not URL_AVCB:
    raise RuntimeError(
        "AVCB_SHAREPOINT_URL não encontrada no arquivo .env."
    )


# =========================================================
# PASTAS
# =========================================================

PASTA_BRONZE = (
    PASTA_PROJETO
    / "data"
    / "bronze"
    / "avcb"
)

ARQUIVO_FINAL = (
    PASTA_BRONZE
    / "Controle de AVCB.xlsx"
)

ARQUIVO_TEMPORARIO = (
    PASTA_BRONZE
    / "Controle de AVCB_novo.xlsx"
)

# Perfil persistente do Edge usado pelo Playwright.
# Mantém a autenticação corporativa entre execuções.
PROFILE_DIR = (
    PASTA_PROJETO
    / ".edge_sso_playwright"
)


# =========================================================
# PREPARA URL DIRETA DE DOWNLOAD DO SHAREPOINT
# =========================================================

def preparar_download_sharepoint(
    url_original
):
    """
    Extrai da URL original:

    - tenant SharePoint
    - site
    - sourcedoc / UniqueId

    E monta a URL direta:

    /_layouts/15/download.aspx?UniqueId=...
    """

    parsed = urlparse(
        url_original
    )

    if not parsed.scheme or not parsed.netloc:
        raise RuntimeError(
            "A URL do AVCB não é válida."
        )

    if (
        "sharepoint.com"
        not in parsed.netloc.lower()
    ):
        raise RuntimeError(
            "A URL do AVCB não pertence ao SharePoint."
        )

    # -----------------------------------------------------
    # UNIQUE ID
    # -----------------------------------------------------

    parametros = parse_qs(
        parsed.query
    )

    sourcedoc = parametros.get(
        "sourcedoc"
    )

    if not sourcedoc:
        raise RuntimeError(
            "A URL do AVCB não possui "
            "o parâmetro sourcedoc."
        )

    unique_id = unquote(
        sourcedoc[0]
    ).strip()

    if not unique_id:
        raise RuntimeError(
            "Não foi possível identificar "
            "o UniqueId do arquivo AVCB."
        )

    # -----------------------------------------------------
    # IDENTIFICA O SITE SHAREPOINT
    #
    # Exemplo:
    #
    # /:x:/r/sites/BR_Compartilha/_layouts/...
    #
    # Resultado:
    #
    # /sites/BR_Compartilha
    # -----------------------------------------------------

    match = re.search(
        r"(/sites/[^/]+)",
        parsed.path,
        re.IGNORECASE,
    )

    if not match:
        raise RuntimeError(
            "Não foi possível identificar "
            "o site SharePoint na URL do AVCB."
        )

    site_path = match.group(1)

    origem = (
        f"{parsed.scheme}://"
        f"{parsed.netloc}"
    )

    site_url = (
        origem
        + site_path
    )

    unique_id_encoded = quote(
        unique_id,
        safe=""
    )

    download_url = (
        f"{site_url}"
        f"/_layouts/15/download.aspx"
        f"?UniqueId={unique_id_encoded}"
    )

    return (
        unique_id,
        site_url,
        download_url,
    )


# =========================================================
# VALIDA ESTRUTURA XLSX
# =========================================================

def validar_estrutura_xlsx(
    arquivo
):
    """
    Um arquivo .xlsx é internamente
    um arquivo ZIP.

    Esta validação impede que uma página
    HTML de login ou AccessDenied seja
    confundida com um arquivo Excel.
    """

    if not arquivo.exists():
        raise RuntimeError(
            "O arquivo baixado não foi encontrado."
        )

    tamanho = (
        arquivo.stat().st_size
    )

    if tamanho == 0:
        raise RuntimeError(
            "O arquivo baixado possui tamanho zero."
        )

    if not zipfile.is_zipfile(
        arquivo
    ):
        raise RuntimeError(
            "O arquivo recebido não é "
            "um XLSX válido."
        )

    with zipfile.ZipFile(
        arquivo,
        "r"
    ) as arquivo_zip:

        nomes = (
            arquivo_zip.namelist()
        )

        if (
            "[Content_Types].xml"
            not in nomes
        ):
            raise RuntimeError(
                "O arquivo recebido possui "
                "estrutura ZIP, mas não parece "
                "ser um XLSX válido."
            )

    return tamanho


# =========================================================
# VALIDA O CONTEÚDO DO EXCEL
# =========================================================

def validar_excel(
    arquivo
):

    print(
        f"Validando arquivo: {arquivo.name}"
    )

    # Primeiro valida a estrutura física
    tamanho = validar_estrutura_xlsx(
        arquivo
    )

    # Depois valida o conteúdo esperado
    workbook = load_workbook(
        arquivo,
        read_only=True,
        data_only=True
    )

    try:

        abas = (
            workbook.sheetnames
        )

    finally:

        workbook.close()

    if NOME_ABA_AVCB not in abas:

        raise ValueError(
            f"A aba '{NOME_ABA_AVCB}' "
            "nao foi encontrada. "
            f"Abas encontradas: {abas}"
        )

    print(
        "Arquivo Excel válido."
    )

    print(
        f"Aba '{NOME_ABA_AVCB}' encontrada."
    )

    print(
        f"Tamanho do arquivo: "
        f"{tamanho:,} bytes"
    )


# =========================================================
# DOWNLOAD DIRETO PELO EDGE
# =========================================================

def baixar_avcb():

    (
        unique_id,
        site_url,
        download_url,
    ) = preparar_download_sharepoint(
        URL_AVCB
    )

    print(
        "Arquivo AVCB identificado no SharePoint."
    )

    print(
        f"UniqueId: {unique_id}"
    )

    print(
        f"Site: {site_url}"
    )

    PASTA_BRONZE.mkdir(
        parents=True,
        exist_ok=True
    )

    # Remove eventual arquivo temporário
    # deixado por execução anterior.
    if ARQUIVO_TEMPORARIO.exists():

        ARQUIVO_TEMPORARIO.unlink()

    with sync_playwright() as p:

        print(
            "Abrindo sessão autenticada "
            "do Microsoft Edge..."
        )

        context = (
            p.chromium.launch_persistent_context(
                user_data_dir=str(
                    PROFILE_DIR
                ),

                channel="msedge",

                headless=False,

                accept_downloads=True,

                viewport=None,

                chromium_sandbox=True,

                args=[
                    "--start-maximized",
                ],
            )
        )

        try:

            page = (
                context.pages[0]
                if context.pages
                else context.new_page()
            )

            # Diagnóstico: ajuda a identificar se alguma
            # política corporativa fechou o Edge inesperadamente.
            context.on(
                "close",
                lambda:
                    print(
                        "[AVISO] Contexto do Edge foi fechado."
                    )
            )

            page.on(
                "close",
                lambda:
                    print(
                        "[AVISO] Página do Edge foi fechada."
                    )
            )

            # -------------------------------------------------
            # ABRE SOMENTE O SITE SHAREPOINT
            #
            # Não abre o Excel Online.
            #
            # O objetivo é executar o fetch dentro de uma página
            # do próprio domínio celeogroup.sharepoint.com.
            # Assim usamos:
            #
            # - DNS do Edge corporativo
            # - proxy corporativo
            # - cookies da sessão autenticada
            #
            # e não usamos APIRequestContext nem Download.save_as.
            # -------------------------------------------------

            print(
                "Preparando sessão no site SharePoint..."
            )

            try:

                page.goto(
                    site_url,
                    wait_until="domcontentloaded",
                    timeout=60000,
                )

            except PlaywrightTimeoutError:

                # A home do SharePoint pode continuar carregando
                # conteúdo dinâmico. Para o fetch, basta termos
                # uma página viva no mesmo domínio.
                print(
                    "[AVISO] O SharePoint demorou para concluir "
                    "o carregamento, mas a sessão continuará."
                )

            if page.is_closed():

                raise RuntimeError(
                    "O Edge fechou a página antes da extração."
                )

            # -------------------------------------------------
            # DOWNLOAD VIA FETCH DENTRO DO PRÓPRIO EDGE
            #
            # Diferentemente de context.request, esse fetch usa
            # a pilha de rede do navegador corporativo.
            #
            # Diferentemente de expect_download/save_as, o arquivo
            # não fica preso ao artefato temporário do Playwright.
            # Os bytes são entregues ao Python antes do navegador
            # ser fechado.
            # -------------------------------------------------

            print(
                "Baixando arquivo pelo navegador "
                "sem abrir o Excel Online..."
            )

            try:

                resultado = page.evaluate(
                    """
                    async (url) => {

                        const response = await fetch(
                            url,
                            {
                                method: "GET",
                                credentials: "include",
                                cache: "no-store",
                                redirect: "follow"
                            }
                        );

                        const contentType =
                            response.headers.get("content-type") || "";

                        const contentDisposition =
                            response.headers.get("content-disposition") || "";

                        if (!response.ok) {

                            throw new Error(
                                "SharePoint retornou HTTP "
                                + response.status
                                + " "
                                + response.statusText
                            );
                        }

                        const buffer =
                            await response.arrayBuffer();

                        const bytes =
                            new Uint8Array(buffer);

                        let binary = "";

                        const chunkSize = 0x8000;

                        for (
                            let i = 0;
                            i < bytes.length;
                            i += chunkSize
                        ) {

                            const chunk =
                                bytes.subarray(
                                    i,
                                    Math.min(
                                        i + chunkSize,
                                        bytes.length
                                    )
                                );

                            binary +=
                                String.fromCharCode(
                                    ...chunk
                                );
                        }

                        return {
                            status: response.status,
                            finalUrl: response.url,
                            contentType: contentType,
                            contentDisposition: contentDisposition,
                            tamanho: bytes.length,
                            base64: btoa(binary)
                        };
                    }
                    """,
                    download_url,
                )

            except Exception as erro:

                raise RuntimeError(
                    "O Edge não conseguiu obter o arquivo "
                    "diretamente do SharePoint. "
                    f"Detalhe: {erro}"
                ) from erro

            # -------------------------------------------------
            # VALIDA RESPOSTA HTTP ANTES DE GRAVAR
            # -------------------------------------------------

            status = resultado.get(
                "status"
            )

            content_type = (
                resultado.get(
                    "contentType",
                    ""
                )
                or ""
            )

            tamanho_recebido = (
                resultado.get(
                    "tamanho",
                    0
                )
                or 0
            )

            print(
                f"HTTP recebido: {status}"
            )

            print(
                f"Content-Type: {content_type or 'não informado'}"
            )

            print(
                f"Tamanho recebido: "
                f"{tamanho_recebido:,} bytes"
            )

            if status != 200:

                raise RuntimeError(
                    f"O SharePoint retornou HTTP {status}."
                )

            if tamanho_recebido <= 0:

                raise RuntimeError(
                    "O SharePoint retornou um arquivo vazio."
                )

            dados_base64 = (
                resultado.get(
                    "base64"
                )
            )

            if not dados_base64:

                raise RuntimeError(
                    "O navegador não retornou "
                    "o conteúdo do arquivo."
                )

            # -------------------------------------------------
            # GRAVA OS BYTES DIRETAMENTE NA BRONZE TEMPORÁRIA
            # -------------------------------------------------

            try:

                dados = base64.b64decode(
                    dados_base64,
                    validate=True
                )

            except Exception as erro:

                raise RuntimeError(
                    "Não foi possível decodificar "
                    "o arquivo recebido do SharePoint."
                ) from erro

            if len(dados) != tamanho_recebido:

                raise RuntimeError(
                    "O tamanho recebido do navegador "
                    "não corresponde ao tamanho decodificado."
                )

            ARQUIVO_TEMPORARIO.write_bytes(
                dados
            )

            if not ARQUIVO_TEMPORARIO.exists():

                raise RuntimeError(
                    "O arquivo temporário não foi criado "
                    "na camada Bronze."
                )

            tamanho_salvo = (
                ARQUIVO_TEMPORARIO.stat().st_size
            )

            if tamanho_salvo != tamanho_recebido:

                raise RuntimeError(
                    "O arquivo salvo ficou com tamanho "
                    "diferente do conteúdo recebido."
                )

            print(
                "Nova versão salva temporariamente "
                "na camada Bronze."
            )

            print(
                f"Arquivo temporário: "
                f"{ARQUIVO_TEMPORARIO}"
            )

        finally:

            # Só fecha o navegador DEPOIS de os bytes
            # já estarem persistidos no disco pelo Python.
            try:

                context.close()

            except Exception:

                pass

    return ARQUIVO_TEMPORARIO


# =========================================================
# ATUALIZA A CAMADA BRONZE
# =========================================================

def atualizar_bronze(
    arquivo_baixado
):
    """
    Só substitui o arquivo oficial depois
    que o novo arquivo passou por todas
    as validações.
    """

    if not arquivo_baixado.exists():

        raise RuntimeError(
            "Arquivo temporário não encontrado "
            "para atualização da Bronze."
        )

    print(
        "Substituindo versão anterior "
        "da camada Bronze..."
    )

    os.replace(
        arquivo_baixado,
        ARQUIVO_FINAL
    )

    print(
        "Camada Bronze atualizada "
        "com sucesso:"
    )

    print(
        ARQUIVO_FINAL
    )


# =========================================================
# LIMPEZA EM CASO DE ERRO
# =========================================================

def limpar_temporario():

    try:

        if ARQUIVO_TEMPORARIO.exists():

            ARQUIVO_TEMPORARIO.unlink()

    except Exception:

        pass


# =========================================================
# EXTRAÇÃO DO AVCB
# =========================================================

def extrair_avcb():

    print(
        "=" * 70
    )

    print(
        "EXTRAÇÃO AVCB - SHAREPOINT"
    )

    print(
        "=" * 70
    )

    print()

    print(
        "Método: Playwright + "
        "download direto SharePoint"
    )

    print(
        "Excel Online: não utilizado"
    )

    print(
        "PyAutoGUI: não utilizado"
    )

    print()

    try:

        # -----------------------------------------------
        # 1. DOWNLOAD
        # -----------------------------------------------

        arquivo = baixar_avcb()

        # -----------------------------------------------
        # 2. VALIDAÇÃO
        # -----------------------------------------------

        validar_excel(
            arquivo
        )

        # -----------------------------------------------
        # 3. PROMOÇÃO PARA BRONZE OFICIAL
        # -----------------------------------------------

        atualizar_bronze(
            arquivo
        )

        print()

        print(
            "=" * 70
        )

        print(
            "EXTRAÇÃO AVCB CONCLUÍDA COM SUCESSO"
        )

        print(
            "=" * 70
        )

    except Exception:

        # Nunca deixa um arquivo incompleto
        # virar a Bronze oficial.

        limpar_temporario()

        raise


# =========================================================
# EXECUÇÃO
# =========================================================

if __name__ == "__main__":
    extrair_avcb()
