from pathlib import Path
from datetime import datetime, date
import numbers
import os

import pandas as pd
from docx import Document


# =========================================================
# CONFIGURAÇÕES
# =========================================================

PASTA_PROJETO = Path(__file__).resolve().parents[2]

ARQUIVO_AVCB = (
    PASTA_PROJETO
    / "data"
    / "bronze"
    / "avcb"
    / "Controle de AVCB.xlsx"
)

PASTA_RELATORIOS = (
    PASTA_PROJETO
    / "relatorios"
)

NOME_ABA_AVCB = "Oficial"

NUMERO_COLUNAS_ESPERADO = 16


# Colunas que obrigatoriamente devem ser texto
# A, B, C, D, E, G, H, I, J, K, L, N, O, P
COLUNAS_STRING = {
    0, 1, 2, 3, 4,
    6, 7, 8, 9, 10, 11,
    13, 14, 15
}

# F
COLUNA_INTEIRO = 5

# M
COLUNA_DATA = 12


# =========================================================
# FUNÇÕES AUXILIARES
# =========================================================

def letra_coluna(indice):
    return chr(65 + indice)


def esta_vazio(valor):

    if pd.isna(valor):
        return True

    if isinstance(valor, str):
        return valor.strip() == ""

    return False


def eh_inteiro_sem_decimal(valor):

    if isinstance(valor, bool):
        return False

    if not isinstance(valor, numbers.Number):
        return False

    try:
        return float(valor).is_integer()
    except (TypeError, ValueError):
        return False


def eh_data(valor):

    return isinstance(
        valor,
        (
            pd.Timestamp,
            datetime,
            date
        )
    )


# =========================================================
# VALIDAÇÃO DO AVCB
# =========================================================

def validar_avcb():

    print("Validando qualidade do AVCB...")

    df = pd.read_excel(
        ARQUIVO_AVCB,
        sheet_name=NOME_ABA_AVCB
    )

    quantidade_colunas = len(df.columns)

    print(f"Registros analisados: {len(df)}")
    print(f"Colunas encontradas: {quantidade_colunas}")

    problemas_vazios = []
    problemas_tipo = []


    # =====================================================
    # AVCB-01 — ESTRUTURA
    # =====================================================

    estrutura_ok = (
        quantidade_colunas
        == NUMERO_COLUNAS_ESPERADO
    )


    # =====================================================
    # AVCB-02 — CAMPOS VAZIOS
    # AVCB-03/04/05 — TIPOS
    # =====================================================

    for indice_coluna, coluna in enumerate(df.columns):

        for indice_linha, valor in df[coluna].items():

            linha_excel = indice_linha + 2
            coluna_excel = letra_coluna(indice_coluna)

            # ---------------------------------------------
            # CAMPO VAZIO
            # ---------------------------------------------

            if esta_vazio(valor):

                problemas_vazios.append({
                    "linha": linha_excel,
                    "coluna": coluna_excel,
                    "campo": coluna
                })

                # Se está vazio, não faz sentido
                # também apontar erro de tipo.
                continue


            # ---------------------------------------------
            # STRING
            # ---------------------------------------------

            if indice_coluna in COLUNAS_STRING:

                if not isinstance(valor, str):

                    problemas_tipo.append({
                        "linha": linha_excel,
                        "coluna": coluna_excel,
                        "campo": coluna,
                        "valor": valor,
                        "esperado": "TEXTO"
                    })


            # ---------------------------------------------
            # INTEIRO SEM DECIMAL
            # ---------------------------------------------

            elif indice_coluna == COLUNA_INTEIRO:

                if not eh_inteiro_sem_decimal(valor):

                    problemas_tipo.append({
                        "linha": linha_excel,
                        "coluna": coluna_excel,
                        "campo": coluna,
                        "valor": valor,
                        "esperado": "NÚMERO INTEIRO SEM DECIMAL"
                    })


            # ---------------------------------------------
            # DATA
            # ---------------------------------------------

            elif indice_coluna == COLUNA_DATA:

                if not eh_data(valor):

                    problemas_tipo.append({
                        "linha": linha_excel,
                        "coluna": coluna_excel,
                        "campo": coluna,
                        "valor": valor,
                        "esperado": "DATA"
                    })


    return {
        "arquivo": ARQUIVO_AVCB.name,
        "aba": NOME_ABA_AVCB,
        "registros": len(df),
        "colunas_encontradas": quantidade_colunas,
        "colunas_esperadas": NUMERO_COLUNAS_ESPERADO,
        "estrutura_ok": estrutura_ok,
        "problemas_vazios": problemas_vazios,
        "problemas_tipo": problemas_tipo
    }


# =========================================================
# RELATÓRIO WORD
# =========================================================

def gerar_relatorio(resultado):

    PASTA_RELATORIOS.mkdir(
        parents=True,
        exist_ok=True
    )

    agora = datetime.now()

    caminho_relatorio = (
        PASTA_RELATORIOS
        / (
            "relatorio_qualidade_"
            f"{agora.strftime('%Y-%m-%d')}.docx"
        )
    )

    documento = Document()

    documento.add_heading(
        "Relatório de Qualidade de Dados",
        level=1
    )

    documento.add_paragraph(
        f"Data da validação: "
        f"{agora.strftime('%d/%m/%Y %H:%M')}"
    )

    documento.add_paragraph(
        f"Arquivo: {resultado['arquivo']}"
    )

    documento.add_paragraph(
        f"Aba analisada: {resultado['aba']}"
    )

    documento.add_paragraph(
        f"Registros analisados: "
        f"{resultado['registros']}"
    )


    # =====================================================
    # 1. ESTRUTURA
    # =====================================================

    documento.add_heading(
        "1. Estrutura da base",
        level=2
    )

    documento.add_paragraph(
        f"Colunas esperadas: "
        f"{resultado['colunas_esperadas']}"
    )

    documento.add_paragraph(
        f"Colunas encontradas: "
        f"{resultado['colunas_encontradas']}"
    )

    if resultado["estrutura_ok"]:
        documento.add_paragraph(
            "Status da estrutura: OK"
        )
    else:
        documento.add_paragraph(
            "Status da estrutura: ERRO"
        )


    # =====================================================
    # 2. CAMPOS VAZIOS
    # =====================================================

    documento.add_heading(
        "2. Campos em branco",
        level=2
    )

    vazios = resultado[
        "problemas_vazios"
    ]

    documento.add_paragraph(
        f"Total de campos em branco: "
        f"{len(vazios)}"
    )

    if vazios:

        tabela = documento.add_table(
            rows=1,
            cols=3
        )

        tabela.style = "Table Grid"

        cabecalho = tabela.rows[0].cells

        cabecalho[0].text = "Linha"
        cabecalho[1].text = "Coluna"
        cabecalho[2].text = "Campo"

        for problema in vazios:

            linha = tabela.add_row().cells

            linha[0].text = str(
                problema["linha"]
            )

            linha[1].text = str(
                problema["coluna"]
            )

            linha[2].text = str(
                problema["campo"]
            )


    # =====================================================
    # 3. TIPO DE DADO
    # =====================================================

    documento.add_heading(
        "3. Tipos de dados",
        level=2
    )

    tipos = resultado[
        "problemas_tipo"
    ]

    documento.add_paragraph(
        f"Total de incompatibilidades: "
        f"{len(tipos)}"
    )

    if tipos:

        tabela = documento.add_table(
            rows=1,
            cols=5
        )

        tabela.style = "Table Grid"

        cabecalho = tabela.rows[0].cells

        cabecalho[0].text = "Linha"
        cabecalho[1].text = "Coluna"
        cabecalho[2].text = "Campo"
        cabecalho[3].text = "Valor encontrado"
        cabecalho[4].text = "Esperado"

        for problema in tipos:

            linha = tabela.add_row().cells

            linha[0].text = str(
                problema["linha"]
            )

            linha[1].text = str(
                problema["coluna"]
            )

            linha[2].text = str(
                problema["campo"]
            )

            linha[3].text = str(
                problema["valor"]
            )

            linha[4].text = str(
                problema["esperado"]
            )


    # =====================================================
    # STATUS FINAL
    # =====================================================

    documento.add_heading(
        "Status final",
        level=2
    )

    possui_problema = (
        not resultado["estrutura_ok"]
        or len(vazios) > 0
        or len(tipos) > 0
    )

    if possui_problema:

        documento.add_paragraph(
            "REPROVADO — Existem pendências "
            "que devem ser verificadas na fonte."
        )

    else:

        documento.add_paragraph(
            "APROVADO — Nenhuma pendência "
            "foi encontrada."
        )

    documento.save(
        caminho_relatorio
    )

    return caminho_relatorio


# =========================================================
# EXECUÇÃO
# =========================================================

def main():

    resultado = validar_avcb()

    total_vazios = len(
        resultado["problemas_vazios"]
    )

    total_tipos = len(
        resultado["problemas_tipo"]
    )

    possui_problema = (
        not resultado["estrutura_ok"]
        or total_vazios > 0
        or total_tipos > 0
    )

    print()
    print("=" * 50)

    if possui_problema:

        print("VALIDAÇÃO REPROVADA")

    else:

        print("VALIDAÇÃO APROVADA")

    print("=" * 50)

    print(
        f"Campos vazios: {total_vazios}"
    )

    print(
        f"Erros de tipo: {total_tipos}"
    )

    if not resultado["estrutura_ok"]:

        print(
            "ERRO: quantidade de colunas "
            "diferente de 16."
        )


    # Só gera Word quando há problema.
    if possui_problema:

        relatorio = gerar_relatorio(
            resultado
        )

        print()
        print("Relatório gerado em:")
        print(relatorio)

        print()
        print("Abrindo relatório no Word...")

        os.startfile(
            str(relatorio)
        )

    else:

        print()
        print(
            "A base está liberada "
            "para a camada Prata."
        )


if __name__ == "__main__":
    main()
