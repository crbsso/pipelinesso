from pathlib import Path
import os
import re
import time
import zipfile

from openpyxl import load_workbook
from playwright.sync_api import (
    sync_playwright,
    TimeoutError as PlaywrightTimeoutError,
)


# =========================================================
# CONFIGURACOES DO PROJETO
# =========================================================

PASTA_PROJETO = Path(__file__).resolve().parents[2]

URL_PREVSIS = os.getenv(
    "PREVSIS_URL",
    "https://celeoqa.prevsis.com/",
)

PROFILE_DIR = (
    PASTA_PROJETO
    / ".edge_sso_playwright"
)


# =========================================================
# BRONZE - INSPECOES / QUANTIDADE
# =========================================================

PASTA_BRONZE = (
    PASTA_PROJETO
    / "data"
    / "bronze"
    / "inspecoes"
    / "quantidade"
)

ARQUIVO_FINAL = (
    PASTA_BRONZE
    / "Detalle_Checklist.xlsx"
)

ARQUIVO_TEMPORARIO = (
    PASTA_BRONZE
    / "Detalle_Checklist_novo.xlsx"
)


# =========================================================
# CAMPOS MINIMOS ESPERADOS
# =========================================================

COLUNAS_OBRIGATORIAS = {
    "ID resposta",
    "Nome formulário",
    "Data da inspeção",
    "Área organizacional",
    "Local",
    "Empresa",
    "Criador de checklist",
}


# =========================================================
# UTILITARIOS PLAYWRIGHT
# =========================================================

def frames_da_pagina(page):
    """
    Retorna o frame principal e todos os iframes carregados.

    O dashboard do Prevsys utiliza Amazon QuickSight,
    portanto parte dos elementos pode estar dentro de iframe.
    """
    return page.frames


def localizar_texto(
    page,
    texto,
    exact=False,
):
    """
    Procura o texto no frame principal e nos iframes.

    Quando existem elementos duplicados, retorna
    somente um elemento realmente visivel.
    """

    for frame in frames_da_pagina(page):

        try:
            locators = frame.get_by_text(
                texto,
                exact=exact,
            )

            quantidade = locators.count()

            for indice in range(quantidade):

                locator = locators.nth(indice)

                try:
                    if locator.is_visible():
                        return frame, locator
                except Exception:
                    continue

        except Exception:
            continue

    return None, None



def aguardar_texto(
    page,
    texto,
    timeout=60000,
    exact=False,
):
    """
    Aguarda um texto aparecer em qualquer frame.
    """

    limite = time.time() + timeout / 1000

    while time.time() < limite:

        frame, locator = localizar_texto(
            page,
            texto,
            exact=exact,
        )

        if locator is not None:

            try:
                if locator.is_visible():
                    return frame, locator
            except Exception:
                pass

        time.sleep(0.5)

    raise RuntimeError(
        f"O texto '{texto}' não apareceu "
        f"dentro de {timeout / 1000:.0f}s."
    )


# =========================================================
# LOGIN
# =========================================================

def garantir_login(page):
    """
    Abre o Prevsys.

    Se aparecer Azure AD Login, clica nele.

    Na primeira execucao o usuario pode precisar
    concluir manualmente o login corporativo/MFA.
    O perfil persistente conservara a sessao.
    """

    print("Abrindo Prevsys...")

    page.goto(
        URL_PREVSIS,
        wait_until="domcontentloaded",
        timeout=60000,
    )

    time.sleep(2)

    _, botao_login = localizar_texto(
        page,
        "Azure AD Login",
        exact=True,
    )

    if botao_login is not None:

        try:
            if botao_login.is_visible():

                print(
                    "Tela de login encontrada. "
                    "Clicando em Azure AD Login..."
                )

                botao_login.click()

        except Exception:
            pass

    print(
        "Aguardando a sessão autenticada do Prevsys..."
    )

    try:
        page.wait_for_url(
            re.compile(
                r"https://celeo\.prevsis\.com/.*"
            ),
            timeout=120000,
        )

    except PlaywrightTimeoutError:

        raise RuntimeError(
            "Não foi possível confirmar a sessão "
            "autenticada no Prevsys. "
            "Se aparecer Microsoft/Azure AD no Edge, "
            "conclua o login e execute novamente."
        )

    print(
        f"Sessão autenticada: {page.url}"
    )


# =========================================================
# ABRE REPORTE 360
# =========================================================

def abrir_reporte_360(page):

    print("Procurando Reporte 360...")

    _, locator = aguardar_texto(
        page,
        "Reporte 360",
        timeout=60000,
        exact=False,
    )

    print("Reporte 360 visivel encontrado.")

    locator.scroll_into_view_if_needed()

    locator.click(
        timeout=15000,
    )

    print(
        "Reporte 360 acionado. "
        "Aguardando pagina do relatorio..."
    )

    try:
        page.wait_for_url(
            re.compile(
                r".*/KPI/VistaKPI.*",
                re.IGNORECASE,
            ),
            timeout=120000,
        )
    except PlaywrightTimeoutError:
        raise RuntimeError(
            "O clique em Reporte 360 ocorreu, "
            "mas a pagina KPI/VistaKPI nao foi carregada."
        )

    print(
        f"Reporte 360 carregado: {page.url}"
    )

    # Aguarda a interface do QuickSight iniciar.
    time.sleep(5)


# =========================================================
# ABRE ANALISE DA LISTA DE VERIFICACAO
# =========================================================

def abrir_lista_verificacao(page):

    print(
        "Procurando Analise da lista de verificacao..."
    )

    padrao = re.compile(
        r"lista de verifica",
        re.IGNORECASE,
    )

    _, locator = aguardar_texto(
        page,
        padrao,
        timeout=120000,
        exact=False,
    )

    print(
        "Opcao de lista de verificacao encontrada."
    )

    locator.scroll_into_view_if_needed()

    locator.click(
        timeout=15000,
    )

    print(
        "Analise da lista de verificacao acionada."
    )

    # O filtro Divisao fica no topo da analise e
    # funciona como marcador de que o dashboard carregou.
    aguardar_texto(
        page,
        re.compile(
            r"Divis",
            re.IGNORECASE,
        ),
        timeout=120000,
        exact=False,
    )

    print(
        "Dashboard da lista de verificacao carregado."
    )


# =========================================================
# FILTRO DIVISAO = CELEO BRASIL
# =========================================================


# =========================================================
# LOCALIZA MENU DO DETALLE CHECKLIST
# =========================================================

def abrir_menu_detalle(page):

    print(
        "Procurando visual Detalle Checklist..."
    )

    # =====================================================
    # 1. LOCALIZA O FRAME DO QUICKSIGHT
    # =====================================================

    frame_quicksight = None

    for frame in frames_da_pagina(page):
        try:
            if "quicksight.aws.amazon.com" in frame.url:
                frame_quicksight = frame
                break
        except Exception:
            continue

    if frame_quicksight is None:
        raise RuntimeError(
            "Frame do Amazon QuickSight nao encontrado."
        )

    print(
        "Frame do QuickSight encontrado."
    )

    # =====================================================
    # 2. PROCURA DETALLE CHECKLIST E ROLA O DASHBOARD
    # =====================================================

    titulo = None
    limite = time.time() + 90
    tentativa = 0

    while time.time() < limite:

        tentativa += 1

        # -------------------------------------------------
        # Primeiro procura o texto no DOM, mesmo que ainda
        # esteja fora da area visivel.
        # -------------------------------------------------

        for exact in [True, False]:

            try:

                candidatos = (
                    frame_quicksight.get_by_text(
                        "Detalle Checklist",
                        exact=exact,
                    )
                )

                for indice in range(
                    candidatos.count()
                ):

                    candidato = candidatos.nth(
                        indice
                    )

                    try:

                        candidato.evaluate(
                            """
                            (el) => {
                                el.scrollIntoView({
                                    block: "center",
                                    inline: "nearest"
                                });
                            }
                            """
                        )

                        time.sleep(0.5)

                        if candidato.is_visible():

                            titulo = candidato
                            break

                    except Exception:
                        continue

            except Exception:
                pass

            if titulo is not None:
                break

        if titulo is not None:
            break

        # -------------------------------------------------
        # Se o titulo ainda nao existe/nao esta visivel,
        # avanca todos os containers rolaveis do QuickSight.
        # -------------------------------------------------

        try:

            quantidade_scroll = (
                frame_quicksight.evaluate(
                    """
                    () => {

                        const elementos =
                            Array.from(
                                document.querySelectorAll("*")
                            )
                            .filter((el) => {

                                if (
                                    !el ||
                                    el.clientHeight < 150
                                ) {
                                    return false;
                                }

                                return (
                                    el.scrollHeight >
                                    el.clientHeight + 100
                                );
                            });

                        let alterados = 0;

                        for (const el of elementos) {

                            const antes =
                                el.scrollTop;

                            const passo =
                                Math.max(
                                    700,
                                    el.clientHeight * 0.8
                                );

                            el.scrollTop =
                                Math.min(
                                    el.scrollTop + passo,
                                    el.scrollHeight
                                );

                            if (
                                el.scrollTop !== antes
                            ) {
                                alterados++;
                            }
                        }

                        window.scrollBy(
                            0,
                            Math.max(
                                700,
                                window.innerHeight * 0.8
                            )
                        );

                        return alterados;
                    }
                    """
                )
            )

        except Exception:
            quantidade_scroll = 0

        if tentativa == 1:

            print(
                "Detalle Checklist ainda nao esta visivel. "
                "Rolando o dashboard do QuickSight..."
            )

        if tentativa % 5 == 0:

            print(
                f"  Tentativa {tentativa}: "
                f"{quantidade_scroll} containers rolados."
            )

        time.sleep(1)

    if titulo is None:

        # Salva evidencia caso ainda falhe.
        try:

            pasta_diag = (
                PASTA_PROJETO
                / "data"
                / "diagnosticos"
            )

            pasta_diag.mkdir(
                parents=True,
                exist_ok=True,
            )

            screenshot = (
                pasta_diag
                / "prevsis_detalle_nao_encontrado.png"
            )

            page.screenshot(
                path=str(screenshot),
                full_page=False,
            )

            print(
                f"Screenshot salvo em: {screenshot}"
            )

        except Exception:
            pass

        raise RuntimeError(
            "Nao foi possivel localizar o visual "
            "'Detalle Checklist' no QuickSight."
        )

    print(
        "Visual Detalle Checklist encontrado."
    )

    try:
        titulo.scroll_into_view_if_needed()
    except Exception:
        pass

    titulo.hover()

    time.sleep(1)

    # =====================================================
    # 3. LOCALIZA O MENU DO VISUAL
    # =====================================================

    print(
        "Localizando menu do visual Detalle Checklist..."
    )

    nomes_menu = [
        re.compile(
            r"more options",
            re.IGNORECASE,
        ),
        re.compile(
            r"visual.*menu",
            re.IGNORECASE,
        ),
        re.compile(
            r"options",
            re.IGNORECASE,
        ),
        re.compile(
            r"menu",
            re.IGNORECASE,
        ),
    ]

    for nome in nomes_menu:

        try:

            botoes = (
                frame_quicksight.get_by_role(
                    "button",
                    name=nome,
                )
            )

            for indice in reversed(
                range(botoes.count())
            ):

                botao = botoes.nth(
                    indice
                )

                try:

                    if not botao.is_visible():
                        continue

                    botao.click(
                        timeout=10000
                    )

                    time.sleep(1)

                    export_excel = (
                        frame_quicksight.get_by_text(
                            "Export to Excel",
                            exact=True,
                        )
                    )

                    for j in range(
                        export_excel.count()
                    ):

                        if (
                            export_excel
                            .nth(j)
                            .is_visible()
                        ):

                            print(
                                "Menu de exportacao aberto."
                            )

                            return frame_quicksight

                except Exception:
                    continue

        except Exception:
            pass

    # =====================================================
    # 4. FALLBACK PELOS ATRIBUTOS ARIA/TITLE
    # =====================================================

    candidatos = (
        frame_quicksight.locator(
            "[aria-label], [title]"
        )
    )

    for indice in range(
        candidatos.count()
    ):

        elemento = candidatos.nth(
            indice
        )

        try:

            if not elemento.is_visible():
                continue

            aria = (
                elemento.get_attribute(
                    "aria-label"
                )
                or ""
            )

            title = (
                elemento.get_attribute(
                    "title"
                )
                or ""
            )

            descricao = (
                aria + " " + title
            ).lower()

            if not any(
                termo in descricao
                for termo in [
                    "more",
                    "option",
                    "menu",
                ]
            ):
                continue

            elemento.click(
                timeout=10000
            )

            time.sleep(1)

            export_excel = (
                frame_quicksight.get_by_text(
                    "Export to Excel",
                    exact=True,
                )
            )

            for j in range(
                export_excel.count()
            ):

                if (
                    export_excel
                    .nth(j)
                    .is_visible()
                ):

                    print(
                        "Menu de exportacao aberto."
                    )

                    return frame_quicksight

        except Exception:
            continue

    raise RuntimeError(
        "O visual Detalle Checklist foi encontrado, "
        "mas o menu de exportacao nao foi localizado."
    )


# =========================================================
# EXPORTA EXCEL
# =========================================================

def exportar_excel(page):

    import base64

    frame = abrir_menu_detalle(
        page
    )

    print(
        "Selecionando Export to Excel..."
    )

    exportar = frame.get_by_text(
        "Export to Excel",
        exact=True,
    )

    if exportar.count() == 0:

        _, exportar = aguardar_texto(
            page,
            "Export to Excel",
            timeout=10000,
            exact=True,
        )

    PASTA_BRONZE.mkdir(
        parents=True,
        exist_ok=True,
    )

    if ARQUIVO_TEMPORARIO.exists():
        ARQUIVO_TEMPORARIO.unlink()

    estado_blob = {
        "capturado": False,
        "nome": "",
        "tamanho": 0,
        "erro": "",
    }

    def salvar_blob_quicksight(payload):

        try:

            conteudo_base64 = payload.get(
                "base64",
                ""
            )

            nome = payload.get(
                "nome",
                ""
            )

            dados = base64.b64decode(
                conteudo_base64,
                validate=True,
            )

            if not dados:

                raise RuntimeError(
                    "Blob recebido sem dados."
                )

            ARQUIVO_TEMPORARIO.write_bytes(
                dados
            )

            estado_blob[
                "capturado"
            ] = True

            estado_blob[
                "nome"
            ] = nome

            estado_blob[
                "tamanho"
            ] = len(dados)

            print(
                "Blob do Excel capturado "
                "diretamente do QuickSight."
            )

            print(
                f"Tamanho capturado: "
                f"{len(dados):,} bytes"
            )

            return True

        except Exception as erro:

            estado_blob[
                "erro"
            ] = str(erro)

            return False

    page.expose_function(
        "__salvar_excel_quicksight",
        salvar_blob_quicksight,
    )

    script_interceptador = r'''
() => {

    if (window.__qs_blob_interceptor_installed) {
        return true;
    }

    window.__qs_blob_interceptor_installed = true;

    const blobs = new Map();

    window.__qs_blob_map = blobs;

    const originalCreateObjectURL =
        URL.createObjectURL.bind(URL);

    URL.createObjectURL = function(objeto) {

        const url =
            originalCreateObjectURL(
                objeto
            );

        try {

            if (objeto instanceof Blob) {

                blobs.set(
                    url,
                    objeto
                );

            }

        } catch (erro) {
        }

        return url;
    };


    async function capturarAnchor(anchor) {

        try {

            const href =
                anchor.href
                ||
                anchor.getAttribute(
                    "href"
                )
                ||
                "";

            if (!href.startsWith("blob:")) {
                return false;
            }

            const blob =
                blobs.get(
                    href
                );

            if (!blob) {
                return false;
            }

            if (window.__qs_blob_capture_running) {
                return true;
            }

            window.__qs_blob_capture_running =
                true;

            const buffer =
                await blob.arrayBuffer();

            const bytes =
                new Uint8Array(
                    buffer
                );

            let binario = "";

            const tamanhoBloco =
                32768;

            for (
                let inicio = 0;
                inicio < bytes.length;
                inicio += tamanhoBloco
            ) {

                const parte =
                    bytes.subarray(
                        inicio,
                        Math.min(
                            inicio + tamanhoBloco,
                            bytes.length
                        )
                    );

                binario +=
                    String.fromCharCode(
                        ...parte
                    );

            }

            const base64 =
                btoa(
                    binario
                );

            const nome =
                anchor.download
                ||
                "Detalle_Checklist.xlsx";

            const resultado =
                await window
                    .__salvar_excel_quicksight(
                        {
                            base64: base64,
                            nome: nome,
                            tipo: blob.type || "",
                            tamanho: blob.size || 0,
                        }
                    );

            window.__qs_blob_capture_ok =
                Boolean(
                    resultado
                );

            return Boolean(
                resultado
            );

        } catch (erro) {

            window.__qs_blob_capture_error =
                String(
                    erro
                );

            return false;

        } finally {

            window.__qs_blob_capture_running =
                false;

        }
    }


    const originalAnchorClick =
        HTMLAnchorElement
            .prototype
            .click;

    HTMLAnchorElement
        .prototype
        .click = function() {

            const href =
                this.href
                ||
                this.getAttribute(
                    "href"
                )
                ||
                "";

            if (
                href.startsWith("blob:")
                &&
                blobs.has(href)
            ) {

                capturarAnchor(
                    this
                );

                return;
            }

            return originalAnchorClick
                .apply(
                    this,
                    arguments
                );
        };


    document.addEventListener(
        "click",
        function(evento) {

            let anchor =
                evento.target;

            try {

                if (
                    anchor
                    &&
                    anchor.closest
                ) {

                    anchor =
                        anchor.closest(
                            "a"
                        );

                }

            } catch (erro) {

                return;

            }

            if (!anchor) {
                return;
            }

            const href =
                anchor.href
                ||
                anchor.getAttribute(
                    "href"
                )
                ||
                "";

            if (
                href.startsWith("blob:")
                &&
                blobs.has(href)
            ) {

                evento.preventDefault();

                evento.stopImmediatePropagation();

                capturarAnchor(
                    anchor
                );
            }

        },
        true
    );


    const originalDispatch =
        EventTarget
            .prototype
            .dispatchEvent;

    EventTarget
        .prototype
        .dispatchEvent =
        function(evento) {

            try {

                if (
                    this instanceof HTMLAnchorElement
                    &&
                    evento
                    &&
                    evento.type === "click"
                ) {

                    const href =
                        this.href
                        ||
                        this.getAttribute(
                            "href"
                        )
                        ||
                        "";

                    if (
                        href.startsWith("blob:")
                        &&
                        blobs.has(href)
                    ) {

                        capturarAnchor(
                            this
                        );

                        return true;
                    }
                }

            } catch (erro) {
            }

            return originalDispatch
                .apply(
                    this,
                    arguments
                );
        };


    return true;
}
'''

    frames_instaladas = 0

    for frame_atual in page.frames:

        try:

            resultado = frame_atual.evaluate(
                script_interceptador
            )

            if resultado:
                frames_instaladas += 1

        except Exception:
            pass

    print(
        "Interceptador de Blob instalado "
        f"em {frames_instaladas} frame(s)."
    )

    estado_download = {
        "detectado": False,
        "erro": "",
    }

    def ao_download(download):

        estado_download[
            "detectado"
        ] = True

        try:

            print(
                "Download nativo detectado: "
                f"{download.suggested_filename}"
            )

            download.save_as(
                str(
                    ARQUIVO_TEMPORARIO
                )
            )

            print(
                "Download persistido "
                "pelo Playwright."
            )

        except Exception as erro:

            estado_download[
                "erro"
            ] = str(erro)

            print(
                "[AVISO] O download nativo "
                "foi detectado, mas save_as "
                "falhou."
            )

    page.on(
        "download",
        ao_download,
    )

    exportar.first.click()

    limite = time.time() + 180

    while time.time() < limite:

        if (
            ARQUIVO_TEMPORARIO.exists()
            and
            ARQUIVO_TEMPORARIO.stat().st_size > 0
        ):

            break

        try:

            page.wait_for_timeout(
                250
            )

        except Exception:

            time.sleep(
                0.25
            )

        if (
            estado_blob["erro"]
            and
            estado_download["erro"]
        ):

            break

    if not ARQUIVO_TEMPORARIO.exists():

        detalhes = []

        if estado_blob["erro"]:

            detalhes.append(
                "Blob: "
                + estado_blob["erro"]
            )

        if estado_download["erro"]:

            detalhes.append(
                "Playwright: "
                + estado_download["erro"]
            )

        detalhe = (
            " | ".join(
                detalhes
            )
            or
            "nenhum arquivo foi persistido"
        )

        raise RuntimeError(
            "Falha ao persistir o Excel "
            "exportado pelo QuickSight. "
            + detalhe
        )

    tamanho = (
        ARQUIVO_TEMPORARIO
        .stat()
        .st_size
    )

    if tamanho <= 0:

        raise RuntimeError(
            "O arquivo temporario "
            "possui tamanho zero."
        )

    print(
        "Arquivo temporario persistido: "
        f"{tamanho:,} bytes"
    )

    print(
        "Download salvo temporariamente em:"
    )

    print(
        ARQUIVO_TEMPORARIO
    )

    return ARQUIVO_TEMPORARIO


# =========================================================
# VALIDACAO XLSX
# =========================================================

def validar_estrutura_xlsx(arquivo):

    if not arquivo.exists():
        raise RuntimeError(
            "O arquivo exportado não foi encontrado."
        )

    tamanho = arquivo.stat().st_size

    if tamanho <= 0:
        raise RuntimeError(
            "O arquivo exportado possui tamanho zero."
        )

    if not zipfile.is_zipfile(
        arquivo
    ):
        raise RuntimeError(
            "O arquivo recebido não é um XLSX válido."
        )

    return tamanho


def validar_excel_inspecoes(arquivo):

    print(
        "Validando Excel exportado..."
    )

    tamanho = validar_estrutura_xlsx(
        arquivo
    )

    workbook = load_workbook(
        arquivo,
        read_only=True,
        data_only=True,
    )

    try:

        if not workbook.sheetnames:
            raise RuntimeError(
                "O Excel não possui abas."
            )

        worksheet = workbook[
            workbook.sheetnames[0]
        ]

        linha_cabecalho = None
        cabecalho = None

        # O arquivo do QuickSight possui título
        # antes do cabeçalho. Procuramos o cabeçalho
        # nas primeiras 10 linhas.
        for numero_linha, valores in enumerate(
            worksheet.iter_rows(
                min_row=1,
                max_row=10,
                values_only=True,
            ),
            start=1,
        ):

            valores_texto = {
                str(valor).strip()
                for valor in valores
                if valor is not None
            }

            if "ID resposta" in valores_texto:
                linha_cabecalho = numero_linha
                cabecalho = valores_texto
                break

        if linha_cabecalho is None:
            raise RuntimeError(
                "Não foi encontrado o cabeçalho "
                "'ID resposta' no Excel."
            )

        ausentes = (
            COLUNAS_OBRIGATORIAS
            - cabecalho
        )

        if ausentes:
            raise RuntimeError(
                "O Excel não contém todas as colunas "
                "mínimas esperadas. Ausentes: "
                + ", ".join(
                    sorted(ausentes)
                )
            )

        ids = []

        for valores in worksheet.iter_rows(
            min_row=linha_cabecalho + 1,
            values_only=True,
        ):

            id_resposta = valores[0]

            if id_resposta is None:
                continue

            ids.append(
                str(id_resposta).strip()
            )

    finally:

        workbook.close()

    if not ids:
        raise RuntimeError(
            "Nenhuma inspeção foi encontrada "
            "no arquivo exportado."
        )

    total_linhas = len(ids)
    total_ids_unicos = len(
        set(ids)
    )

    print(
        "Excel de inspeções válido."
    )

    print(
        f"Linhas de dados: {total_linhas}"
    )

    print(
        f"IDs de resposta únicos: "
        f"{total_ids_unicos}"
    )

    if total_linhas != total_ids_unicos:

        print(
            "[AVISO] Existem IDs de resposta "
            "duplicados na Bronze."
        )

        print(
            "[AVISO] A Bronze será preservada "
            "sem deduplicação."
        )

    print(
        f"Tamanho: {tamanho:,} bytes"
    )


# =========================================================
# PROMOCAO PARA BRONZE
# =========================================================

def atualizar_bronze(
    arquivo_baixado
):

    if not arquivo_baixado.exists():
        raise RuntimeError(
            "Arquivo temporário não encontrado."
        )

    os.replace(
        arquivo_baixado,
        ARQUIVO_FINAL,
    )

    print()
    print(
        "Camada Bronze atualizada com sucesso:"
    )
    print(
        ARQUIVO_FINAL
    )


def limpar_temporario():

    try:

        if ARQUIVO_TEMPORARIO.exists():
            ARQUIVO_TEMPORARIO.unlink()

    except Exception:
        pass


# =========================================================
# EXECUCAO PRINCIPAL
# =========================================================

def extrair_quantidade_inspecoes():

    print(
        "=" * 70
    )

    print(
        "EXTRACAO INSPECOES - QUANTIDADE - PREVSYS"
    )

    print(
        "=" * 70
    )

    print()

    PASTA_BRONZE.mkdir(
        parents=True,
        exist_ok=True,
    )

    try:

        with sync_playwright() as p:

            print(
                "Abrindo Microsoft Edge com "
                "perfil persistente..."
            )

            context = (
                p.chromium.launch_persistent_context(
                    user_data_dir=str(
                        PROFILE_DIR
                    ),
                    channel="msedge",
                    headless=False,
                    accept_downloads=True,
                    viewport=None,
                    chromium_sandbox=True,
                    args=[
                        "--start-maximized",
                    ],
                )
            )

            try:

                page = (
                    context.pages[0]
                    if context.pages
                    else context.new_page()
                )

                garantir_login(
                    page
                )

                abrir_reporte_360(
                    page
                )

                abrir_lista_verificacao(
                    page
                )

                arquivo = exportar_excel(
                    page
                )

                validar_excel_inspecoes(
                    arquivo
                )

                atualizar_bronze(
                    arquivo
                )

            finally:

                try:
                    context.close()
                except Exception:
                    pass

        print()
        print(
            "=" * 70
        )

        print(
            "EXTRACAO DE INSPECOES CONCLUIDA COM SUCESSO"
        )

        print(
            "=" * 70
        )

    except Exception:

        limpar_temporario()

        raise


if __name__ == "__main__":
    extrair_quantidade_inspecoes()
