from pathlib import Path
from collections import Counter, defaultdict
from datetime import datetime, date
import csv
import json
import os
import re
import unicodedata

from openpyxl import load_workbook, Workbook


RAIZ = Path(__file__).resolve().parents[2]

ARQUIVO_BRONZE = (
    RAIZ
    / "data"
    / "bronze"
    / "inspecoes"
    / "quantidade"
    / "Detalle_Checklist.xlsx"
)

ARQUIVO_CONFIG = (
    RAIZ
    / "backend"
    / "inspecoes"
    / "config"
    / "areas_escopo.csv"
)

PASTA_PRATA = (
    RAIZ
    / "data"
    / "prata"
    / "inspecoes"
    / "quantidade"
)

ARQUIVO_PRATA = (
    PASTA_PRATA
    / "Inspecoes_Quantidade.xlsx"
)

ARQUIVO_TEMP = (
    PASTA_PRATA
    / "Inspecoes_Quantidade_novo.xlsx"
)

ARQUIVO_IDENTIDADES_COLABORADORES = (
    RAIZ
    / "backend"
    / "inspecoes"
    / "config"
    / "identidades_colaboradores.json"
)


def normalizar(valor):

    if valor is None:
        return ""

    texto = str(valor).strip()

    texto = unicodedata.normalize(
        "NFKD",
        texto,
    )

    texto = "".join(
        caractere
        for caractere in texto
        if not unicodedata.combining(caractere)
    )

    return (
        texto
        .lower()
        .strip()
    )


def slug(valor):

    texto = normalizar(valor)

    texto = re.sub(
        r"[^a-z0-9]+",
        "_",
        texto,
    )

    return texto.strip("_")


def id_normalizado(valor):

    if valor is None:
        return ""

    if isinstance(valor, float):

        if valor.is_integer():
            return str(int(valor))

    if isinstance(valor, int):
        return str(valor)

    texto = str(valor).strip()

    if texto.endswith(".0"):

        parte = texto[:-2]

        if parte.isdigit():
            return parte

    return texto


def converter_data(valor):

    if valor is None:
        return None

    if isinstance(valor, datetime):
        return valor

    if isinstance(valor, date):

        return datetime(
            valor.year,
            valor.month,
            valor.day,
        )

    texto = str(valor).strip()

    formatos = [
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y",
    ]

    for formato in formatos:

        try:
            return datetime.strptime(
                texto,
                formato,
            )
        except ValueError:
            continue

    return None


def extrair_identificador_criador(valor):

    texto = str(
        valor or ""
    ).strip()

    if not texto:

        raise RuntimeError(
            "Criador de checklist vazio."
        )

    identificador = (
        texto
        .split(
            maxsplit=1
        )[0]
        .strip()
        .upper()
    )

    if "-" not in identificador:

        raise RuntimeError(
            "Formato inesperado de criador: "
            f"{texto}"
        )

    return identificador


def carregar_identidades_colaboradores():

    if not ARQUIVO_IDENTIDADES_COLABORADORES.exists():

        raise FileNotFoundError(
            "Configuracao de identidades "
            "de Inspecoes nao encontrada: "
            f"{ARQUIVO_IDENTIDADES_COLABORADORES}"
        )

    with ARQUIVO_IDENTIDADES_COLABORADORES.open(
        "r",
        encoding="utf-8",
    ) as arquivo:

        configuracao = json.load(
            arquivo
        )

    if configuracao.get(
        "versao"
    ) != 1:

        raise RuntimeError(
            "Versao invalida da configuracao "
            "de identidades."
        )

    sistema_origem = str(
        configuracao.get(
            "sistema_origem",
            "",
        )
    ).strip().upper()

    if sistema_origem != "INSPECOES":

        raise RuntimeError(
            "sistema_origem invalido "
            "na configuracao de identidades: "
            f"{sistema_origem}"
        )

    registros = configuracao.get(
        "identidades"
    )

    if not isinstance(
        registros,
        list,
    ):

        raise RuntimeError(
            "Campo identidades deve ser "
            "uma lista."
        )

    obrigatorias = {
        "codigo_inspecao",
        "cod_colaborador",
        "nome_exibicao",
        "status_vinculo",
        "metodo_vinculo",
    }

    identidades = {}

    for item in registros:

        if not isinstance(
            item,
            dict,
        ):

            raise RuntimeError(
                "Identidade invalida "
                "na configuracao."
            )

        faltantes = (
            obrigatorias
            - set(
                item.keys()
            )
        )

        if faltantes:

            raise RuntimeError(
                "Campos ausentes na identidade: "
                + ", ".join(
                    sorted(
                        faltantes
                    )
                )
            )

        codigo_inspecao = str(
            item.get(
                "codigo_inspecao",
                "",
            )
        ).strip().upper()

        cod_colaborador = str(
            item.get(
                "cod_colaborador",
                "",
            )
        ).strip()

        nome_exibicao = str(
            item.get(
                "nome_exibicao",
                "",
            )
        ).strip()

        status_vinculo = str(
            item.get(
                "status_vinculo",
                "",
            )
        ).strip().upper()

        metodo_vinculo = str(
            item.get(
                "metodo_vinculo",
                "",
            )
        ).strip().upper()

        if not codigo_inspecao:

            raise RuntimeError(
                "codigo_inspecao vazio "
                "na configuracao."
            )

        if codigo_inspecao in identidades:

            raise RuntimeError(
                "codigo_inspecao duplicado: "
                f"{codigo_inspecao}"
            )

        if status_vinculo == "VINCULADO":

            if not cod_colaborador:

                raise RuntimeError(
                    "Identidade VINCULADA sem "
                    "cod_colaborador: "
                    f"{codigo_inspecao}"
                )

            if not nome_exibicao:

                raise RuntimeError(
                    "Identidade VINCULADA sem "
                    "nome_exibicao: "
                    f"{codigo_inspecao}"
                )

            if metodo_vinculo not in {
                "EXATO",
                "SIMILARIDADE_VALIDADA",
            }:

                raise RuntimeError(
                    "Metodo invalido para "
                    "identidade VINCULADA: "
                    f"{codigo_inspecao} -> "
                    f"{metodo_vinculo}"
                )

        elif status_vinculo == "NAO_VINCULADO":

            if cod_colaborador:

                raise RuntimeError(
                    "Identidade NAO_VINCULADA "
                    "possui cod_colaborador: "
                    f"{codigo_inspecao}"
                )

            if nome_exibicao:

                raise RuntimeError(
                    "Identidade NAO_VINCULADA "
                    "possui nome_exibicao: "
                    f"{codigo_inspecao}"
                )

            if metodo_vinculo != "NAO_VINCULADO":

                raise RuntimeError(
                    "Metodo invalido para "
                    "identidade NAO_VINCULADA: "
                    f"{codigo_inspecao}"
                )

        else:

            raise RuntimeError(
                "status_vinculo invalido: "
                f"{codigo_inspecao} -> "
                f"{status_vinculo}"
            )

        identidades[
            codigo_inspecao
        ] = {
            "cod_colaborador":
                cod_colaborador,
            "nome_exibicao":
                nome_exibicao,
            "status_vinculo":
                status_vinculo,
            "metodo_vinculo":
                metodo_vinculo,
        }

    if not identidades:

        raise RuntimeError(
            "Nenhuma identidade cadastrada "
            "para Inspecoes."
        )

    return identidades


def carregar_dimensao():

    if not ARQUIVO_CONFIG.exists():

        raise FileNotFoundError(
            f"Dimensao nao encontrada: "
            f"{ARQUIVO_CONFIG}"
        )

    dimensao = {}

    with ARQUIVO_CONFIG.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as arquivo:

        leitor = csv.DictReader(
            arquivo,
            delimiter=";",
        )

        for linha in leitor:

            if (
                str(
                    linha.get(
                        "ativo",
                        "",
                    )
                ).strip().upper()
                != "SIM"
            ):
                continue

            chave = (
                linha.get(
                    "chave_normalizada",
                    ""
                )
                or normalizar(
                    linha.get(
                        "area_organizacional",
                        ""
                    )
                )
            )

            chave = normalizar(
                chave
            )

            escopo = (
                linha.get(
                    "escopo",
                    ""
                )
                .strip()
                .upper()
            )

            if escopo not in {
                "BRASIL",
                "FORA_BRASIL",
            }:

                raise RuntimeError(
                    "Escopo invalido na dimensao: "
                    f"{escopo!r}"
                )

            if chave in dimensao:

                if (
                    dimensao[chave]
                    != escopo
                ):

                    raise RuntimeError(
                        "Classificacao conflitante "
                        f"para area: {chave}"
                    )

            dimensao[chave] = escopo

    if not dimensao:

        raise RuntimeError(
            "Dimensao de areas esta vazia."
        )

    return dimensao


def ler_bronze():

    if not ARQUIVO_BRONZE.exists():

        raise FileNotFoundError(
            f"Bronze nao encontrada: "
            f"{ARQUIVO_BRONZE}"
        )

    wb = load_workbook(
        ARQUIVO_BRONZE,
        read_only=True,
        data_only=True,
    )

    ws = wb[
        wb.sheetnames[0]
    ]

    linha_cabecalho = None
    cabecalho = None

    for numero_linha, linha in enumerate(
        ws.iter_rows(
            min_row=1,
            max_row=10,
            values_only=True,
        ),
        start=1,
    ):

        valores = [
            normalizar(valor)
            for valor in linha
        ]

        if "id resposta" in valores:

            linha_cabecalho = (
                numero_linha
            )

            cabecalho = [
                str(valor).strip()
                if valor is not None
                else ""
                for valor in linha
            ]

            break

    if linha_cabecalho is None:

        wb.close()

        raise RuntimeError(
            "Cabecalho real da Bronze "
            "nao foi localizado."
        )

    colunas = [
        slug(coluna)
        for coluna in cabecalho
    ]

    obrigatorias = {
        "id_resposta",
        "nome_formulario",
        "data_da_inspecao",
        "area_organizacional",
        "local",
        "criador_de_checklist",
    }

    faltantes = (
        obrigatorias
        - set(colunas)
    )

    if faltantes:

        wb.close()

        raise RuntimeError(
            "Colunas obrigatorias ausentes: "
            + ", ".join(
                sorted(faltantes)
            )
        )

    registros = []

    for numero_linha, linha in enumerate(
        ws.iter_rows(
            min_row=linha_cabecalho + 1,
            values_only=True,
        ),
        start=linha_cabecalho + 1,
    ):

        registro = {
            coluna: (
                linha[indice]
                if indice < len(linha)
                else None
            )
            for indice, coluna
            in enumerate(colunas)
        }

        registro[
            "_linha_bronze"
        ] = numero_linha

        registros.append(
            registro
        )

    wb.close()

    return (
        colunas,
        registros,
    )


def extrair_prefixo_formulario(
    formulario,
):

    texto = str(
        formulario
        or ""
    ).strip()

    match = re.match(
        r"^\s*([A-Za-z]{2,5})\s*-\s+",
        texto,
    )

    if not match:
        return ""

    return (
        match
        .group(1)
        .upper()
    )


def carregar_regras_compostas():

    arquivo_regras = (
        RAIZ
        / "backend"
        / "inspecoes"
        / "config"
        / "regras_escopo.csv"
    )

    if not arquivo_regras.exists():

        raise FileNotFoundError(
            "Arquivo de regras compostas "
            f"nao encontrado: {arquivo_regras}"
        )

    regras = {}

    with arquivo_regras.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as arquivo:

        leitor = csv.DictReader(
            arquivo,
            delimiter=";",
        )

        for linha in leitor:

            if (
                str(
                    linha.get(
                        "ativo",
                        ""
                    )
                )
                .strip()
                .upper()
                != "SIM"
            ):
                continue

            chave_area = normalizar(
                linha.get(
                    "chave_area_normalizada",
                    ""
                )
                or
                linha.get(
                    "area_organizacional",
                    ""
                )
            )

            prefixo = (
                str(
                    linha.get(
                        "prefixo_formulario",
                        ""
                    )
                )
                .strip()
                .upper()
            )

            escopo = (
                str(
                    linha.get(
                        "escopo",
                        ""
                    )
                )
                .strip()
                .upper()
            )

            if not chave_area:
                raise RuntimeError(
                    "Regra composta sem area."
                )

            if not prefixo:
                raise RuntimeError(
                    "Regra composta sem prefixo."
                )

            if escopo not in {
                "BRASIL",
                "FORA_BRASIL",
            }:

                raise RuntimeError(
                    "Escopo invalido em "
                    "regras_escopo.csv: "
                    f"{escopo!r}"
                )

            chave = (
                chave_area,
                prefixo,
            )

            if chave in regras:

                if (
                    regras[chave]
                    != escopo
                ):

                    raise RuntimeError(
                        "Regra composta conflitante: "
                        f"{chave}"
                    )

            regras[
                chave
            ] = escopo

    if not regras:

        raise RuntimeError(
            "Nenhuma regra composta ativa."
        )

    return regras


def classificar_registros(
    registros,
    dimensao,
    regras_compostas,
):

    desconhecidas = Counter()

    exemplos = defaultdict(
        set
    )

    classificadas_area = 0
    classificadas_regra = 0

    for registro in registros:

        area = registro.get(
            "area_organizacional"
        )

        formulario = registro.get(
            "nome_formulario"
        )

        chave_area = normalizar(
            area
        )

        prefixo = (
            extrair_prefixo_formulario(
                formulario
            )
        )

        registro[
            "_prefixo_formulario"
        ] = prefixo

        # ---------------------------------------------
        # REGRA 1
        # Area historicamente validada pelo filtro
        # oficial do Prevsys.
        # ---------------------------------------------

        if chave_area in dimensao:

            registro[
                "_escopo_calculado"
            ] = dimensao[
                chave_area
            ]

            registro[
                "_fonte_classificacao"
            ] = "AREA_HISTORICA"

            classificadas_area += 1

            continue

        # ---------------------------------------------
        # REGRA 2
        # Area nova somente entra se existir regra
        # composta explicitamente validada.
        # ---------------------------------------------

        chave_regra = (
            chave_area,
            prefixo,
        )

        if (
            prefixo
            and
            chave_regra
            in regras_compostas
        ):

            registro[
                "_escopo_calculado"
            ] = regras_compostas[
                chave_regra
            ]

            registro[
                "_fonte_classificacao"
            ] = "AREA_PREFIXO"

            classificadas_regra += 1

            continue

        # ---------------------------------------------
        # QUALITY GATE
        # ---------------------------------------------

        area_exibicao = (
            str(area).strip()
            if area is not None
            else "<AREA VAZIA>"
        )

        prefixo_exibicao = (
            prefixo
            or "<SEM PREFIXO>"
        )

        chave_erro = (
            area_exibicao,
            prefixo_exibicao,
        )

        desconhecidas[
            chave_erro
        ] += 1

        exemplos[
            chave_erro
        ].add(
            str(
                formulario
                or ""
            )
        )

    if desconhecidas:

        print()
        print("=" * 70)
        print("QUALITY GATE BLOQUEADO")
        print("=" * 70)

        print(
            "Foram encontradas combinacoes "
            "sem classificacao:"
        )

        for (
            area_prefixo,
            quantidade,
        ) in desconhecidas.most_common():

            area, prefixo = (
                area_prefixo
            )

            print()

            print(
                f"Area: {area}"
            )

            print(
                f"Prefixo: {prefixo}"
            )

            print(
                f"Registros: {quantidade}"
            )

            for formulario in sorted(
                exemplos[
                    area_prefixo
                ]
            )[:5]:

                if formulario:

                    print(
                        f"  - {formulario}"
                    )

        print()

        print(
            "A Prata NAO sera atualizada."
        )

        raise RuntimeError(
            "Existem combinacoes de "
            "area/prefixo nao classificadas."
        )

    print(
        "Quality Gate de escopo: APROVADO"
    )

    print(
        "Classificados por area historica: "
        f"{classificadas_area}"
    )

    print(
        "Classificados por regra "
        "area + prefixo: "
        f"{classificadas_regra}"
    )


def selecionar_brasil(
    registros,
):

    brasil = []
    fora = 0

    for registro in registros:

        escopo = registro.get(
            "_escopo_calculado"
        )

        if escopo == "BRASIL":

            brasil.append(
                registro
            )

        elif escopo == "FORA_BRASIL":

            fora += 1

        else:

            raise RuntimeError(
                "Registro sem escopo calculado."
            )

    return brasil, fora


def deduplicar(
    registros,
):

    grupos = defaultdict(list)

    for registro in registros:

        identificador = (
            id_normalizado(
                registro.get(
                    "id_resposta"
                )
            )
        )

        if not identificador:

            raise RuntimeError(
                "Registro Brasil sem "
                "ID resposta encontrado na "
                f"linha Bronze "
                f"{registro['_linha_bronze']}."
            )

        registro[
            "id_resposta"
        ] = identificador

        grupos[
            identificador
        ].append(
            registro
        )

    campos_criticos = [
        "nome_formulario",
        "contrato",
        "tipo_de_registro",
        "data_da_inspecao",
        "area_organizacional",
        "local",
        "empresa",
    ]

    conflitos = []
    resultado = []

    ids_duplicados = 0
    linhas_duplicadas = 0

    for identificador, grupo in (
        grupos.items()
    ):

        if len(grupo) > 1:

            ids_duplicados += 1

            linhas_duplicadas += (
                len(grupo) - 1
            )

            for campo in campos_criticos:

                valores = {
                    normalizar(
                        registro.get(
                            campo
                        )
                    )
                    for registro in grupo
                }

                if len(valores) > 1:

                    conflitos.append(
                        (
                            identificador,
                            campo,
                            valores,
                        )
                    )

        escolhido = dict(
            grupo[0]
        )

        escolhido[
            "ocorrencias_id_origem"
        ] = len(grupo)

        escolhido[
            "duplicado_origem"
        ] = (
            "SIM"
            if len(grupo) > 1
            else "NAO"
        )

        resultado.append(
            escolhido
        )

    if conflitos:

        print()
        print("=" * 70)
        print("QUALITY GATE DE DUPLICIDADE BLOQUEADO")
        print("=" * 70)

        for (
            identificador,
            campo,
            valores,
        ) in conflitos[:20]:

            print(
                f"ID {identificador} | "
                f"campo {campo}: "
                f"{sorted(valores)}"
            )

        raise RuntimeError(
            "Existem IDs duplicados com "
            "dados criticos conflitantes."
        )

    print(
        "Quality Gate de duplicidade: APROVADO"
    )

    print(
        f"IDs duplicados encontrados: "
        f"{ids_duplicados}"
    )

    print(
        f"Linhas excedentes removidas: "
        f"{linhas_duplicadas}"
    )

    return resultado


def preparar_prata(
    registros,
    identidades_colaboradores,
):

    resultado = []

    for registro in registros:

        novo = {
            chave: valor
            for chave, valor
            in registro.items()
            if not chave.startswith("_")
        }

        identificador_criador = (
            extrair_identificador_criador(
                registro.get(
                    "criador_de_checklist"
                )
            )
        )

        identidade = (
            identidades_colaboradores.get(
                identificador_criador
            )
        )

        if identidade is None:

            raise RuntimeError(
                "Criador ainda nao cadastrado "
                "em identidades_colaboradores.json: "
                f"{identificador_criador}"
            )

        novo[
            "identificador_criador"
        ] = identificador_criador

        novo[
            "cod_colaborador"
        ] = identidade[
            "cod_colaborador"
        ]

        novo[
            "status_vinculo_colaborador"
        ] = identidade[
            "status_vinculo"
        ]

        data_inspecao = converter_data(
            registro.get(
                "data_da_inspecao"
            )
        )

        novo[
            "pais_escopo"
        ] = "BRASIL"

        novo[
            "prefixo_formulario"
        ] = registro.get(
            "_prefixo_formulario",
            "",
        )

        novo[
            "fonte_classificacao_escopo"
        ] = registro.get(
            "_fonte_classificacao",
            "",
        )

        novo[
            "data_inspecao_normalizada"
        ] = data_inspecao

        if data_inspecao:

            novo[
                "ano_inspecao"
            ] = data_inspecao.year

            novo[
                "mes_inspecao"
            ] = data_inspecao.month

            novo[
                "ano_mes"
            ] = (
                f"{data_inspecao.year:04d}-"
                f"{data_inspecao.month:02d}"
            )

        else:

            novo[
                "ano_inspecao"
            ] = None

            novo[
                "mes_inspecao"
            ] = None

            novo[
                "ano_mes"
            ] = None

        novo[
            "quantidade_inspecao"
        ] = 1

        resultado.append(
            novo
        )

    return resultado


def escrever_prata(
    registros,
):

    PASTA_PRATA.mkdir(
        parents=True,
        exist_ok=True,
    )

    if ARQUIVO_TEMP.exists():
        ARQUIVO_TEMP.unlink()

    wb = Workbook()

    ws = wb.active
    ws.title = "inspecoes"

    if not registros:

        raise RuntimeError(
            "Nenhum registro Brasil "
            "foi encontrado."
        )

    colunas = list(
        registros[0].keys()
    )

    ws.append(
        colunas
    )

    for registro in registros:

        ws.append(
            [
                registro.get(
                    coluna
                )
                for coluna in colunas
            ]
        )

    ws.freeze_panes = "A2"

    ws.auto_filter.ref = (
        ws.dimensions
    )

    wb.save(
        ARQUIVO_TEMP
    )

    wb.close()

    return colunas


def validar_prata(
    quantidade_esperada,
):

    wb = load_workbook(
        ARQUIVO_TEMP,
        read_only=True,
        data_only=True,
    )

    ws = wb[
        wb.sheetnames[0]
    ]

    linhas = list(
        ws.iter_rows(
            values_only=True
        )
    )

    wb.close()

    if not linhas:

        raise RuntimeError(
            "Prata temporaria esta vazia."
        )

    cabecalho = [
        str(valor)
        if valor is not None
        else ""
        for valor in linhas[0]
    ]

    dados = linhas[1:]

    if (
        len(dados)
        != quantidade_esperada
    ):

        raise RuntimeError(
            "Quantidade da Prata diverge "
            "da quantidade esperada."
        )

    if "id_resposta" not in cabecalho:

        raise RuntimeError(
            "Prata sem coluna id_resposta."
        )

    indice_id = cabecalho.index(
        "id_resposta"
    )

    ids = [
        id_normalizado(
            linha[indice_id]
        )
        for linha in dados
    ]

    if len(ids) != len(set(ids)):

        raise RuntimeError(
            "A Prata ainda possui "
            "IDs resposta duplicados."
        )

    if "pais_escopo" not in cabecalho:

        raise RuntimeError(
            "Prata sem coluna pais_escopo."
        )

    indice_pais = cabecalho.index(
        "pais_escopo"
    )

    paises = {
        str(
            linha[indice_pais]
        ).strip().upper()
        for linha in dados
    }

    if paises != {"BRASIL"}:

        raise RuntimeError(
            "A Prata contem registros "
            "fora do escopo Brasil."
        )

    # =====================================================
    # VALIDACAO DA IDENTIDADE DO COLABORADOR
    # =====================================================

    colunas_identidade = {
        "identificador_criador",
        "cod_colaborador",
        "status_vinculo_colaborador",
    }

    faltantes_identidade = (
        colunas_identidade
        - set(cabecalho)
    )

    if faltantes_identidade:

        raise RuntimeError(
            "Prata sem colunas de identidade: "
            + ", ".join(
                sorted(
                    faltantes_identidade
                )
            )
        )

    indice_identificador = cabecalho.index(
        "identificador_criador"
    )

    indice_colaborador = cabecalho.index(
        "cod_colaborador"
    )

    indice_status = cabecalho.index(
        "status_vinculo_colaborador"
    )

    vinculadas = 0
    nao_vinculadas = 0

    for numero_linha, linha in enumerate(
        dados,
        start=2,
    ):

        identificador = str(
            linha[indice_identificador]
            or ""
        ).strip().upper()

        cod_colaborador = str(
            linha[indice_colaborador]
            or ""
        ).strip()

        status = str(
            linha[indice_status]
            or ""
        ).strip().upper()

        if not identificador:

            raise RuntimeError(
                "Prata possui identificador_criador "
                "vazio na linha "
                f"{numero_linha}."
            )

        if not re.fullmatch(
            r"[0-9]+-[0-9K]",
            identificador,
        ):

            raise RuntimeError(
                "identificador_criador invalido "
                "na linha "
                f"{numero_linha}: "
                f"{identificador}"
            )

        if status == "VINCULADO":

            vinculadas += 1

            if not cod_colaborador:

                raise RuntimeError(
                    "Registro VINCULADO sem "
                    "cod_colaborador na linha "
                    f"{numero_linha}."
                )

            if not re.fullmatch(
                r"[0-9]+-[0-9]+",
                cod_colaborador,
            ):

                raise RuntimeError(
                    "cod_colaborador invalido "
                    "na linha "
                    f"{numero_linha}: "
                    f"{cod_colaborador}"
                )

        elif status == "NAO_VINCULADO":

            nao_vinculadas += 1

            if cod_colaborador:

                raise RuntimeError(
                    "Registro NAO_VINCULADO possui "
                    "cod_colaborador na linha "
                    f"{numero_linha}: "
                    f"{cod_colaborador}"
                )

        else:

            raise RuntimeError(
                "status_vinculo_colaborador "
                "invalido na linha "
                f"{numero_linha}: "
                f"{status}"
            )

    if (
        vinculadas
        + nao_vinculadas
        != len(dados)
    ):

        raise RuntimeError(
            "Total de vinculos nao corresponde "
            "ao total de registros da Prata."
        )

    print(
        "Inspecoes vinculadas: "
        f"{vinculadas}"
    )

    print(
        "Inspecoes nao vinculadas: "
        f"{nao_vinculadas}"
    )

    print(
        "Validacao da Prata: APROVADA"
    )


def atualizar_prata():

    try:

        dimensao = carregar_dimensao()

        regras_compostas = (
            carregar_regras_compostas()
        )

        print(
            "Areas classificadas na dimensao: "
            f"{len(dimensao)}"
        )

        print(
            "Regras compostas ativas: "
            f"{len(regras_compostas)}"
        )

        (
            colunas_bronze,
            registros,
        ) = ler_bronze()

        print(
            "Registros recebidos da Bronze: "
            f"{len(registros)}"
        )

        classificar_registros(
            registros,
            dimensao,
            regras_compostas,
        )

        brasil, fora = selecionar_brasil(
            registros
        )

        print(
            "Registros classificados BRASIL: "
            f"{len(brasil)}"
        )

        print(
            "Registros classificados FORA_BRASIL: "
            f"{fora}"
        )

        deduplicados = deduplicar(
            brasil
        )

        print(
            "Inspecoes Brasil unicas: "
            f"{len(deduplicados)}"
        )

        identidades_colaboradores = (
            carregar_identidades_colaboradores()
        )

        prata = preparar_prata(
            deduplicados,
            identidades_colaboradores,
        )

        escrever_prata(
            prata
        )

        validar_prata(
            len(prata)
        )

        os.replace(
            ARQUIVO_TEMP,
            ARQUIVO_PRATA,
        )

        print()

        print("=" * 70)
        print(
            "CAMADA PRATA ATUALIZADA COM SUCESSO"
        )
        print("=" * 70)

        print(
            f"Arquivo: {ARQUIVO_PRATA}"
        )

        print(
            f"Inspecoes Brasil: {len(prata)}"
        )

    except Exception:

        if ARQUIVO_TEMP.exists():

            try:
                ARQUIVO_TEMP.unlink()
            except Exception:
                pass

        raise


if __name__ == "__main__":

    print("=" * 70)
    print("TRATAMENTO INSPECOES - QUANTIDADE")
    print("=" * 70)
    print()

    atualizar_prata()
