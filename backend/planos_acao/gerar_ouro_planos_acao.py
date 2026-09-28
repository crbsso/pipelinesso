from pathlib import Path
from openpyxl import load_workbook, Workbook
from collections import Counter
from datetime import datetime, date
import unicodedata
import os


# =========================================================
# CAMINHOS
# =========================================================

RAIZ = Path(
    __file__
).resolve().parents[2]


ARQUIVO_PRATA = (
    RAIZ
    / "data"
    / "prata"
    / "planos_acao"
    / "Planos_de_acao.xlsx"
)


PASTA_OURO = (
    RAIZ
    / "data"
    / "ouro"
    / "planos_acao"
)


ARQUIVO_OURO = (
    PASTA_OURO
    / "mart_planos_acao.xlsx"
)


ARQUIVO_TEMP = (
    PASTA_OURO
    / "mart_planos_acao_novo.xlsx"
)


# =========================================================
# NORMALIZACAO
# =========================================================

def normalizar(valor):

    if valor is None:
        return ""

    texto = unicodedata.normalize(
        "NFKD",
        str(valor),
    )

    texto = "".join(
        caractere
        for caractere in texto
        if not unicodedata.combining(
            caractere
        )
    )

    return texto.strip().lower()


# =========================================================
# LEITURA DA PRATA
# =========================================================

def carregar_prata():

    if not ARQUIVO_PRATA.exists():

        raise RuntimeError(
            "Arquivo da Prata nao encontrado: "
            f"{ARQUIVO_PRATA}"
        )


    wb = load_workbook(
        ARQUIVO_PRATA,
        read_only=True,
        data_only=True,
    )


    try:

        ws = wb[
            wb.sheetnames[0]
        ]


        cabecalho = [
            celula.value
            for celula in ws[1]
        ]


        indices = {
            nome: indice
            for indice, nome
            in enumerate(cabecalho)
        }


        obrigatorios = {
            "plano_id",
            "id_resposta",
            "nome_plano",
            "modulo_origem",
            "inicio_plano",
            "criacao_plano",
            "cronograma_plano",
            "encerramento_plano",
            "nivel_organizacional",
            "divisao_plano",
            "estado_plano",
            "criador_plano",
            "criador_plano_variantes",
            "status_criador",
            "gestor_plano",
            "descricao_plano",
            "contrato_plano",
            "grupo_sso",
            "metodo_classificacao",
            "registros_origem",
            "duplicado_origem",
        }


        ausentes = (
            obrigatorios
            - set(indices)
        )


        if ausentes:

            raise RuntimeError(
                "Prata sem colunas obrigatorias: "
                + ", ".join(
                    sorted(ausentes)
                )
            )


        registros = []


        for linha in ws.iter_rows(
            min_row=2,
            values_only=True,
        ):

            if not any(
                valor is not None
                and str(valor).strip()
                for valor in linha
            ):

                continue


            registro = {
                coluna: (
                    linha[
                        indices[coluna]
                    ]
                    if indices[coluna] < len(linha)
                    else None
                )
                for coluna
                in obrigatorios
            }


            registros.append(
                registro
            )


    finally:

        wb.close()


    return registros


# =========================================================
# AUXILIAR DE DATA
# =========================================================

def obter_data(valor):

    if valor is None:

        return None


    if isinstance(
        valor,
        datetime,
    ):

        return valor.date()


    if isinstance(
        valor,
        date,
    ):

        return valor


    return None


# =========================================================
# CONSTRUCAO DA MART
# =========================================================

def construir_mart(
    prata,
):

    hoje = date.today()

    mart = []


    for registro in prata:

        estado_fonte = (
            registro[
                "estado_plano"
            ]
        )


        estado_norm = normalizar(
            estado_fonte
        )


        ativo_fonte = (
            estado_norm
            in {
                "aberto",
                "atrasado",
            }
        )


        prazo = obter_data(
            registro[
                "cronograma_plano"
            ]
        )


        dias_para_prazo = None
        dias_em_atraso = None
        situacao_prazo = (
            "NAO_APLICAVEL"
        )

        flag_vencido_calculado = 0
        flag_em_prazo_calculado = 0
        flag_vence_hoje = 0


        if ativo_fonte:

            if prazo is None:

                situacao_prazo = (
                    "SEM_PRAZO"
                )

            else:

                dias_para_prazo = (
                    prazo
                    - hoje
                ).days


                if dias_para_prazo < 0:

                    situacao_prazo = (
                        "VENCIDO"
                    )

                    dias_em_atraso = abs(
                        dias_para_prazo
                    )

                    flag_vencido_calculado = 1


                elif dias_para_prazo == 0:

                    situacao_prazo = (
                        "VENCE_HOJE"
                    )

                    dias_em_atraso = 0
                    flag_vence_hoje = 1


                else:

                    situacao_prazo = (
                        "EM_PRAZO"
                    )

                    dias_em_atraso = 0
                    flag_em_prazo_calculado = 1


        # -------------------------------------------------
        # Divergencia entre status PREVSYS e data calculada
        # -------------------------------------------------

        divergencia = 0


        if (
            estado_norm == "atrasado"
            and
            situacao_prazo
            in {
                "EM_PRAZO",
                "VENCE_HOJE",
            }
        ):

            divergencia = 1


        if (
            estado_norm == "aberto"
            and
            situacao_prazo
            == "VENCIDO"
        ):

            divergencia = 1


        mart.append(
            {
                "plano_id":
                    registro[
                        "plano_id"
                    ],

                "id_resposta":
                    registro[
                        "id_resposta"
                    ],

                "nome_plano":
                    registro[
                        "nome_plano"
                    ],

                "modulo_origem":
                    registro[
                        "modulo_origem"
                    ],

                "inicio_plano":
                    registro[
                        "inicio_plano"
                    ],

                "criacao_plano":
                    registro[
                        "criacao_plano"
                    ],

                "cronograma_plano":
                    registro[
                        "cronograma_plano"
                    ],

                "encerramento_plano":
                    registro[
                        "encerramento_plano"
                    ],

                "nivel_organizacional":
                    registro[
                        "nivel_organizacional"
                    ],

                "divisao_plano":
                    registro[
                        "divisao_plano"
                    ],

                # -----------------------------------------
                # STATUS OFICIAL DA FONTE
                # -----------------------------------------

                "estado_plano_fonte":
                    estado_fonte,

                "flag_ativo_fonte":
                    int(
                        ativo_fonte
                    ),

                "flag_aberto_fonte":
                    int(
                        estado_norm
                        == "aberto"
                    ),

                "flag_atrasado_fonte":
                    int(
                        estado_norm
                        == "atrasado"
                    ),

                "flag_fechado_fonte":
                    int(
                        estado_norm
                        == "fechado"
                    ),

                "flag_fechado_apos_prazo_fonte":
                    int(
                        estado_norm
                        == "fechado apos o prazo"
                    ),

                # -----------------------------------------
                # PRAZO CALCULADO
                # -----------------------------------------

                "data_referencia":
                    hoje,

                "dias_para_prazo":
                    dias_para_prazo,

                "dias_em_atraso_calculado":
                    dias_em_atraso,

                "situacao_prazo_calculada":
                    situacao_prazo,

                "flag_vencido_calculado":
                    flag_vencido_calculado,

                "flag_em_prazo_calculado":
                    flag_em_prazo_calculado,

                "flag_vence_hoje":
                    flag_vence_hoje,

                "divergencia_status_prazo":
                    divergencia,

                # -----------------------------------------
                # RESPONSAVEIS
                # -----------------------------------------

                "criador_plano":
                    registro[
                        "criador_plano"
                    ],

                "criador_plano_variantes":
                    registro[
                        "criador_plano_variantes"
                    ],

                "status_criador":
                    registro[
                        "status_criador"
                    ],

                "gestor_plano":
                    registro[
                        "gestor_plano"
                    ],

                # -----------------------------------------
                # CLASSIFICACAO SSO
                # -----------------------------------------

                "grupo_sso":
                    registro[
                        "grupo_sso"
                    ],

                "metodo_classificacao":
                    registro[
                        "metodo_classificacao"
                    ],

                # -----------------------------------------
                # DETALHES
                # -----------------------------------------

                "descricao_plano":
                    registro[
                        "descricao_plano"
                    ],

                "contrato_plano":
                    registro[
                        "contrato_plano"
                    ],

                # -----------------------------------------
                # LINHAGEM
                # -----------------------------------------

                "registros_origem":
                    registro[
                        "registros_origem"
                    ],

                "duplicado_origem":
                    registro[
                        "duplicado_origem"
                    ],
            }
        )


    return mart


# =========================================================
# QUALITY GATE DA OURO
# =========================================================

def validar_mart(
    mart,
):

    if not mart:

        raise RuntimeError(
            "Mart vazia."
        )


    ids = [
        str(
            registro[
                "plano_id"
            ]
        )
        for registro
        in mart
    ]


    if len(ids) != len(
        set(ids)
    ):

        raise RuntimeError(
            "Existem plano_id duplicados "
            "na Ouro."
        )


    if len(mart) != 204:

        raise RuntimeError(
            "Quantidade inesperada na Ouro. "
            f"Esperado: 204. Recebido: {len(mart)}"
        )


    grupos = Counter(
        registro[
            "grupo_sso"
        ]
        for registro
        in mart
    )


    if groups := (
        grupos
        != {
            "SEGURANCA": 107,
            "OUTRAS_AREAS": 97,
        }
    ):

        raise RuntimeError(
            "Distribuicao de grupos SSO "
            "diferente do contrato esperado: "
            f"{dict(grupos)}"
        )


    estados = Counter(
        normalizar(
            registro[
                "estado_plano_fonte"
            ]
        )
        for registro
        in mart
    )


    esperado_estados = {
        "aberto": 47,
        "atrasado": 75,
        "fechado": 48,
        "fechado apos o prazo": 34,
    }


    if estados != esperado_estados:

        raise RuntimeError(
            "Distribuicao de estados diferente "
            "do contrato esperado: "
            f"{dict(estados)}"
        )


    ativos_sem_prazo = [
        registro[
            "plano_id"
        ]
        for registro
        in mart
        if (
            registro[
                "flag_ativo_fonte"
            ]
            == 1
            and
            registro[
                "cronograma_plano"
            ]
            is None
        )
    ]


    if ativos_sem_prazo:

        raise RuntimeError(
            "Existem planos ativos sem prazo: "
            + ", ".join(
                ativos_sem_prazo
            )
        )


    divergentes = {
        str(
            registro[
                "plano_id"
            ]
        )
        for registro
        in mart
        if registro[
            "divergencia_status_prazo"
        ]
        == 1
    }


    if divergentes != {
        "509",
        "290",
    }:

        raise RuntimeError(
            "Conjunto inesperado de divergencias "
            "status x prazo: "
            f"{sorted(divergentes)}"
        )


    print(
        "Quality Gate da Ouro: APROVADO"
    )


# =========================================================
# ESCRITA ATOMICA
# =========================================================

def escrever_mart(
    mart,
):

    PASTA_OURO.mkdir(
        parents=True,
        exist_ok=True,
    )


    if ARQUIVO_TEMP.exists():

        ARQUIVO_TEMP.unlink()


    wb = Workbook()

    ws = wb.active

    ws.title = (
        "mart_planos_acao"
    )


    colunas = list(
        mart[0].keys()
    )


    ws.append(
        colunas
    )


    for registro in mart:

        ws.append(
            [
                registro.get(
                    coluna
                )
                for coluna
                in colunas
            ]
        )


    wb.save(
        ARQUIVO_TEMP
    )


    os.replace(
        ARQUIVO_TEMP,
        ARQUIVO_OURO,
    )


# =========================================================
# EXECUCAO
# =========================================================

def main():

    print("=" * 70)
    print("GERACAO OURO - PLANOS DE ACAO")
    print("=" * 70)


    prata = carregar_prata()


    print()
    print(
        "Registros recebidos da Prata:",
        len(prata),
    )


    mart = construir_mart(
        prata
    )


    validar_mart(
        mart
    )


    grupos = Counter(
        registro[
            "grupo_sso"
        ]
        for registro
        in mart
    )


    estados = Counter(
        normalizar(
            registro[
                "estado_plano_fonte"
            ]
        )
        for registro
        in mart
    )


    situacoes = Counter(
        registro[
            "situacao_prazo_calculada"
        ]
        for registro
        in mart
    )


    divergencias = [
        registro
        for registro
        in mart
        if registro[
            "divergencia_status_prazo"
        ]
        == 1
    ]


    print()
    print("Grupos SSO:")

    for chave, qtd in sorted(
        grupos.items()
    ):

        print(
            f"  {chave}: {qtd}"
        )


    print()
    print("Estados da fonte:")

    for chave, qtd in sorted(
        estados.items()
    ):

        print(
            f"  {chave}: {qtd}"
        )


    print()
    print("Situacao de prazo calculada:")

    for chave, qtd in sorted(
        situacoes.items()
    ):

        print(
            f"  {chave}: {qtd}"
        )


    print()
    print(
        "Divergencias status x prazo:",
        len(divergencias),
    )


    for registro in divergencias:

        print(
            "  Plano "
            f"{registro['plano_id']} | "
            f"fonte={registro['estado_plano_fonte']} | "
            f"prazo={registro['cronograma_plano']} | "
            f"situacao={registro['situacao_prazo_calculada']}"
        )


    escrever_mart(
        mart
    )


    print()
    print("=" * 70)
    print(
        "CAMADA OURO DE PLANOS DE ACAO "
        "ATUALIZADA COM SUCESSO"
    )
    print("=" * 70)

    print(
        "Arquivo:",
        ARQUIVO_OURO,
    )

    print(
        "Planos:",
        len(mart),
    )


if __name__ == "__main__":

    main()
