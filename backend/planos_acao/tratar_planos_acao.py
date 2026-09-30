from pathlib import Path
from openpyxl import load_workbook, Workbook
from collections import Counter
from datetime import datetime
import unicodedata
import csv
import json
import os
import re


# =========================================================
# CAMINHOS
# =========================================================

RAIZ = Path(
    __file__
).resolve().parents[2]


ARQUIVO_BRONZE = (
    RAIZ
    / "data"
    / "bronze"
    / "planos_acao"
    / "Planos_de_acao_detalhados.xlsx"
)


PASTA_PRATA = (
    RAIZ
    / "data"
    / "prata"
    / "planos_acao"
)


ARQUIVO_PRATA = (
    PASTA_PRATA
    / "Planos_de_acao.xlsx"
)


ARQUIVO_TEMP = (
    PASTA_PRATA
    / "Planos_de_acao_novo.xlsx"
)


ARQUIVO_MAPA_COLABORADORES = (
    RAIZ
    / "data"
    / "master"
    / "colaboradores"
    / "mapa_colaboradores.csv"
)

ARQUIVO_ALIASES = (
    RAIZ
    / "backend"
    / "planos_acao"
    / "config"
    / "aliases_gestores.json"
)


# =========================================================
# NORMALIZACAO
# =========================================================

def normalizar(valor):

    if valor is None:
        return ""

    texto = str(valor)

    texto = texto.replace(
        "\xa0",
        " ",
    )

    texto = unicodedata.normalize(
        "NFKD",
        texto,
    )

    texto = "".join(
        caractere
        for caractere in texto
        if not unicodedata.combining(
            caractere
        )
    )

    texto = re.sub(
        r"\s+",
        " ",
        texto,
    )

    return texto.strip().lower()


def texto_limpo(valor):

    if valor is None:
        return None

    if not isinstance(
        valor,
        str,
    ):
        return valor

    valor = valor.replace(
        "\xa0",
        " ",
    )

    valor = re.sub(
        r"\s+",
        " ",
        valor,
    )

    valor = valor.strip()

    return valor or None


# =========================================================
# CONFIGURACAO
# =========================================================

def carregar_mapa_colaboradores():

    if not ARQUIVO_MAPA_COLABORADORES.exists():

        raise RuntimeError(
            "Mapa de colaboradores "
            "nao encontrado: "
            f"{ARQUIVO_MAPA_COLABORADORES}"
        )


    with ARQUIVO_MAPA_COLABORADORES.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as arquivo_csv:

        leitor = csv.DictReader(
            arquivo_csv,
            delimiter=";",
        )

        campos = set(
            leitor.fieldnames
            or []
        )

        obrigatorios = {
            "cod_colaborador",
            "nome_normalizado",
            "nome_completo_rm",
            "secao_rm",
            "elegivel_planos_acao",
            "origem_registro",
            "data_demissao",
            "data_limite_elegibilidade",
        }

        ausentes = (
            obrigatorios
            - campos
        )

        if ausentes:

            raise RuntimeError(
                "Mapa de colaboradores "
                "sem colunas obrigatorias: "
                + ", ".join(
                    sorted(ausentes)
                )
            )


        mapa = {}


        for linha in leitor:

            nome = normalizar(
                linha.get(
                    "nome_normalizado"
                )
                or linha.get(
                    "nome_completo_rm"
                )
            )


            if not nome:

                raise RuntimeError(
                    "Mapa de colaboradores "
                    "possui nome vazio."
                )


            if nome in mapa:

                raise RuntimeError(
                    "Mapa de colaboradores "
                    "possui nome duplicado: "
                    f"{nome}"
                )


            mapa[nome] = {
                "cod_colaborador":
                    texto_limpo(
                        linha.get(
                            "cod_colaborador"
                        )
                    ),

                "nome_completo_rm":
                    texto_limpo(
                        linha.get(
                            "nome_completo_rm"
                        )
                    ),

                "secao_rm":
                    texto_limpo(
                        linha.get(
                            "secao_rm"
                        )
                    ),

                "elegivel_planos_acao":
                    texto_limpo(
                        linha.get(
                            "elegivel_planos_acao"
                        )
                    ),

                "origem_registro":
                    texto_limpo(
                        linha.get(
                            "origem_registro"
                        )
                    ),

                "data_demissao":
                    texto_limpo(
                        linha.get(
                            "data_demissao"
                        )
                    ),

                "data_limite_elegibilidade":
                    texto_limpo(
                        linha.get(
                            "data_limite_elegibilidade"
                        )
                    ),
            }


    if not mapa:

        raise RuntimeError(
            "Mapa de colaboradores vazio."
        )


    return mapa


def carregar_aliases_gestores(
    mapa_colaboradores,
):

    if not ARQUIVO_ALIASES.exists():

        raise RuntimeError(
            "Arquivo de aliases "
            "nao encontrado: "
            f"{ARQUIVO_ALIASES}"
        )


    with ARQUIVO_ALIASES.open(
        "r",
        encoding="utf-8",
    ) as arquivo_json:

        dados = json.load(
            arquivo_json
        )


    aliases = {}


    for item in dados.get(
        "aliases",
        [],
    ):

        origem = normalizar(
            item.get(
                "nome_prevsys"
            )
        )

        destino = normalizar(
            item.get(
                "nome_rm"
            )
        )


        if not origem:

            raise RuntimeError(
                "Alias possui "
                "nome PREVSYS vazio."
            )


        if not destino:

            raise RuntimeError(
                "Alias possui "
                "nome RM vazio."
            )


        if origem in aliases:

            raise RuntimeError(
                "Alias PREVSYS duplicado: "
                f"{origem}"
            )


        if destino not in mapa_colaboradores:

            raise RuntimeError(
                "Destino de alias nao existe "
                "no mapa RM: "
                f"{destino}"
            )


        aliases[origem] = destino


    return aliases

def carregar_bronze():

    if not ARQUIVO_BRONZE.exists():

        raise RuntimeError(
            "Bronze nao encontrada: "
            f"{ARQUIVO_BRONZE}"
        )


    wb = load_workbook(
        ARQUIVO_BRONZE,
        read_only=True,
        data_only=True,
    )


    try:

        ws = wb[
            wb.sheetnames[0]
        ]


        cabecalho = [
            celula.value
            for celula in ws[3]
        ]


        indices = {
            normalizar(nome): indice
            for indice, nome
            in enumerate(cabecalho)
            if nome is not None
        }


        obrigatorios = {
            "plano de identificacao",
            "id da resposta",
            "nome do plano de acao",
            "modulo de origem",
            "inicio do plano",
            "criacao do plano",
            "cronograma do plano",
            "encerramento do plano",
            "nivel organizacional",
            "divisao do plano",
            "estado do plano",
            "criador do plano",
            "gestor de plano",
            "descricao do plano de acao",
            "contrato de plano de acao",
        }


        ausentes = (
            obrigatorios
            - set(indices)
        )


        if ausentes:

            raise RuntimeError(
                "Bronze sem colunas obrigatorias: "
                + ", ".join(
                    sorted(ausentes)
                )
            )


        def valor(
            linha,
            coluna,
        ):

            indice = indices[
                coluna
            ]

            if indice >= len(linha):
                return None

            return linha[
                indice
            ]


        registros = []


        for linha in ws.iter_rows(
            min_row=4,
            values_only=True,
        ):

            if not any(
                item is not None
                and str(item).strip()
                for item in linha
            ):
                continue


            registro = {
                coluna: valor(
                    linha,
                    coluna,
                )
                for coluna
                in obrigatorios
            }


            registros.append(
                registro
            )


    finally:

        wb.close()


    return registros


# =========================================================
# ESCOPO BRASIL
# =========================================================

def filtrar_brasil(
    registros,
):

    brasil = [
        registro
        for registro
        in registros
        if normalizar(
            registro[
                "divisao do plano"
            ]
        )
        == "celeo brasil"
    ]


    if not brasil:

        raise RuntimeError(
            "Nenhum registro Celeo Brasil "
            "foi encontrado."
        )


    return brasil


# =========================================================
# DEDUPLICACAO
# =========================================================

def deduplicar(
    registros,
):

    grupos = {}


    for registro in registros:

        plano = registro[
            "plano de identificacao"
        ]


        if plano is None:

            raise RuntimeError(
                "Plano sem identificador "
                "na Bronze."
            )


        plano_id = str(
            plano
        ).strip()


        grupos.setdefault(
            plano_id,
            [],
        ).append(
            registro
        )


    resultado = []

    ids_duplicados = 0
    linhas_excedentes = 0


    campos_comparacao = {
        "plano de identificacao",
        "id da resposta",
        "nome do plano de acao",
        "modulo de origem",
        "inicio do plano",
        "criacao do plano",
        "cronograma do plano",
        "encerramento do plano",
        "nivel organizacional",
        "divisao do plano",
        "estado do plano",
        "gestor de plano",
        "descricao do plano de acao",
        "contrato de plano de acao",
    }


    for plano_id, linhas in grupos.items():

        base = linhas[0]


        if len(linhas) > 1:

            ids_duplicados += 1

            linhas_excedentes += (
                len(linhas) - 1
            )


            for outra in linhas[1:]:

                diferentes = []


                for campo in campos_comparacao:

                    if (
                        texto_limpo(
                            base.get(
                                campo
                            )
                        )
                        !=
                        texto_limpo(
                            outra.get(
                                campo
                            )
                        )
                    ):

                        diferentes.append(
                            campo
                        )


                if diferentes:

                    raise RuntimeError(
                        "Quality Gate de duplicidade "
                        f"bloqueado no plano {plano_id}. "
                        "Campos divergentes: "
                        + ", ".join(
                            sorted(
                                diferentes
                            )
                        )
                    )


        criadores = []


        for linha in linhas:

            criador = texto_limpo(
                linha.get(
                    "criador do plano"
                )
            )


            if (
                criador
                and
                criador not in criadores
            ):

                criadores.append(
                    criador
                )


        if len(criadores) == 1:

            criador_plano = (
                criadores[0]
            )

            status_criador = (
                "UNICO"
            )

        elif len(criadores) > 1:

            # Nao escolhemos arbitrariamente
            # uma das identidades da fonte.
            criador_plano = None

            status_criador = (
                "MULTIPLAS_VARIANTES"
            )

        else:

            criador_plano = None

            status_criador = (
                "AUSENTE"
            )


        resultado.append(
            {
                "registro":
                    base,

                "plano_id":
                    plano_id,

                "criador_plano":
                    criador_plano,

                "criador_plano_variantes":
                    " | ".join(
                        criadores
                    )
                    if criadores
                    else None,

                "status_criador":
                    status_criador,

                "registros_origem":
                    len(linhas),

                "duplicado_origem":
                    (
                        "SIM"
                        if len(linhas) > 1
                        else "NAO"
                    ),
            }
        )


    print(
        "Quality Gate de duplicidade: APROVADO"
    )

    print(
        f"IDs duplicados encontrados: "
        f"{ids_duplicados}"
    )

    print(
        f"Linhas excedentes removidas: "
        f"{linhas_excedentes}"
    )


    return resultado


# =========================================================
# CLASSIFICACAO SSO
# =========================================================

def classificar_sso(
    gestor,
    mapa_colaboradores,
    aliases_gestores,
):

    gestor_norm = normalizar(
        gestor
    )


    if not gestor_norm:

        return (
            None,
            "SEM_GESTOR",
        )


    if gestor_norm in mapa_colaboradores:

        nome_rm = gestor_norm
        metodo = "MAPA_RM_EXATO"

    elif gestor_norm in aliases_gestores:

        nome_rm = (
            aliases_gestores[
                gestor_norm
            ]
        )

        metodo = "MAPA_RM_ALIAS"

    else:

        return (
            None,
            "SEM_IDENTIDADE_RM",
        )


    cadastro = (
        mapa_colaboradores[
            nome_rm
        ]
    )


    if (
        cadastro[
            "elegivel_planos_acao"
        ]
        != "1"
    ):

        return (
            None,
            "NAO_ELEGIVEL_RM",
        )


    secao = texto_limpo(
        cadastro[
            "secao_rm"
        ]
    )


    if not secao:

        return (
            None,
            "SECAO_RM_NAO_RESOLVIDA",
        )


    secao_norm = normalizar(
        secao
    )


    if (
        secao_norm
        == "seguranca e saude ocupacional"
    ):

        grupo = "SEGURANCA"

    else:

        grupo = "OUTRAS_AREAS"


    return (
        grupo,
        metodo,
    )

def preparar_prata(
    registros_deduplicados,
    mapa_colaboradores,
    aliases_gestores,
):

    prata = []
    excluidos = []


    for item in registros_deduplicados:

        origem = item[
            "registro"
        ]


        grupo_sso, metodo = (
            classificar_sso(
                origem.get(
                    "gestor de plano"
                ),
                mapa_colaboradores,
                aliases_gestores,
            )
        )


        if grupo_sso is None:

            excluidos.append(
                {
                    "plano_id":
                        item.get(
                            "plano_id",
                            "",
                        ),

                    "gestor_plano":
                        texto_limpo(
                            origem.get(
                                "gestor de plano"
                            )
                        ),

                    "motivo":
                        metodo,
                }
            )

            continue


        prata.append(
            {
                "plano_id":
                    item[
                        "plano_id"
                    ],

                "id_resposta":
                    origem.get(
                        "id da resposta"
                    ),

                "nome_plano":
                    texto_limpo(
                        origem.get(
                            "nome do plano de acao"
                        )
                    ),

                "modulo_origem":
                    texto_limpo(
                        origem.get(
                            "modulo de origem"
                        )
                    ),

                "inicio_plano":
                    origem.get(
                        "inicio do plano"
                    ),

                "criacao_plano":
                    origem.get(
                        "criacao do plano"
                    ),

                "cronograma_plano":
                    origem.get(
                        "cronograma do plano"
                    ),

                "encerramento_plano":
                    origem.get(
                        "encerramento do plano"
                    ),

                "nivel_organizacional":
                    texto_limpo(
                        origem.get(
                            "nivel organizacional"
                        )
                    ),

                "divisao_plano":
                    texto_limpo(
                        origem.get(
                            "divisao do plano"
                        )
                    ),

                "estado_plano":
                    texto_limpo(
                        origem.get(
                            "estado do plano"
                        )
                    ),

                "criador_plano":
                    item[
                        "criador_plano"
                    ],

                "criador_plano_variantes":
                    item[
                        "criador_plano_variantes"
                    ],

                "status_criador":
                    item[
                        "status_criador"
                    ],

                "gestor_plano":
                    texto_limpo(
                        origem.get(
                            "gestor de plano"
                        )
                    ),

                "descricao_plano":
                    texto_limpo(
                        origem.get(
                            "descricao do plano de acao"
                        )
                    ),

                "contrato_plano":
                    texto_limpo(
                        origem.get(
                            "contrato de plano de acao"
                        )
                    ),

                "grupo_sso":
                    grupo_sso,

                "metodo_classificacao":
                    metodo,

                "registros_origem":
                    item[
                        "registros_origem"
                    ],

                "duplicado_origem":
                    item[
                        "duplicado_origem"
                    ],
            }
        )


    print()
    print("=" * 70)
    print(
        "FILTRO DE IDENTIDADE RM "
        "- PLANOS DE ACAO"
    )
    print("=" * 70)

    print(
        "Planos recebidos:",
        len(
            registros_deduplicados
        ),
    )

    print(
        "Planos elegiveis:",
        len(prata),
    )

    print(
        "Planos excluidos:",
        len(excluidos),
    )


    if excluidos:

        motivos = Counter(
            item[
                "motivo"
            ]
            for item
            in excluidos
        )

        print()
        print(
            "Motivos de exclusao:"
        )

        for motivo, quantidade in (
            sorted(
                motivos.items()
            )
        ):

            print(
                f"  {motivo}: "
                f"{quantidade}"
            )


        print()
        print(
            "Planos excluidos:"
        )

        for item in excluidos:

            print(
                "  "
                f"plano_id="
                f"{item['plano_id']} | "
                f"gestor="
                f"{item['gestor_plano']} | "
                f"motivo="
                f"{item['motivo']}"
            )


    return prata


# =========================================================
# VALIDACAO PRATA
# =========================================================

def validar_prata(
    prata,
):

    if not prata:

        raise RuntimeError(
            "Prata vazia."
        )


    ids = [
        registro[
            "plano_id"
        ]
        for registro
        in prata
    ]


    if len(ids) != len(
        set(ids)
    ):

        raise RuntimeError(
            "Existem plano_id duplicados "
            "na Prata."
        )


    grupos_validos = {
        "SEGURANCA",
        "OUTRAS_AREAS",
    }


    invalidos = [
        registro
        for registro
        in prata
        if registro[
            "grupo_sso"
        ]
        not in grupos_validos
    ]


    if invalidos:

        raise RuntimeError(
            "Existem registros sem "
            "classificacao SSO valida."
        )


    if any(
        normalizar(
            registro[
                "divisao_plano"
            ]
        )
        != "celeo brasil"
        for registro
        in prata
    ):

        raise RuntimeError(
            "A Prata contem registros "
            "fora de Celeo Brasil."
        )


    print(
        "Validacao da Prata: APROVADA"
    )


# =========================================================
# ESCRITA ATOMICA
# =========================================================

def escrever_prata(
    registros,
):

    PASTA_PRATA.mkdir(
        parents=True,
        exist_ok=True,
    )


    if ARQUIVO_TEMP.exists():

        ARQUIVO_TEMP.unlink()


    wb = Workbook()

    ws = wb.active

    ws.title = (
        "planos_acao"
    )


    colunas = list(
        registros[0].keys()
    )


    ws.append(
        colunas
    )


    for registro in registros:

        ws.append(
            [
                registro.get(
                    coluna
                )
                for coluna
                in colunas
            ]
        )


    wb.save(
        ARQUIVO_TEMP
    )

    os.replace(
        ARQUIVO_TEMP,
        ARQUIVO_PRATA,
    )


# =========================================================
# EXECUCAO
# =========================================================

def main():

    print("=" * 70)
    print("TRATAMENTO PLANOS DE ACAO - PRATA")
    print("=" * 70)


    mapa_colaboradores = (
        carregar_mapa_colaboradores()
    )

    aliases_gestores = (
        carregar_aliases_gestores(
            mapa_colaboradores
        )
    )


    print()
    print(
        "Colaboradores disponiveis no mapa RM:",
        len(mapa_colaboradores),
    )


    print(
        "Aliases de gestores:",
        len(aliases_gestores),
    )


    bronze = carregar_bronze()


    print(
        "Registros recebidos da Bronze:",
        len(bronze),
    )


    brasil = filtrar_brasil(
        bronze
    )


    print(
        "Registros Celeo Brasil:",
        len(brasil),
    )


    deduplicados = deduplicar(
        brasil
    )


    print(
        "Planos Brasil unicos:",
        len(deduplicados),
    )


    prata = preparar_prata(
        deduplicados,
        mapa_colaboradores,
        aliases_gestores,
    )


    validar_prata(
        prata
    )


    grupos = Counter(
        registro[
            "grupo_sso"
        ]
        for registro
        in prata
    )


    metodos = Counter(
        registro[
            "metodo_classificacao"
        ]
        for registro
        in prata
    )


    estados = Counter(
        registro[
            "estado_plano"
        ]
        for registro
        in prata
    )


    criadores = Counter(
        registro[
            "status_criador"
        ]
        for registro
        in prata
    )


    print()
    print("Classificacao SSO:")

    for chave, quantidade in sorted(
        grupos.items()
    ):

        print(
            f"  {chave}: {quantidade}"
        )


    print()
    print("Metodos de classificacao:")

    for chave, quantidade in sorted(
        metodos.items()
    ):

        print(
            f"  {chave}: {quantidade}"
        )


    print()
    print("Estados:")

    for chave, quantidade in sorted(
        estados.items()
    ):

        print(
            f"  {chave}: {quantidade}"
        )


    print()
    print("Status do criador:")

    for chave, quantidade in sorted(
        criadores.items()
    ):

        print(
            f"  {chave}: {quantidade}"
        )


    escrever_prata(
        prata
    )


    print()
    print(
        "Camada Prata atualizada com sucesso:"
    )

    print(
        ARQUIVO_PRATA
    )


    print()
    print(
        "Data de processamento:",
        datetime.now().isoformat(
            timespec="seconds"
        ),
    )


if __name__ == "__main__":

    main()
