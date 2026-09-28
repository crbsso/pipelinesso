import argparse
import os
import unicodedata
import uuid
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import load_workbook
from supabase import create_client


# =========================================================
# CONFIGURACAO
# =========================================================

RAIZ = Path(
    __file__
).resolve().parents[2]


ARQUIVO_OURO = (
    RAIZ
    / "data"
    / "ouro"
    / "planos_acao"
    / "mart_planos_acao.xlsx"
)


NOME_ABA = (
    "mart_planos_acao"
)


TABELA_MART = (
    "mart_planos_acao"
)


TABELA_STAGING = (
    "stg_mart_planos_acao"
)


RPC_PROMOVER = (
    "promover_mart_planos_acao"
)


TAMANHO_LOTE = 500

TAMANHO_PAGINA = 1000


COLUNAS_ESPERADAS = [
    "plano_id",
    "id_resposta",
    "nome_plano",
    "modulo_origem",
    "inicio_plano",
    "criacao_plano",
    "cronograma_plano",
    "encerramento_plano",
    "nivel_organizacional",
    "divisao_plano",
    "estado_plano_fonte",
    "flag_ativo_fonte",
    "flag_aberto_fonte",
    "flag_atrasado_fonte",
    "flag_fechado_fonte",
    "flag_fechado_apos_prazo_fonte",
    "data_referencia",
    "dias_para_prazo",
    "dias_em_atraso_calculado",
    "situacao_prazo_calculada",
    "flag_vencido_calculado",
    "flag_em_prazo_calculado",
    "flag_vence_hoje",
    "divergencia_status_prazo",
    "criador_plano",
    "criador_plano_variantes",
    "status_criador",
    "gestor_plano",
    "grupo_sso",
    "metodo_classificacao",
    "descricao_plano",
    "contrato_plano",
    "registros_origem",
    "duplicado_origem",
]


COLUNAS_DATA = {
    "inicio_plano",
    "criacao_plano",
    "cronograma_plano",
    "encerramento_plano",
    "data_referencia",
}


COLUNAS_INTEIRO = {
    "id_resposta",
    "flag_ativo_fonte",
    "flag_aberto_fonte",
    "flag_atrasado_fonte",
    "flag_fechado_fonte",
    "flag_fechado_apos_prazo_fonte",
    "dias_para_prazo",
    "dias_em_atraso_calculado",
    "flag_vencido_calculado",
    "flag_em_prazo_calculado",
    "flag_vence_hoje",
    "divergencia_status_prazo",
    "registros_origem",
}


# =========================================================
# CONFIGURACAO DO SUPABASE
# =========================================================

def carregar_configuracao():

    load_dotenv(
        RAIZ / ".env"
    )


    supabase_url = os.getenv(
        "SUPABASE_URL"
    )


    supabase_secret_key = os.getenv(
        "SUPABASE_SECRET_KEY"
    )


    if not supabase_url:

        raise RuntimeError(
            "SUPABASE_URL nao encontrada "
            "no arquivo .env."
        )


    if not supabase_secret_key:

        raise RuntimeError(
            "SUPABASE_SECRET_KEY nao encontrada "
            "no arquivo .env."
        )


    return (
        supabase_url,
        supabase_secret_key,
    )


# =========================================================
# NORMALIZACAO
# =========================================================

def normalizar(
    valor,
):

    if valor is None:
        return ""


    texto = unicodedata.normalize(
        "NFKD",
        str(valor),
    )


    texto = "".join(
        caractere
        for caractere in texto
        if not unicodedata.combining(
            caractere
        )
    )


    return (
        " ".join(
            texto.split()
        )
        .strip()
        .lower()
    )


# =========================================================
# SERIALIZACAO
# =========================================================

def serializar_valor(
    coluna,
    valor,
):

    if valor is None:
        return None


    if coluna in COLUNAS_DATA:

        if isinstance(
            valor,
            datetime,
        ):

            return valor.date().isoformat()


        if isinstance(
            valor,
            date,
        ):

            return valor.isoformat()


        raise RuntimeError(
            "Valor de data inesperado em "
            f"{coluna}: {valor!r}"
        )


    if coluna in COLUNAS_INTEIRO:

        if valor is None:
            return None

        return int(
            valor
        )


    if isinstance(
        valor,
        str,
    ):

        valor = valor.strip()

        return (
            valor
            if valor
            else None
        )


    return valor


# =========================================================
# LEITURA DA OURO
# =========================================================

def carregar_ouro():

    if not ARQUIVO_OURO.exists():

        raise FileNotFoundError(
            "Arquivo Ouro nao encontrado: "
            f"{ARQUIVO_OURO}"
        )


    wb = load_workbook(
        ARQUIVO_OURO,
        read_only=True,
        data_only=True,
    )


    try:

        if NOME_ABA not in wb.sheetnames:

            raise RuntimeError(
                "Aba esperada nao encontrada "
                "na camada Ouro: "
                f"{NOME_ABA}"
            )


        ws = wb[
            NOME_ABA
        ]


        cabecalho = [
            celula.value
            for celula
            in ws[1]
        ]


        if (
            cabecalho
            != COLUNAS_ESPERADAS
        ):

            raise RuntimeError(
                "Contrato da Ouro diferente "
                "do esperado.\n"
                f"Esperado: {COLUNAS_ESPERADAS}\n"
                f"Recebido: {cabecalho}"
            )


        registros = []


        for linha in ws.iter_rows(
            min_row=2,
            values_only=True,
        ):

            if not any(
                valor is not None
                and str(valor).strip()
                for valor
                in linha
            ):

                continue


            registro = {}


            for indice, coluna in enumerate(
                COLUNAS_ESPERADAS
            ):

                valor = (
                    linha[indice]
                    if indice < len(linha)
                    else None
                )


                registro[
                    coluna
                ] = serializar_valor(
                    coluna,
                    valor,
                )


            registros.append(
                registro
            )


    finally:

        wb.close()


    return registros


# =========================================================
# QUALITY GATE LOCAL
# =========================================================

def validar_local(
    registros,
):

    if not registros:

        raise RuntimeError(
            "A camada Ouro esta vazia."
        )


    ids = [
        str(
            registro[
                "plano_id"
            ]
        )
        for registro
        in registros
    ]


    if len(ids) != len(
        set(ids)
    ):

        raise RuntimeError(
            "Existem plano_id duplicados "
            "na camada Ouro."
        )


    if any(
        not plano_id.strip()
        for plano_id
        in ids
    ):

        raise RuntimeError(
            "Existe plano_id vazio "
            "na camada Ouro."
        )


    grupos_validos = {
        "SEGURANCA",
        "OUTRAS_AREAS",
    }


    grupos = Counter(
        registro[
            "grupo_sso"
        ]
        for registro
        in registros
    )


    if (
        set(
            grupos
        )
        - grupos_validos
    ):

        raise RuntimeError(
            "Grupo SSO inesperado: "
            f"{dict(grupos)}"
        )


    estados_validos = {
        "aberto",
        "atrasado",
        "fechado",
        "fechado apos o prazo",
    }


    estados = Counter(
        normalizar(
            registro[
                "estado_plano_fonte"
            ]
        )
        for registro
        in registros
    )


    if (
        set(
            estados
        )
        - estados_validos
    ):

        raise RuntimeError(
            "Estado da fonte inesperado: "
            f"{dict(estados)}"
        )


    situacoes_validas = {
        "EM_PRAZO",
        "VENCIDO",
        "VENCE_HOJE",
        "SEM_PRAZO",
        "NAO_APLICAVEL",
    }


    situacoes = Counter(
        registro[
            "situacao_prazo_calculada"
        ]
        for registro
        in registros
    )


    if (
        set(
            situacoes
        )
        - situacoes_validas
    ):

        raise RuntimeError(
            "Situacao de prazo inesperada: "
            f"{dict(situacoes)}"
        )


    ativos_sem_prazo = [
        registro[
            "plano_id"
        ]
        for registro
        in registros
        if (
            registro[
                "flag_ativo_fonte"
            ]
            == 1
            and
            registro[
                "cronograma_plano"
            ]
            is None
        )
    ]


    if ativos_sem_prazo:

        raise RuntimeError(
            "Existem planos ativos "
            "sem cronograma: "
            + ", ".join(
                map(
                    str,
                    ativos_sem_prazo,
                )
            )
        )


    divergentes = {
        str(
            registro[
                "plano_id"
            ]
        )
        for registro
        in registros
        if registro[
            "divergencia_status_prazo"
        ]
        == 1
    }


    return {
        "total":
            len(registros),

        "ids":
            set(ids),

        "grupos":
            grupos,

        "estados":
            estados,

        "situacoes":
            situacoes,

        "divergentes":
            divergentes,
    }


# =========================================================
# CONTAGEM REMOTA
# =========================================================

def contar_remoto(
    cliente,
    tabela,
    carga_id=None,
):

    consulta = (
        cliente
        .table(
            tabela
        )
        .select(
            "plano_id",
            count="exact",
        )
    )


    if carga_id is not None:

        consulta = (
            consulta
            .eq(
                "carga_id",
                carga_id,
            )
        )


    resposta = (
        consulta
        .limit(1)
        .execute()
    )


    return (
        resposta.count
        if resposta.count is not None
        else 0
    )


# =========================================================
# LEITURA REMOTA DA MART
# =========================================================

def carregar_mart_remota(
    cliente,
):

    registros = []

    inicio = 0


    while True:

        fim = (
            inicio
            + TAMANHO_PAGINA
            - 1
        )


        resposta = (
            cliente
            .table(
                TABELA_MART
            )
            .select(
                ",".join(
                    [
                        "plano_id",
                        "grupo_sso",
                        "estado_plano_fonte",
                        "situacao_prazo_calculada",
                        "divergencia_status_prazo",
                    ]
                )
            )
            .range(
                inicio,
                fim,
            )
            .execute()
        )


        pagina = (
            resposta.data
            or []
        )


        registros.extend(
            pagina
        )


        if len(
            pagina
        ) < TAMANHO_PAGINA:

            break


        inicio += (
            TAMANHO_PAGINA
        )


    return registros


# =========================================================
# VALIDACAO REMOTA DA MART
# =========================================================

def validar_mart_remota(
    cliente,
    indicadores_locais,
):

    remoto = (
        carregar_mart_remota(
            cliente
        )
    )


    if len(
        remoto
    ) != indicadores_locais[
        "total"
    ]:

        raise RuntimeError(
            "Quantidade da Mart remota "
            "difere da Ouro local. "
            f"Local: "
            f"{indicadores_locais['total']} | "
            f"Remoto: {len(remoto)}"
        )


    ids_remotos = {
        str(
            registro[
                "plano_id"
            ]
        )
        for registro
        in remoto
    }


    if (
        ids_remotos
        != indicadores_locais[
            "ids"
        ]
    ):

        faltantes = (
            indicadores_locais[
                "ids"
            ]
            - ids_remotos
        )


        excedentes = (
            ids_remotos
            - indicadores_locais[
                "ids"
            ]
        )


        raise RuntimeError(
            "IDs da Mart remota "
            "diferem da Ouro local. "
            f"Faltantes: "
            f"{sorted(faltantes)} | "
            f"Excedentes: "
            f"{sorted(excedentes)}"
        )


    grupos_remotos = Counter(
        registro[
            "grupo_sso"
        ]
        for registro
        in remoto
    )


    if (
        grupos_remotos
        != indicadores_locais[
            "grupos"
        ]
    ):

        raise RuntimeError(
            "Distribuicao de grupos "
            "diferente apos publicacao. "
            f"Local: "
            f"{dict(indicadores_locais['grupos'])} | "
            f"Remoto: "
            f"{dict(grupos_remotos)}"
        )


    estados_remotos = Counter(
        normalizar(
            registro[
                "estado_plano_fonte"
            ]
        )
        for registro
        in remoto
    )


    if (
        estados_remotos
        != indicadores_locais[
            "estados"
        ]
    ):

        raise RuntimeError(
            "Distribuicao de estados "
            "diferente apos publicacao."
        )


    situacoes_remotas = Counter(
        registro[
            "situacao_prazo_calculada"
        ]
        for registro
        in remoto
    )


    if (
        situacoes_remotas
        != indicadores_locais[
            "situacoes"
        ]
    ):

        raise RuntimeError(
            "Distribuicao das situacoes "
            "de prazo diferente apos "
            "publicacao."
        )


    divergentes_remotos = {
        str(
            registro[
                "plano_id"
            ]
        )
        for registro
        in remoto
        if int(
            registro[
                "divergencia_status_prazo"
            ]
        )
        == 1
    }


    if (
        divergentes_remotos
        != indicadores_locais[
            "divergentes"
        ]
    ):

        raise RuntimeError(
            "Divergencias status x prazo "
            "diferentes apos publicacao."
        )


# =========================================================
# INSERCAO NA STAGING
# =========================================================

def carregar_staging(
    cliente,
    registros,
    carga_id,
):

    registros_staging = []


    for registro in registros:

        item = dict(
            registro
        )

        item[
            "carga_id"
        ] = carga_id

        registros_staging.append(
            item
        )


    for inicio in range(
        0,
        len(
            registros_staging
        ),
        TAMANHO_LOTE,
    ):

        lote = (
            registros_staging[
                inicio:
                inicio
                + TAMANHO_LOTE
            ]
        )


        (
            cliente
            .table(
                TABELA_STAGING
            )
            .insert(
                lote
            )
            .execute()
        )


# =========================================================
# LIMPEZA DE UMA CARGA
# =========================================================

def limpar_staging(
    cliente,
    carga_id,
):

    (
        cliente
        .table(
            TABELA_STAGING
        )
        .delete()
        .eq(
            "carga_id",
            carga_id,
        )
        .execute()
    )


    restante = contar_remoto(
        cliente,
        TABELA_STAGING,
        carga_id=carga_id,
    )


    if restante != 0:

        raise RuntimeError(
            "A carga permaneceu na "
            "staging apos a limpeza."
        )


# =========================================================
# PUBLICACAO
# =========================================================

def publicar(
    dry_run=False,
):

    print("=" * 78)
    print("PUBLICACAO OURO - PLANOS DE ACAO")
    print("=" * 78)


    print()
    print("1. Validando Mart local...")


    registros = carregar_ouro()


    indicadores = validar_local(
        registros
    )


    print(
        "   Registros:",
        indicadores[
            "total"
        ],
    )


    print(
        "   Planos unicos:",
        len(
            indicadores[
                "ids"
            ]
        ),
    )


    print(
        "   Grupos:",
        dict(
            indicadores[
                "grupos"
            ]
        ),
    )


    print(
        "   Estados:",
        dict(
            indicadores[
                "estados"
            ]
        ),
    )


    print(
        "   Situacao de prazo:",
        dict(
            indicadores[
                "situacoes"
            ]
        ),
    )


    print(
        "   Divergencias status x prazo:",
        sorted(
            indicadores[
                "divergentes"
            ]
        ),
    )


    print()
    print("2. Conectando ao Supabase...")


    (
        supabase_url,
        supabase_secret_key,
    ) = carregar_configuracao()


    cliente = create_client(
        supabase_url,
        supabase_secret_key,
    )


    print(
        "   Conexao configurada."
    )


    print()
    print(
        "3. Validando infraestrutura remota..."
    )


    quantidade_mart = contar_remoto(
        cliente,
        TABELA_MART,
    )


    quantidade_staging = contar_remoto(
        cliente,
        TABELA_STAGING,
    )


    print(
        "   mart_planos_acao:",
        quantidade_mart,
    )


    print(
        "   stg_mart_planos_acao:",
        quantidade_staging,
    )


    if dry_run:

        print()
        print("=" * 78)
        print(
            "DRY-RUN APROVADO - "
            "NENHUM DADO FOI ALTERADO"
        )
        print("=" * 78)

        return


    # =====================================================
    # PUBLICACAO REAL
    # =====================================================

    carga_id = str(
        uuid.uuid4()
    )

    promovida = False


    print()
    print(
        "4. Nova carga:"
    )

    print(
        "   carga_id:",
        carga_id,
    )


    try:

        print()
        print(
            "5. Carregando staging..."
        )


        carregar_staging(
            cliente,
            registros,
            carga_id,
        )


        quantidade_carga = (
            contar_remoto(
                cliente,
                TABELA_STAGING,
                carga_id=carga_id,
            )
        )


        print(
            "   Registros na carga:",
            quantidade_carga,
        )


        if (
            quantidade_carga
            != indicadores[
                "total"
            ]
        ):

            raise RuntimeError(
                "Quantidade na staging "
                "difere da Ouro local. "
                f"Esperado: "
                f"{indicadores['total']} | "
                f"Recebido: "
                f"{quantidade_carga}"
            )


        print(
            "   Quality Gate da staging: "
            "APROVADO"
        )


        print()
        print(
            "6. Promovendo carga "
            "atomicamente..."
        )


        resposta = (
            cliente
            .rpc(
                RPC_PROMOVER,
                {
                    "p_carga_id":
                        carga_id,

                    "p_quantidade_esperada":
                        indicadores[
                            "total"
                        ],
                },
            )
            .execute()
        )


        promovida = True


        print(
            "   RPC executada."
        )


        print(
            "   Retorno:",
            resposta.data,
        )


        print()
        print(
            "7. Validando Mart oficial..."
        )


        validar_mart_remota(
            cliente,
            indicadores,
        )


        print(
            "   Quality Gate remoto: "
            "APROVADO"
        )


        print()
        print(
            "8. Limpando staging "
            "da carga..."
        )


        limpar_staging(
            cliente,
            carga_id,
        )


        print(
            "   Staging da carga limpa."
        )


    except Exception:

        if not promovida:

            try:

                limpar_staging(
                    cliente,
                    carga_id,
                )

            except Exception:

                pass

        else:

            print()
            print(
                "ATENCAO: a RPC foi "
                "executada antes da falha."
            )

            print(
                "A staging desta carga "
                "foi preservada para "
                "diagnostico:"
            )

            print(
                carga_id
            )


        raise


    print()
    print("=" * 78)
    print(
        "PUBLICACAO DE PLANOS DE ACAO "
        "CONCLUIDA COM SUCESSO"
    )
    print("=" * 78)

    print(
        "Tabela:",
        f"public.{TABELA_MART}",
    )

    print(
        "Registros:",
        indicadores[
            "total"
        ],
    )


# =========================================================
# CLI
# =========================================================

def main():

    parser = argparse.ArgumentParser()


    parser.add_argument(
        "--dry-run",
        action="store_true",
        help=(
            "Valida a Mart local e "
            "a infraestrutura remota "
            "sem publicar dados."
        ),
    )


    args = (
        parser.parse_args()
    )


    publicar(
        dry_run=args.dry_run,
    )


if __name__ == "__main__":

    main()
