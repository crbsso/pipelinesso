from pathlib import Path
import os
import re
import sys
import time
import unicodedata
from collections import Counter

from openpyxl import load_workbook
from playwright.sync_api import (
    sync_playwright,
)

RAIZ = Path(__file__).resolve().parents[2]

if str(RAIZ) not in sys.path:
    sys.path.insert(
        0,
        str(RAIZ),
    )

from backend.planos_acao import (
    extrair_planos_acao as planos,
)


# =========================================================
# CONFIGURACOES
# =========================================================

PASTA_PROJETO = (
    Path(__file__)
    .resolve()
    .parents[2]
)

URL_PREVSIS = planos.URL_PREVSIS
PROFILE_DIR = planos.PROFILE_DIR

PASTA_BRONZE = (
    PASTA_PROJETO
    / "data"
    / "bronze"
    / "acoes_corretivas"
)

ARQUIVO_FINAL = (
    PASTA_BRONZE
    / "Detalhes_das_acoes_corretivas.xlsx"
)

ARQUIVO_TEMPORARIO = (
    PASTA_BRONZE
    / "Detalhes_das_acoes_corretivas_novo.xlsx"
)


# Contrato observado no arquivo real exportado
# do visual "Detalhes das acoes corretivas".
COLUNAS_OBRIGATORIAS = {
    "id da acao corretiva",
    "data de inicio da acao corretiva",
    "prazo para acao corretiva",
    "modulo de origem",
    "id da resposta",
    "nome do formulario",
    "plano de identificacao",
    "nome do plano",
    "gestor de plano",
    "estado do plano",
    "nivel organizacional",
    "perguntar",
    "descricao acao corretiva",
    "responsavel pelas acoes corretivas",
    "estado: acao corretiva",
    "estado de aprovacao",
    "oficial de aprovacao",
    "descricao final",
    "proxima verificacao",
}


# =========================================================
# UTILITARIOS
# =========================================================

def normalizar(valor):

    if valor is None:
        return ""

    texto = str(valor)

    texto = texto.replace(
        "\xa0",
        " ",
    )

    texto = unicodedata.normalize(
        "NFKD",
        texto,
    )

    texto = "".join(
        caractere
        for caractere in texto
        if not unicodedata.combining(
            caractere
        )
    )

    texto = re.sub(
        r"\s+",
        " ",
        texto,
    )

    return (
        texto
        .strip()
        .lower()
    )


def obter_contextos(page):

    contextos = [
        (
            "PAGINA_PRINCIPAL",
            page,
        )
    ]

    try:

        for indice, frame in enumerate(
            page.frames
        ):

            try:

                frame.locator(
                    "body"
                ).count()

                contextos.append(
                    (
                        f"FRAME_{indice}",
                        frame,
                    )
                )

            except Exception:
                continue

    except Exception:
        pass

    return contextos


# =========================================================
# ABA ACOES CORRETIVAS
# =========================================================

def abrir_acoes_corretivas(page):

    print(
        "Procurando aba "
        "'Acoes corretivas'..."
    )

    padrao = re.compile(
        r"^\s*A..es\s+corretivas\s*$",
        re.IGNORECASE,
    )

    limite = (
        time.time()
        + 90
    )

    while time.time() < limite:

        for (
            nome_contexto,
            contexto,
        ) in obter_contextos(page):

            try:

                candidatos = (
                    contexto.get_by_text(
                        padrao
                    )
                )

                for indice in range(
                    candidatos.count()
                ):

                    candidato = (
                        candidatos.nth(
                            indice
                        )
                    )

                    try:

                        if not candidato.is_visible():
                            continue

                    except Exception:
                        continue

                    try:

                        candidato.scroll_into_view_if_needed(
                            timeout=10000,
                        )

                    except Exception:
                        pass

                    print(
                        "Aba 'Acoes corretivas' "
                        "encontrada."
                    )

                    print(
                        "Contexto:",
                        nome_contexto,
                    )

                    try:

                        candidato.click(
                            timeout=15000,
                        )

                    except Exception:

                        candidato.evaluate(
                            "(el) => el.click()"
                        )

                    print(
                        "Aba 'Acoes corretivas' "
                        "acionada."
                    )

                    time.sleep(
                        6
                    )

                    return

            except Exception:
                continue

        time.sleep(
            1
        )

    raise RuntimeError(
        "A aba 'Acoes corretivas' "
        "nao foi localizada no dashboard "
        "Analise Plano de Acao."
    )


# =========================================================
# VISUAL DETALHES DAS ACOES CORRETIVAS
# =========================================================

def abrir_menu_acoes_corretivas(page):

    padrao_titulo = re.compile(
        r"Detalhes\s+das\s+a..es\s+corretivas",
        re.IGNORECASE,
    )

    # Primeiro tenta o motor generico,
    # que continua sendo o comportamento padrao.
    try:

        return planos.abrir_menu_planos_acao(
            page,
            padrao_titulo=(
                padrao_titulo
            ),
            nome_visual=(
                "Detalhes das acoes corretivas"
            ),
        )

    except RuntimeError as erro:

        mensagem = str(
            erro
        )

        if (
            "botao dos tres pontos nao apareceu"
            not in mensagem
        ):

            raise

        print()
        print(
            "[FALLBACK] O container generico "
            "nao revelou o menu."
        )

        print(
            "[FALLBACK] Procurando o "
            "analysis_visual exato..."
        )


    limite = (
        time.time()
        + 60
    )


    while time.time() < limite:

        for (
            nome_contexto,
            contexto,
        ) in obter_contextos(
            page
        ):

            try:

                titulos = (
                    contexto.get_by_text(
                        padrao_titulo
                    )
                )

            except Exception:
                continue


            for indice in range(
                titulos.count()
            ):

                titulo = (
                    titulos.nth(
                        indice
                    )
                )

                try:

                    if not titulo.is_visible():
                        continue

                except Exception:
                    continue


                try:

                    titulo.scroll_into_view_if_needed(
                        timeout=10000,
                    )

                except Exception:
                    pass


                # O titulo precisa pertencer
                # ao visual do QuickSight.
                visual = titulo.locator(
                    "xpath=ancestor::*"
                    "[@data-automation-id="
                    "'analysis_visual'][1]"
                )


                try:

                    quantidade_visuais = (
                        visual.count()
                    )

                except Exception:

                    quantidade_visuais = 0


                if quantidade_visuais == 0:

                    continue


                visual = visual.first


                try:

                    caixa = (
                        visual.bounding_box()
                    )

                except Exception:

                    caixa = None


                print(
                    "[FALLBACK] analysis_visual "
                    "encontrado."
                )

                print(
                    "[FALLBACK] Contexto:",
                    nome_contexto,
                )


                if caixa:

                    print(
                        "[FALLBACK] Dimensoes: "
                        f"{int(caixa['width'])} x "
                        f"{int(caixa['height'])}"
                    )


                # Hover no proprio visual.
                try:

                    if caixa:

                        visual.hover(
                            position={
                                "x": max(
                                    5,
                                    int(
                                        caixa[
                                            "width"
                                        ]
                                    )
                                    - 20,
                                ),
                                "y": 15,
                            },
                            force=True,
                            timeout=10000,
                        )

                    else:

                        visual.hover(
                            force=True,
                            timeout=10000,
                        )

                except Exception:

                    try:

                        visual.evaluate(
                            """(el) => {
                                el.dispatchEvent(
                                    new MouseEvent(
                                        "mouseover",
                                        {
                                            bubbles: true
                                        }
                                    )
                                );
                            }"""
                        )

                    except Exception:
                        pass


                time.sleep(
                    1
                )


                # Primeiro procura dentro
                # do visual correto.
                botao = visual.locator(
                    '[data-automation-id='
                    '"analysis_visual_dropdown_menu_button"]'
                )


                if botao.count() > 0:

                    for indice_botao in range(
                        botao.count()
                    ):

                        candidato = (
                            botao.nth(
                                indice_botao
                            )
                        )

                        try:

                            if not candidato.is_visible():
                                continue

                        except Exception:
                            continue


                        aria = (
                            candidato.get_attribute(
                                "aria-label"
                            )
                            or
                            ""
                        )

                        print(
                            "[FALLBACK] Botao encontrado."
                        )

                        print(
                            "[FALLBACK] aria-label:",
                            aria,
                        )


                        try:

                            candidato.click(
                                timeout=10000,
                            )

                        except Exception:

                            candidato.evaluate(
                                "(el) => el.click()"
                            )


                        time.sleep(
                            1
                        )


                        # Confirma que o menu
                        # realmente foi aberto.
                        try:

                            exportar = (
                                contexto.get_by_text(
                                    "Export to Excel",
                                    exact=True,
                                )
                            )

                            if (
                                exportar.count() > 0
                            ):

                                print(
                                    "[FALLBACK] Menu de "
                                    "exportacao aberto."
                                )

                                return contexto

                        except Exception:
                            pass


                        # Mesmo que o overlay tenha
                        # sido renderizado em outro
                        # contexto, exportar_excel()
                        # possui busca global depois.
                        print(
                            "[FALLBACK] Clique no menu "
                            "executado."
                        )

                        return contexto


        time.sleep(
            1
        )


    raise RuntimeError(
        "O visual 'Detalhes das acoes "
        "corretivas' foi encontrado, "
        "mas o menu especifico do visual "
        "nao foi localizado."
    )


# =========================================================
# EXPORTACAO
# =========================================================

def exportar_excel(page):

    return planos.exportar_excel(
        page,
        abrir_menu_func=(
            abrir_menu_acoes_corretivas
        ),
        pasta_bronze=(
            PASTA_BRONZE
        ),
        arquivo_temporario=(
            ARQUIVO_TEMPORARIO
        ),
    )


# =========================================================
# VALIDACAO
# =========================================================

def validar_excel_acoes_corretivas(
    arquivo
):

    print(
        "Validando Excel de "
        "Acoes Corretivas..."
    )

    tamanho = (
        planos.validar_estrutura_xlsx(
            arquivo
        )
    )

    workbook = load_workbook(
        arquivo,
        read_only=True,
        data_only=True,
    )

    try:

        if not workbook.sheetnames:

            raise RuntimeError(
                "O Excel nao possui abas."
            )

        worksheet = workbook[
            workbook.sheetnames[0]
        ]

        linha_cabecalho = None
        cabecalho_normalizado = None

        for (
            numero_linha,
            valores,
        ) in enumerate(
            worksheet.iter_rows(
                min_row=1,
                max_row=15,
                values_only=True,
            ),
            start=1,
        ):

            normalizados = [
                normalizar(
                    valor
                )
                for valor in valores
            ]

            if (
                "id da acao corretiva"
                in normalizados
            ):

                linha_cabecalho = (
                    numero_linha
                )

                cabecalho_normalizado = (
                    normalizados
                )

                break

        if linha_cabecalho is None:

            raise RuntimeError(
                "Nao foi encontrado o "
                "cabecalho de Acoes Corretivas."
            )

        cabecalho_set = {
            valor
            for valor
            in cabecalho_normalizado
            if valor
        }

        ausentes = (
            COLUNAS_OBRIGATORIAS
            - cabecalho_set
        )

        if ausentes:

            raise RuntimeError(
                "O Excel nao contem todas "
                "as colunas esperadas. "
                "Ausentes: "
                + ", ".join(
                    sorted(
                        ausentes
                    )
                )
            )

        indice_id = (
            cabecalho_normalizado.index(
                "id da acao corretiva"
            )
        )

        indice_estado = (
            cabecalho_normalizado.index(
                "estado: acao corretiva"
            )
        )

        ids = []
        estados = Counter()
        linhas_dados = 0

        for valores in worksheet.iter_rows(
            min_row=(
                linha_cabecalho
                + 1
            ),
            values_only=True,
        ):

            if not any(
                valor is not None
                and str(valor).strip()
                for valor in valores
            ):

                continue

            linhas_dados += 1

            valor_id = (
                valores[indice_id]
                if indice_id < len(valores)
                else None
            )

            valor_estado = (
                valores[indice_estado]
                if indice_estado < len(valores)
                else None
            )

            id_texto = (
                str(valor_id).strip()
                if valor_id is not None
                else ""
            )

            estado_texto = (
                str(valor_estado).strip()
                if valor_estado is not None
                else ""
            )

            if id_texto:

                ids.append(
                    id_texto
                )

            if estado_texto:

                estados[
                    estado_texto
                ] += 1

    finally:

        workbook.close()

    if linhas_dados == 0:

        raise RuntimeError(
            "Nenhuma Acao Corretiva "
            "foi encontrada no arquivo."
        )

    if not ids:

        raise RuntimeError(
            "Nenhum ID de Acao Corretiva "
            "foi encontrado."
        )

    total_ids = len(
        ids
    )

    ids_unicos = len(
        set(
            ids
        )
    )

    print()
    print(
        "Excel de Acoes Corretivas valido."
    )

    print(
        "Linha do cabecalho:",
        linha_cabecalho,
    )

    print(
        "Colunas validadas:",
        len(
            COLUNAS_OBRIGATORIAS
        ),
    )

    print(
        "Linhas de dados:",
        linhas_dados,
    )

    print(
        "IDs identificados:",
        total_ids,
    )

    print(
        "IDs unicos:",
        ids_unicos,
    )

    print(
        "Estados encontrados:"
    )

    for (
        estado,
        quantidade,
    ) in sorted(
        estados.items()
    ):

        print(
            f"  {estado}: "
            f"{quantidade}"
        )

    if (
        total_ids
        != ids_unicos
    ):

        print()
        print(
            "[AVISO] Existem IDs de "
            "Acoes Corretivas duplicados "
            "na Bronze."
        )

        print(
            "[AVISO] A Bronze sera "
            "preservada sem deduplicacao."
        )

    print(
        f"Tamanho: {tamanho:,} bytes"
    )


# =========================================================
# PROMOCAO PARA BRONZE
# =========================================================

def atualizar_bronze(
    arquivo_baixado
):

    if not arquivo_baixado.exists():

        raise RuntimeError(
            "Arquivo temporario "
            "nao encontrado."
        )

    os.replace(
        arquivo_baixado,
        ARQUIVO_FINAL,
    )

    print()
    print(
        "Camada Bronze de Acoes "
        "Corretivas atualizada:"
    )

    print(
        ARQUIVO_FINAL
    )


def limpar_temporario():

    try:

        if ARQUIVO_TEMPORARIO.exists():

            ARQUIVO_TEMPORARIO.unlink()

    except Exception:
        pass


# =========================================================
# EXECUCAO INDEPENDENTE
# =========================================================

def extrair_acoes_corretivas():

    print(
        "=" * 70
    )

    print(
        "EXTRACAO ACOES CORRETIVAS - PREVSYS"
    )

    print(
        "=" * 70
    )

    PASTA_BRONZE.mkdir(
        parents=True,
        exist_ok=True,
    )

    limpar_temporario()

    with sync_playwright() as p:

        print()
        print(
            "Abrindo Microsoft Edge "
            "com perfil persistente..."
        )

        context = (
            p.chromium.launch_persistent_context(
                user_data_dir=str(
                    PROFILE_DIR
                ),
                channel="msedge",
                headless=False,
                accept_downloads=True,
                viewport=None,
                chromium_sandbox=True,
                args=[
                    "--start-maximized",
                ],
            )
        )

        try:

            page = (
                context.pages[0]
                if context.pages
                else context.new_page()
            )

            planos.garantir_login(
                page
            )

            planos.abrir_reporte_360(
                page
            )

            planos.abrir_plano_acao(
                page
            )

            abrir_acoes_corretivas(
                page
            )

            arquivo = (
                exportar_excel(
                    page
                )
            )

            validar_excel_acoes_corretivas(
                arquivo
            )

            atualizar_bronze(
                arquivo
            )

        finally:

            try:
                context.close()

            except Exception:
                pass

    print()
    print(
        "=" * 70
    )

    print(
        "EXTRACAO DE ACOES CORRETIVAS "
        "CONCLUIDA COM SUCESSO"
    )

    print(
        "=" * 70
    )


if __name__ == "__main__":
    extrair_acoes_corretivas()
