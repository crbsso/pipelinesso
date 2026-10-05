from pathlib import Path
from collections import Counter
from datetime import date, datetime
import re
import unicodedata

from openpyxl import (
    Workbook,
    load_workbook,
)


# =========================================================
# CONFIGURACAO
# =========================================================

RAIZ = (
    Path(__file__)
    .resolve()
    .parents[2]
)

ARQUIVO_PRATA = (
    RAIZ
    / "data"
    / "prata"
    / "acoes_corretivas"
    / "Acoes_corretivas.xlsx"
)

PASTA_OURO = (
    RAIZ
    / "data"
    / "ouro"
    / "acoes_corretivas"
)

ARQUIVO_OURO = (
    PASTA_OURO
    / "mart_acoes_corretivas.xlsx"
)

NOME_ABA = (
    "mart_acoes_corretivas"
)


COLUNAS_OURO = [
    "acao_corretiva_id",
    "plano_id",
    "id_resposta",
    "nome_formulario",
    "modulo_origem",
    "data_inicio",
    "prazo",
    "proxima_verificacao",
    "nivel_organizacional",

    "grupo_sso",
    "metodo_classificacao",

    "gestor_plano",
    "responsavel_acao",
    "responsavel_exibicao",
    "metodo_identidade_responsavel",

    "estado_plano",
    "estado_acao_fonte",

    "flag_ativo_fonte",
    "flag_aberto_fonte",
    "flag_atrasado_fonte",
    "flag_fechado_fonte",
    "flag_fechado_apos_prazo_fonte",

    "data_referencia",
    "dias_para_prazo",
    "dias_em_atraso_calculado",
    "situacao_prazo_calculada",
    "situacao_gerencial",

    "flag_vencido_calculado",
    "flag_em_prazo_calculado",
    "flag_vence_hoje",

    "divergencia_status_prazo",

    "estado_aprovacao",
    "oficial_aprovacao",

    "nome_plano",
    "pergunta",
    "descricao_acao",
    "descricao_final",

    "registros_origem",
    "duplicado_origem",
]


# =========================================================
# NORMALIZACAO
# =========================================================

def normalizar(
    valor,
):

    if valor is None:
        return ""

    texto = unicodedata.normalize(
        "NFKD",
        str(valor),
    )

    texto = "".join(
        caractere
        for caractere in texto
        if not unicodedata.combining(
            caractere
        )
    )

    texto = re.sub(
        r"\s+",
        " ",
        texto,
    )

    return (
        texto
        .strip()
        .lower()
    )


def limpar(
    valor,
):

    if valor is None:
        return None

    if isinstance(
        valor,
        str,
    ):

        valor = valor.strip()

        if not valor:
            return None

    return valor


def obter_data(
    valor,
):

    if valor is None:
        return None

    if isinstance(
        valor,
        datetime,
    ):

        return valor.date()

    if isinstance(
        valor,
        date,
    ):

        return valor

    texto = str(
        valor
    ).strip()

    if not texto:
        return None

    formatos = [
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
        "%d/%m/%Y",
    ]

    for formato in formatos:

        try:

            return datetime.strptime(
                texto,
                formato,
            ).date()

        except ValueError:
            pass

    raise RuntimeError(
        "Data nao reconhecida: "
        + texto
    )


# =========================================================
# LEITURA DA PRATA
# =========================================================

def carregar_prata():

    if not ARQUIVO_PRATA.exists():

        raise RuntimeError(
            "Arquivo da Prata nao encontrado: "
            + str(
                ARQUIVO_PRATA
            )
        )

    workbook = load_workbook(
        ARQUIVO_PRATA,
        read_only=True,
        data_only=True,
    )

    try:

        worksheet = workbook[
            workbook.sheetnames[0]
        ]

        cabecalho = list(
            next(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=1,
                    values_only=True,
                )
            )
        )

        registros = []

        for valores in worksheet.iter_rows(
            min_row=2,
            values_only=True,
        ):

            if not any(
                valor is not None
                and str(
                    valor
                ).strip()
                for valor in valores
            ):

                continue

            registros.append(
                {
                    cabecalho[indice]: (
                        valores[indice]
                        if indice
                        < len(
                            valores
                        )
                        else None
                    )
                    for indice in range(
                        len(
                            cabecalho
                        )
                    )
                }
            )

        return registros

    finally:

        workbook.close()


# =========================================================
# TRANSFORMACAO OURO
# =========================================================

def transformar(
    registro,
    data_referencia,
):

    estado_fonte = normalizar(
        registro.get(
            "estado_acao"
        )
    )

    estados_ativos = {
        "aberto",
        "atrasado",
    }

    estados_fechados = {
        "fechado",
        "fechado apos o prazo",
    }

    if (
        estado_fonte
        not in estados_ativos
        and estado_fonte
        not in estados_fechados
    ):

        raise RuntimeError(
            "Estado de Acao Corretiva "
            "nao previsto: "
            + estado_fonte
        )

    ativo_fonte = (
        estado_fonte
        in estados_ativos
    )

    prazo = obter_data(
        registro.get(
            "prazo"
        )
    )

    data_inicio = obter_data(
        registro.get(
            "data_inicio"
        )
    )

    proxima_verificacao = obter_data(
        registro.get(
            "proxima_verificacao"
        )
    )


    dias_para_prazo = None

    dias_em_atraso = None

    situacao_prazo = (
        "NAO_APLICAVEL"
    )

    situacao_gerencial = None

    flag_vencido = 0

    flag_em_prazo = 0

    flag_vence_hoje = 0

    divergencia = 0


    # =====================================================
    # ACOES ATIVAS
    # =====================================================

    if ativo_fonte:

        if prazo is None:

            situacao_prazo = (
                "SEM_PRAZO"
            )

            situacao_gerencial = (
                "SEM_PRAZO"
            )

        elif prazo < data_referencia:

            dias_em_atraso = (
                data_referencia
                - prazo
            ).days

            situacao_prazo = (
                "VENCIDO"
            )

            situacao_gerencial = (
                "VENCIDO"
            )

            flag_vencido = 1

        elif prazo == data_referencia:

            dias_para_prazo = 0

            situacao_prazo = (
                "VENCE_HOJE"
            )

            situacao_gerencial = (
                "VENCE_HOJE"
            )

            flag_em_prazo = 1

            flag_vence_hoje = 1

        else:

            dias_para_prazo = (
                prazo
                - data_referencia
            ).days

            situacao_prazo = (
                "EM_PRAZO"
            )

            situacao_gerencial = (
                "EM_PRAZO"
            )

            flag_em_prazo = 1


        if (
            estado_fonte == "aberto"
            and situacao_prazo
            == "VENCIDO"
        ):

            divergencia = 1


        if (
            estado_fonte == "atrasado"
            and situacao_prazo
            in {
                "EM_PRAZO",
                "VENCE_HOJE",
            }
        ):

            divergencia = 1


    # =====================================================
    # ACOES FECHADAS
    #
    # Nao existe data de encerramento da Acao Corretiva
    # na fonte atual. Portanto:
    #
    # - prazo calculado = NAO_APLICAVEL
    # - classificacao de fechamento vem do PREVSYS
    # =====================================================

    elif (
        estado_fonte
        == "fechado"
    ):

        situacao_gerencial = (
            "FECHADO"
        )

    elif (
        estado_fonte
        == "fechado apos o prazo"
    ):

        situacao_gerencial = (
            "FECHADO_APOS_PRAZO"
        )


    return {
        "acao_corretiva_id":
            limpar(
                registro.get(
                    "acao_corretiva_id"
                )
            ),

        "plano_id":
            limpar(
                registro.get(
                    "plano_id"
                )
            ),

        "id_resposta":
            limpar(
                registro.get(
                    "id_resposta"
                )
            ),

        "nome_formulario":
            limpar(
                registro.get(
                    "nome_formulario"
                )
            ),

        "modulo_origem":
            limpar(
                registro.get(
                    "modulo_origem"
                )
            ),

        "data_inicio":
            data_inicio,

        "prazo":
            prazo,

        "proxima_verificacao":
            proxima_verificacao,

        "nivel_organizacional":
            limpar(
                registro.get(
                    "nivel_organizacional"
                )
            ),

        "grupo_sso":
            limpar(
                registro.get(
                    "grupo_sso"
                )
            ),

        "metodo_classificacao":
            limpar(
                registro.get(
                    "metodo_classificacao"
                )
            ),

        "gestor_plano":
            limpar(
                registro.get(
                    "gestor_plano"
                )
            ),

        "responsavel_acao":
            limpar(
                registro.get(
                    "responsavel_acao"
                )
            ),

        "responsavel_exibicao":
            limpar(
                registro.get(
                    "responsavel_exibicao"
                )
            ),

        "metodo_identidade_responsavel":
            limpar(
                registro.get(
                    "metodo_identidade_responsavel"
                )
            ),

        "estado_plano":
            limpar(
                registro.get(
                    "estado_plano"
                )
            ),

        "estado_acao_fonte":
            limpar(
                registro.get(
                    "estado_acao"
                )
            ),

        "flag_ativo_fonte":
            1
            if ativo_fonte
            else 0,

        "flag_aberto_fonte":
            1
            if estado_fonte
            == "aberto"
            else 0,

        "flag_atrasado_fonte":
            1
            if estado_fonte
            == "atrasado"
            else 0,

        "flag_fechado_fonte":
            1
            if estado_fonte
            == "fechado"
            else 0,

        "flag_fechado_apos_prazo_fonte":
            1
            if estado_fonte
            == "fechado apos o prazo"
            else 0,

        "data_referencia":
            data_referencia,

        "dias_para_prazo":
            dias_para_prazo,

        "dias_em_atraso_calculado":
            dias_em_atraso,

        "situacao_prazo_calculada":
            situacao_prazo,

        "situacao_gerencial":
            situacao_gerencial,

        "flag_vencido_calculado":
            flag_vencido,

        "flag_em_prazo_calculado":
            flag_em_prazo,

        "flag_vence_hoje":
            flag_vence_hoje,

        "divergencia_status_prazo":
            divergencia,

        "estado_aprovacao":
            limpar(
                registro.get(
                    "estado_aprovacao"
                )
            ),

        "oficial_aprovacao":
            limpar(
                registro.get(
                    "oficial_aprovacao"
                )
            ),

        "nome_plano":
            limpar(
                registro.get(
                    "nome_plano"
                )
            ),

        "pergunta":
            limpar(
                registro.get(
                    "pergunta"
                )
            ),

        "descricao_acao":
            limpar(
                registro.get(
                    "descricao_acao"
                )
            ),

        "descricao_final":
            limpar(
                registro.get(
                    "descricao_final"
                )
            ),

        "registros_origem":
            int(
                registro.get(
                    "registros_origem"
                )
                or 0
            ),

        "duplicado_origem":
            int(
                registro.get(
                    "duplicado_origem"
                )
                or 0
            ),
    }


# =========================================================
# QUALITY GATE
# =========================================================

def validar_ouro(
    registros,
):

    if not registros:

        raise RuntimeError(
            "Ouro vazio."
        )

    ids = [
        str(
            registro.get(
                "acao_corretiva_id"
            )
            or
            ""
        ).strip()
        for registro in registros
    ]

    if any(
        not acao_id
        for acao_id in ids
    ):

        raise RuntimeError(
            "Existe Acao Corretiva sem ID."
        )

    if len(
        ids
    ) != len(
        set(
            ids
        )
    ):

        raise RuntimeError(
            "Existem IDs duplicados no Ouro."
        )

    if len(
        registros
    ) != 613:

        raise RuntimeError(
            "Quantidade inesperada "
            "de registros no Ouro: "
            + str(
                len(
                    registros
                )
            )
        )


    divergencias = sum(
        int(
            registro.get(
                "divergencia_status_prazo"
            )
            or 0
        )
        for registro in registros
    )

    if divergencias != 0:

        raise RuntimeError(
            "Existem divergencias entre "
            "estado da fonte e prazo calculado: "
            + str(
                divergencias
            )
        )


    sem_grupo = [
        registro[
            "acao_corretiva_id"
        ]
        for registro in registros
        if registro.get(
            "grupo_sso"
        )
        not in {
            "SEGURANCA",
            "OUTRAS_AREAS",
        }
    ]

    if sem_grupo:

        raise RuntimeError(
            "Acoes sem grupo_sso valido: "
            + str(
                sem_grupo[:20]
            )
        )


# =========================================================
# GRAVACAO
# =========================================================

def salvar_ouro(
    registros,
):

    PASTA_OURO.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = (
        NOME_ABA
    )

    worksheet.append(
        COLUNAS_OURO
    )

    for registro in registros:

        worksheet.append(
            [
                registro.get(
                    coluna
                )
                for coluna
                in COLUNAS_OURO
            ]
        )

    workbook.save(
        ARQUIVO_OURO
    )


# =========================================================
# EXECUCAO
# =========================================================

def main():

    print(
        "=" * 70
    )

    print(
        "GERACAO OURO - ACOES CORRETIVAS"
    )

    print(
        "=" * 70
    )

    prata = carregar_prata()

    print()
    print(
        "Registros recebidos:",
        len(
            prata
        ),
    )

    data_referencia = (
        date.today()
    )

    ouro = [
        transformar(
            registro,
            data_referencia,
        )
        for registro
        in prata
    ]

    validar_ouro(
        ouro
    )


    grupos = Counter(
        registro[
            "grupo_sso"
        ]
        for registro
        in ouro
    )

    estados_fonte = Counter(
        normalizar(
            registro[
                "estado_acao_fonte"
            ]
        )
        for registro
        in ouro
    )

    situacoes = Counter(
        registro[
            "situacao_gerencial"
        ]
        for registro
        in ouro
    )


    ativos = sum(
        registro[
            "flag_ativo_fonte"
        ]
        for registro
        in ouro
    )

    fechados = (
        len(
            ouro
        )
        - ativos
    )

    divergencias = sum(
        registro[
            "divergencia_status_prazo"
        ]
        for registro
        in ouro
    )


    salvar_ouro(
        ouro
    )


    print()
    print(
        "Registros no Ouro:",
        len(
            ouro
        ),
    )

    print()
    print(
        "Grupos:"
    )

    for grupo, quantidade in sorted(
        grupos.items()
    ):

        print(
            f"  {grupo}: "
            f"{quantidade}"
        )


    print()
    print(
        "Estados da fonte:"
    )

    for estado, quantidade in sorted(
        estados_fonte.items()
    ):

        print(
            f"  {estado}: "
            f"{quantidade}"
        )


    print()
    print(
        "Populacao:"
    )

    print(
        "  Ativas:",
        ativos,
    )

    print(
        "  Fechadas:",
        fechados,
    )


    print()
    print(
        "Situacao gerencial:"
    )

    for situacao, quantidade in sorted(
        situacoes.items()
    ):

        print(
            f"  {situacao}: "
            f"{quantidade}"
        )


    print()
    print(
        "Divergencias fonte x prazo:",
        divergencias,
    )


    print()
    print(
        "QUALITY GATE OURO: APROVADO"
    )

    print()
    print(
        "Ouro gerado:"
    )

    print(
        ARQUIVO_OURO
    )

    print()
    print(
        "=" * 70
    )

    print(
        "GERACAO OURO DE ACOES CORRETIVAS "
        "CONCLUIDA COM SUCESSO"
    )

    print(
        "=" * 70
    )


if __name__ == "__main__":
    main()
