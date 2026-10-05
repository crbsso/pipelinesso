import argparse
import os
import uuid
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import load_workbook
from supabase import create_client


# =========================================================
# CONFIGURACAO
# =========================================================

RAIZ = (
    Path(__file__)
    .resolve()
    .parents[2]
)

ARQUIVO_OURO = (
    RAIZ
    / "data"
    / "ouro"
    / "acoes_corretivas"
    / "mart_acoes_corretivas.xlsx"
)

NOME_ABA = (
    "mart_acoes_corretivas"
)

TABELA_MART = (
    "mart_acoes_corretivas"
)

TABELA_STAGING = (
    "stg_mart_acoes_corretivas"
)

RPC_PROMOVER = (
    "promover_mart_acoes_corretivas"
)

PROJECT_REF_ESPERADO = (
    "lkhidajhsqgebjltyicb"
)

TAMANHO_LOTE = 500


COLUNAS_ESPERADAS = [
    "acao_corretiva_id",
    "plano_id",
    "id_resposta",
    "nome_formulario",
    "modulo_origem",
    "data_inicio",
    "prazo",
    "proxima_verificacao",
    "nivel_organizacional",
    "grupo_sso",
    "metodo_classificacao",
    "gestor_plano",
    "responsavel_acao",
    "responsavel_exibicao",
    "metodo_identidade_responsavel",
    "estado_plano",
    "estado_acao_fonte",
    "flag_ativo_fonte",
    "flag_aberto_fonte",
    "flag_atrasado_fonte",
    "flag_fechado_fonte",
    "flag_fechado_apos_prazo_fonte",
    "data_referencia",
    "dias_para_prazo",
    "dias_em_atraso_calculado",
    "situacao_prazo_calculada",
    "situacao_gerencial",
    "flag_vencido_calculado",
    "flag_em_prazo_calculado",
    "flag_vence_hoje",
    "divergencia_status_prazo",
    "estado_aprovacao",
    "oficial_aprovacao",
    "nome_plano",
    "pergunta",
    "descricao_acao",
    "descricao_final",
    "registros_origem",
    "duplicado_origem",
]


COLUNAS_DATA = {
    "data_inicio",
    "prazo",
    "proxima_verificacao",
    "data_referencia",
}


COLUNAS_INTEIRO = {
    "id_resposta",
    "flag_ativo_fonte",
    "flag_aberto_fonte",
    "flag_atrasado_fonte",
    "flag_fechado_fonte",
    "flag_fechado_apos_prazo_fonte",
    "dias_para_prazo",
    "dias_em_atraso_calculado",
    "flag_vencido_calculado",
    "flag_em_prazo_calculado",
    "flag_vence_hoje",
    "divergencia_status_prazo",
    "registros_origem",
    "duplicado_origem",
}


COLUNAS_ID_TEXTO = {
    "acao_corretiva_id",
    "plano_id",
}



# =========================================================
# IDENTIDADE DO RESPONSAVEL
# =========================================================

def validar_identidade_publicacao(
    registros,
):

    esperados = {
        "MAPA_RM_EXATO": 466,
        "MAPA_RM_ALIAS_COMPARTILHADO": 127,
        "MAPA_RM_ALIAS_ACAO": 19,
        "PREVSYS_SEM_RM": 1,
    }


    contadores = {
        chave: 0
        for chave
        in esperados
    }


    sem_exibicao = []
    sem_metodo = []
    metodo_invalido = []


    for registro in registros:

        acao_id = str(
            registro.get(
                "acao_corretiva_id"
            )
            or ""
        ).strip()


        exibicao = str(
            registro.get(
                "responsavel_exibicao"
            )
            or ""
        ).strip()


        metodo = str(
            registro.get(
                "metodo_identidade_responsavel"
            )
            or ""
        ).strip()


        if not exibicao:

            sem_exibicao.append(
                acao_id
            )


        if not metodo:

            sem_metodo.append(
                acao_id
            )

            continue


        if metodo not in esperados:

            metodo_invalido.append(
                (
                    acao_id,
                    metodo,
                )
            )

            continue


        contadores[
            metodo
        ] += 1


    if sem_exibicao:

        raise RuntimeError(
            "Acoes sem responsavel_exibicao: "
            + str(
                sem_exibicao[:20]
            )
        )


    if sem_metodo:

        raise RuntimeError(
            "Acoes sem metodo de identidade: "
            + str(
                sem_metodo[:20]
            )
        )


    if metodo_invalido:

        raise RuntimeError(
            "Metodo de identidade invalido: "
            + str(
                metodo_invalido[:20]
            )
        )


    if contadores != esperados:

        raise RuntimeError(
            "Distribuicao da identidade "
            "divergente. Obtido: "
            + str(
                contadores
            )
            + " | Esperado: "
            + str(
                esperados
            )
        )


    return contadores


# =========================================================
# CONFIGURACAO SUPABASE
# =========================================================

def carregar_configuracao():

    load_dotenv(
        RAIZ / ".env"
    )

    supabase_url = os.getenv(
        "SUPABASE_URL"
    )

    supabase_secret_key = os.getenv(
        "SUPABASE_SECRET_KEY"
    )

    if not supabase_url:

        raise RuntimeError(
            "SUPABASE_URL nao encontrada "
            "no arquivo .env."
        )

    if not supabase_secret_key:

        raise RuntimeError(
            "SUPABASE_SECRET_KEY nao encontrada "
            "no arquivo .env."
        )

    if (
        PROJECT_REF_ESPERADO
        not in supabase_url
    ):

        raise RuntimeError(
            "BLOQUEIO DE SEGURANCA: "
            "SUPABASE_URL nao corresponde "
            "ao projeto Portal SSO. "
            f"Esperado ref: {PROJECT_REF_ESPERADO}. "
            f"Recebido: {supabase_url}"
        )

    return (
        supabase_url,
        supabase_secret_key,
    )


# =========================================================
# SERIALIZACAO
# =========================================================

def serializar_valor(
    coluna,
    valor,
):

    if valor is None:
        return None

    if (
        isinstance(
            valor,
            str,
        )
        and not valor.strip()
    ):
        return None

    if coluna in COLUNAS_DATA:

        if isinstance(
            valor,
            datetime,
        ):
            return (
                valor
                .date()
                .isoformat()
            )

        if isinstance(
            valor,
            date,
        ):
            return valor.isoformat()

        texto = str(
            valor
        ).strip()

        if not texto:
            return None

        return texto


    if coluna in COLUNAS_INTEIRO:

        if valor is None:
            return None

        return int(
            valor
        )


    if coluna in COLUNAS_ID_TEXTO:

        texto = str(
            valor
        ).strip()

        if texto.endswith(
            ".0"
        ):

            try:

                texto = str(
                    int(
                        float(
                            texto
                        )
                    )
                )

            except ValueError:
                pass

        return texto


    if isinstance(
        valor,
        str,
    ):

        valor = valor.strip()

        if not valor:
            return None

    return valor


# =========================================================
# LEITURA DO OURO
# =========================================================

def carregar_ouro():

    if not ARQUIVO_OURO.exists():

        raise RuntimeError(
            "Arquivo Ouro nao encontrado: "
            + str(
                ARQUIVO_OURO
            )
        )

    workbook = load_workbook(
        ARQUIVO_OURO,
        read_only=True,
        data_only=True,
    )

    try:

        if NOME_ABA not in workbook.sheetnames:

            raise RuntimeError(
                "Aba esperada nao encontrada: "
                + NOME_ABA
            )

        worksheet = workbook[
            NOME_ABA
        ]

        cabecalho = list(
            next(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=1,
                    values_only=True,
                )
            )
        )

        if (
            cabecalho
            != COLUNAS_ESPERADAS
        ):

            raise RuntimeError(
                "Contrato do Ouro divergente.\n"
                f"Recebido: {cabecalho}"
            )

        registros = []

        for valores in worksheet.iter_rows(
            min_row=2,
            values_only=True,
        ):

            if not any(
                valor is not None
                and str(
                    valor
                ).strip()
                for valor in valores
            ):
                continue

            registro = {}

            for indice, coluna in enumerate(
                cabecalho
            ):

                valor = (
                    valores[
                        indice
                    ]
                    if indice
                    < len(
                        valores
                    )
                    else None
                )

                registro[
                    coluna
                ] = serializar_valor(
                    coluna,
                    valor,
                )

            registros.append(
                registro
            )

        contadores_identidade_publicacao = (
            validar_identidade_publicacao(
                registros
            )
        )

        return registros

    finally:

        workbook.close()


# =========================================================
# QUALITY GATE LOCAL
# =========================================================

def validar_local(
    registros,
):

    if len(
        registros
    ) != 613:

        raise RuntimeError(
            "Quantidade local inesperada: "
            + str(
                len(
                    registros
                )
            )
        )


    ids = [
        registro.get(
            "acao_corretiva_id"
        )
        for registro
        in registros
    ]

    if any(
        not acao_id
        for acao_id in ids
    ):

        raise RuntimeError(
            "Existe Acao Corretiva sem ID."
        )

    if len(
        ids
    ) != len(
        set(
            ids
        )
    ):

        raise RuntimeError(
            "Existem IDs duplicados "
            "no arquivo Ouro."
        )


    grupos = Counter(
        registro.get(
            "grupo_sso"
        )
        for registro
        in registros
    )

    esperado_grupos = {
        "OUTRAS_AREAS": 388,
        "SEGURANCA": 225,
    }

    if dict(
        grupos
    ) != esperado_grupos:

        raise RuntimeError(
            "Distribuicao grupo_sso divergente: "
            + str(
                dict(
                    grupos
                )
            )
        )


    situacoes = Counter(
        registro.get(
            "situacao_gerencial"
        )
        for registro
        in registros
    )

    esperado_situacoes = {
        "EM_PRAZO": 102,
        "VENCIDO": 126,
        "FECHADO": 295,
        "FECHADO_APOS_PRAZO": 90,
    }

    if dict(
        situacoes
    ) != esperado_situacoes:

        raise RuntimeError(
            "Situacao gerencial divergente: "
            + str(
                dict(
                    situacoes
                )
            )
        )


    divergencias = sum(
        int(
            registro.get(
                "divergencia_status_prazo"
            )
            or 0
        )
        for registro
        in registros
    )

    if divergencias != 0:

        raise RuntimeError(
            "Existem divergencias "
            "fonte x prazo: "
            + str(
                divergencias
            )
        )


    acao_584 = [
        registro
        for registro
        in registros
        if registro.get(
            "acao_corretiva_id"
        )
        == "584"
    ]

    if len(
        acao_584
    ) != 1:

        raise RuntimeError(
            "Acao 584 deveria existir "
            "uma unica vez."
        )

    if int(
        acao_584[
            0
        ].get(
            "registros_origem"
        )
        or 0
    ) != 2:

        raise RuntimeError(
            "Acao 584 perdeu "
            "registros_origem = 2."
        )


# =========================================================
# SUPABASE - CONTAGEM
# =========================================================

def contar_remoto(
    cliente,
    tabela,
    carga_id=None,
):

    consulta = (
        cliente
        .table(
            tabela
        )
        .select(
            "acao_corretiva_id",
            count="exact",
        )
    )

    if carga_id is not None:

        consulta = (
            consulta
            .eq(
                "carga_id",
                carga_id,
            )
        )

    resposta = (
        consulta
        .limit(
            1
        )
        .execute()
    )

    quantidade = getattr(
        resposta,
        "count",
        None,
    )

    if quantidade is None:

        raise RuntimeError(
            "Supabase nao retornou "
            f"contagem para {tabela}."
        )

    return int(
        quantidade
    )


# =========================================================
# SUPABASE - DISTRIBUICOES
# =========================================================

def carregar_resumo_remoto(
    cliente,
    tabela,
    carga_id=None,
):

    consulta = (
        cliente
        .table(
            tabela
        )
        .select(
            ",".join(
                [
                    "acao_corretiva_id",
                    "grupo_sso",
                    "situacao_gerencial",
                    "divergencia_status_prazo",
                    "registros_origem",
                    "duplicado_origem",
                ]
            )
        )
    )

    if carga_id is not None:

        consulta = (
            consulta
            .eq(
                "carga_id",
                carga_id,
            )
        )

    resposta = (
        consulta
        .range(
            0,
            999
        )
        .execute()
    )

    dados = (
        resposta.data
        or []
    )

    return dados


def validar_resumo_remoto(
    registros,
    contexto,
):

    if len(
        registros
    ) != 613:

        raise RuntimeError(
            f"{contexto}: quantidade divergente: "
            + str(
                len(
                    registros
                )
            )
        )


    ids = [
        str(
            registro.get(
                "acao_corretiva_id"
            )
            or
            ""
        ).strip()
        for registro
        in registros
    ]

    if len(
        set(
            ids
        )
    ) != 613:

        raise RuntimeError(
            f"{contexto}: IDs nao sao unicos."
        )


    grupos = Counter(
        registro.get(
            "grupo_sso"
        )
        for registro
        in registros
    )

    if dict(
        grupos
    ) != {
        "OUTRAS_AREAS": 388,
        "SEGURANCA": 225,
    }:

        raise RuntimeError(
            f"{contexto}: grupo_sso divergente: "
            + str(
                dict(
                    grupos
                )
            )
        )


    situacoes = Counter(
        registro.get(
            "situacao_gerencial"
        )
        for registro
        in registros
    )

    if dict(
        situacoes
    ) != {
        "EM_PRAZO": 102,
        "VENCIDO": 126,
        "FECHADO": 295,
        "FECHADO_APOS_PRAZO": 90,
    }:

        raise RuntimeError(
            f"{contexto}: situacao gerencial "
            "divergente: "
            + str(
                dict(
                    situacoes
                )
            )
        )


    divergencias = sum(
        int(
            registro.get(
                "divergencia_status_prazo"
            )
            or 0
        )
        for registro
        in registros
    )

    if divergencias != 0:

        raise RuntimeError(
            f"{contexto}: existem divergencias "
            "fonte x prazo: "
            + str(
                divergencias
            )
        )


    acao_584 = [
        registro
        for registro
        in registros
        if str(
            registro.get(
                "acao_corretiva_id"
            )
            or
            ""
        ).strip()
        == "584"
    ]

    if len(
        acao_584
    ) != 1:

        raise RuntimeError(
            f"{contexto}: acao 584 "
            "nao esta unica."
        )

    if int(
        acao_584[
            0
        ].get(
            "registros_origem"
        )
        or 0
    ) != 2:

        raise RuntimeError(
            f"{contexto}: acao 584 perdeu "
            "registros_origem = 2."
        )


# =========================================================
# STAGING
# =========================================================

def preparar_staging(
    registros,
    carga_id,
):

    preparados = []

    for registro in registros:

        novo = dict(
            registro
        )

        novo[
            "carga_id"
        ] = carga_id

        preparados.append(
            novo
        )

    return preparados


def inserir_staging(
    cliente,
    registros,
):

    for inicio in range(
        0,
        len(
            registros
        ),
        TAMANHO_LOTE,
    ):

        lote = registros[
            inicio:
            inicio
            + TAMANHO_LOTE
        ]

        (
            cliente
            .table(
                TABELA_STAGING
            )
            .insert(
                lote
            )
            .execute()
        )


def limpar_staging(
    cliente,
    carga_id,
):

    (
        cliente
        .table(
            TABELA_STAGING
        )
        .delete()
        .eq(
            "carga_id",
            carga_id,
        )
        .execute()
    )

    restante = contar_remoto(
        cliente,
        TABELA_STAGING,
        carga_id=carga_id,
    )

    if restante != 0:

        raise RuntimeError(
            "Falha ao limpar staging. "
            f"Restaram {restante} registros."
        )


# =========================================================
# PROMOCAO
# =========================================================

def promover(
    cliente,
    carga_id,
    quantidade,
):

    resposta = (
        cliente
        .rpc(
            RPC_PROMOVER,
            {
                "p_carga_id":
                    carga_id,

                "p_quantidade_esperada":
                    quantidade,
            },
        )
        .execute()
    )

    retorno = resposta.data

    if retorno is not None:

        try:

            retorno_int = int(
                retorno
            )

        except (
            TypeError,
            ValueError,
        ):

            retorno_int = None

        if (
            retorno_int is not None
            and retorno_int
            != quantidade
        ):

            raise RuntimeError(
                "RPC retornou quantidade "
                "divergente. "
                f"Esperado: {quantidade}. "
                f"Retornado: {retorno_int}."
            )

    return retorno


# =========================================================
# EXECUCAO
# =========================================================

def executar(
    dry_run=False,
):

    print(
        "=" * 72
    )

    print(
        "PUBLICACAO OURO - ACOES CORRETIVAS"
    )

    print(
        "=" * 72
    )


    print()
    print(
        "1. Carregando e validando Ouro local..."
    )

    registros = carregar_ouro()

    validar_local(
        registros
    )

    print(
        "   Registros locais:",
        len(
            registros
        ),
    )

    print(
        "   Quality Gate local: APROVADO"
    )


    print()
    print(
        "2. Validando configuracao Supabase..."
    )

    (
        supabase_url,
        supabase_secret_key,
    ) = carregar_configuracao()

    print(
        "   Projeto validado:",
        PROJECT_REF_ESPERADO,
    )


    print()
    print(
        "3. Conectando ao Supabase..."
    )

    cliente = create_client(
        supabase_url,
        supabase_secret_key,
    )

    quantidade_mart = contar_remoto(
        cliente,
        TABELA_MART,
    )

    quantidade_staging = contar_remoto(
        cliente,
        TABELA_STAGING,
    )

    print(
        "   mart_acoes_corretivas:",
        quantidade_mart,
    )

    print(
        "   stg_mart_acoes_corretivas:",
        quantidade_staging,
    )


    if dry_run:

        print()
        print(
            "DRY-RUN: nenhuma escrita foi executada."
        )

        print()
        print(
            "PUBLICADOR VALIDADO PARA EXECUCAO."
        )

        return


    carga_id = str(
        uuid.uuid4()
    )

    houve_promocao = False


    try:

        print()
        print(
            "4. Preparando staging..."
        )

        staging = preparar_staging(
            registros,
            carga_id,
        )

        print(
            "   carga_id:",
            carga_id,
        )


        print()
        print(
            "5. Inserindo staging..."
        )

        inserir_staging(
            cliente,
            staging,
        )

        quantidade_staging_carga = (
            contar_remoto(
                cliente,
                TABELA_STAGING,
                carga_id=carga_id,
            )
        )

        print(
            "   Registros na staging:",
            quantidade_staging_carga,
        )

        if (
            quantidade_staging_carga
            != len(
                registros
            )
        ):

            raise RuntimeError(
                "Quantidade da staging "
                "diverge do Ouro local."
            )


        resumo_staging = (
            carregar_resumo_remoto(
                cliente,
                TABELA_STAGING,
                carga_id=carga_id,
            )
        )

        validar_resumo_remoto(
            resumo_staging,
            "STAGING",
        )

        print(
            "   Quality Gate staging: APROVADO"
        )


        print()
        print(
            "6. Promovendo staging -> MART..."
        )

        retorno = promover(
            cliente,
            carga_id,
            len(
                registros
            ),
        )

        houve_promocao = True

        print(
            "   RPC executada."
        )

        print(
            "   Retorno:",
            retorno,
        )


        print()
        print(
            "7. Validando MART oficial..."
        )

        quantidade_mart = contar_remoto(
            cliente,
            TABELA_MART,
        )

        print(
            "   Registros no MART:",
            quantidade_mart,
        )

        if (
            quantidade_mart
            != len(
                registros
            )
        ):

            raise RuntimeError(
                "Quantidade do MART "
                "diverge do Ouro local."
            )


        resumo_mart = (
            carregar_resumo_remoto(
                cliente,
                TABELA_MART,
            )
        )

        validar_resumo_remoto(
            resumo_mart,
            "MART",
        )

        print(
            "   Quality Gate MART: APROVADO"
        )


        print()
        print(
            "8. Limpando staging da carga..."
        )

        limpar_staging(
            cliente,
            carga_id,
        )

        print(
            "   Staging limpa."
        )


        print()
        print(
            "=" * 72
        )

        print(
            "PUBLICACAO CONCLUIDA COM SUCESSO"
        )

        print(
            "=" * 72
        )

        print()
        print(
            "carga_id:",
            carga_id,
        )

        print(
            "Registros publicados:",
            quantidade_mart,
        )


    except Exception:

        print()
        print(
            "ERRO DURANTE A PUBLICACAO."
        )

        if houve_promocao:

            print(
                "ATENCAO: a RPC foi executada "
                "antes da falha."
            )

            print(
                "O MART pode ter sido promovido; "
                "nao limparei a staging "
                "automaticamente para permitir "
                "diagnostico."
            )

        else:

            print(
                "A RPC ainda nao havia sido "
                "executada."
            )

            print(
                "Tentando remover somente "
                "a carga temporaria da staging..."
            )

            try:

                limpar_staging(
                    cliente,
                    carga_id,
                )

                print(
                    "Staging temporaria limpa."
                )

            except Exception as erro_limpeza:

                print(
                    "Falha adicional ao limpar staging:",
                    erro_limpeza,
                )

        raise


def main():

    parser = argparse.ArgumentParser(
        description=(
            "Publica mart_acoes_corretivas "
            "no Supabase."
        )
    )

    parser.add_argument(
        "--dry-run",
        action="store_true",
        help=(
            "Valida arquivo, configuracao "
            "e conexao sem escrever no banco."
        ),
    )

    argumentos = parser.parse_args()

    executar(
        dry_run=argumentos.dry_run,
    )


if __name__ == "__main__":
    main()
