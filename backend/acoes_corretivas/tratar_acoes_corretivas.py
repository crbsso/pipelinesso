from pathlib import Path
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from datetime import datetime
import re
import unicodedata

from openpyxl import (
    Workbook,
    load_workbook,
)


# =========================================================
# CONFIGURACAO
# =========================================================

RAIZ = (
    Path(__file__)
    .resolve()
    .parents[2]
)

ARQUIVO_BRONZE = (
    RAIZ
    / "data"
    / "bronze"
    / "acoes_corretivas"
    / "Detalhes_das_acoes_corretivas.xlsx"
)

ARQUIVO_PLANOS_PRATA = (
    RAIZ
    / "data"
    / "prata"
    / "planos_acao"
    / "Planos_de_acao.xlsx"
)

PASTA_PRATA = (
    RAIZ
    / "data"
    / "prata"
    / "acoes_corretivas"
)

ARQUIVO_PRATA = (
    PASTA_PRATA
    / "Acoes_corretivas.xlsx"
)

NOME_ABA = (
    "acoes_corretivas"
)


COLUNAS_SAIDA = [
    "acao_corretiva_id",
    "data_inicio",
    "prazo",
    "modulo_origem",
    "id_resposta",
    "nome_formulario",
    "plano_id",
    "nome_plano",
    "gestor_plano",
    "estado_plano",
    "nivel_organizacional",
    "pergunta",
    "descricao_acao",
    "responsavel_acao",
    "estado_acao",
    "estado_aprovacao",
    "oficial_aprovacao",
    "descricao_final",
    "proxima_verificacao",
    "grupo_sso",
    "metodo_classificacao",
    "registros_origem",
    "duplicado_origem",
]


# =========================================================
# NORMALIZACAO
# =========================================================

def normalizar_texto(
    valor,
):

    if valor is None:
        return ""

    texto = unicodedata.normalize(
        "NFKD",
        str(valor),
    )

    texto = "".join(
        caractere
        for caractere in texto
        if not unicodedata.combining(
            caractere
        )
    )

    texto = texto.replace(
        "_",
        " ",
    )

    texto = re.sub(
        r"\s+",
        " ",
        texto,
    )

    return (
        texto
        .strip()
        .lower()
    )


def normalizar_id(
    valor,
):

    if valor is None:
        return ""

    texto = str(
        valor
    ).strip()

    if not texto:
        return ""

    try:

        numero = Decimal(
            texto.replace(
                ",",
                ".",
            )
        )

        if (
            numero
            == numero.to_integral()
        ):

            return str(
                int(
                    numero
                )
            )

    except InvalidOperation:
        pass

    return texto


def limpar_valor(
    valor,
):

    if valor is None:
        return None

    if isinstance(
        valor,
        str,
    ):

        valor = valor.strip()

        if not valor:
            return None

    return valor


# =========================================================
# LEITURA GENERICA
# =========================================================

def ler_excel(
    arquivo,
    chave_cabecalho,
):

    if not arquivo.exists():

        raise RuntimeError(
            "Arquivo nao encontrado: "
            + str(
                arquivo
            )
        )

    workbook = load_workbook(
        arquivo,
        read_only=True,
        data_only=True,
    )

    try:

        worksheet = workbook[
            workbook.sheetnames[0]
        ]

        linha_cabecalho = None
        cabecalho = None

        for numero_linha, valores in enumerate(
            worksheet.iter_rows(
                min_row=1,
                max_row=20,
                values_only=True,
            ),
            start=1,
        ):

            normalizados = [
                normalizar_texto(
                    valor
                )
                for valor in valores
            ]

            if (
                chave_cabecalho
                in normalizados
            ):

                linha_cabecalho = (
                    numero_linha
                )

                cabecalho = (
                    normalizados
                )

                break

        if linha_cabecalho is None:

            raise RuntimeError(
                "Cabecalho nao encontrado em: "
                + str(
                    arquivo
                )
            )

        registros = []

        for valores in worksheet.iter_rows(
            min_row=(
                linha_cabecalho
                + 1
            ),
            values_only=True,
        ):

            if not any(
                valor is not None
                and str(
                    valor
                ).strip()
                for valor in valores
            ):

                continue

            registro = {}

            for indice, nome in enumerate(
                cabecalho
            ):

                if not nome:
                    continue

                registro[
                    nome
                ] = (
                    valores[
                        indice
                    ]
                    if indice
                    < len(
                        valores
                    )
                    else None
                )

            registros.append(
                registro
            )

        return registros

    finally:

        workbook.close()


# =========================================================
# PLANOS ELEGIVEIS
# =========================================================

def carregar_planos_elegiveis():

    planos = ler_excel(
        ARQUIVO_PLANOS_PRATA,
        "plano id",
    )

    mapa = {}

    for registro in planos:

        plano_id = normalizar_id(
            registro.get(
                "plano id"
            )
        )

        if not plano_id:
            continue

        grupo = limpar_valor(
            registro.get(
                "grupo sso"
            )
        )

        metodo = limpar_valor(
            registro.get(
                "metodo classificacao"
            )
        )

        if not grupo:

            raise RuntimeError(
                "Plano elegivel sem grupo_sso: "
                + plano_id
            )

        if plano_id in mapa:

            anterior = mapa[
                plano_id
            ]

            if (
                str(
                    anterior[
                        "grupo_sso"
                    ]
                ).strip()
                != str(
                    grupo
                ).strip()
            ):

                raise RuntimeError(
                    "Plano com grupos conflitantes "
                    "na Prata: "
                    + plano_id
                )

        mapa[
            plano_id
        ] = {
            "grupo_sso": grupo,
            "metodo_classificacao": metodo,
        }

    if not mapa:

        raise RuntimeError(
            "Nenhum Plano elegivel "
            "foi encontrado na Prata."
        )

    return mapa


# =========================================================
# CAMPOS DA BRONZE
# =========================================================

MAPEAMENTO = {
    "id da acao corretiva":
        "acao_corretiva_id",

    "data de inicio da acao corretiva":
        "data_inicio",

    "prazo para acao corretiva":
        "prazo",

    "modulo de origem":
        "modulo_origem",

    "id da resposta":
        "id_resposta",

    "nome do formulario":
        "nome_formulario",

    "plano de identificacao":
        "plano_id",

    "nome do plano":
        "nome_plano",

    "gestor de plano":
        "gestor_plano",

    "estado do plano":
        "estado_plano",

    "nivel organizacional":
        "nivel_organizacional",

    "perguntar":
        "pergunta",

    "descricao acao corretiva":
        "descricao_acao",

    "responsavel pelas acoes corretivas":
        "responsavel_acao",

    "estado: acao corretiva":
        "estado_acao",

    "estado de aprovacao":
        "estado_aprovacao",

    "oficial de aprovacao":
        "oficial_aprovacao",

    "descricao final":
        "descricao_final",

    "proxima verificacao":
        "proxima_verificacao",
}


def transformar_registro(
    origem,
    plano,
):

    destino = {}

    for (
        coluna_origem,
        coluna_destino,
    ) in MAPEAMENTO.items():

        destino[
            coluna_destino
        ] = limpar_valor(
            origem.get(
                coluna_origem
            )
        )

    destino[
        "acao_corretiva_id"
    ] = normalizar_id(
        destino[
            "acao_corretiva_id"
        ]
    )

    destino[
        "plano_id"
    ] = normalizar_id(
        destino[
            "plano_id"
        ]
    )

    destino[
        "grupo_sso"
    ] = plano[
        "grupo_sso"
    ]

    destino[
        "metodo_classificacao"
    ] = plano[
        "metodo_classificacao"
    ]

    destino[
        "registros_origem"
    ] = 1

    destino[
        "duplicado_origem"
    ] = 0

    return destino


# =========================================================
# CONSOLIDACAO
# =========================================================

def juntar_textos(
    valores,
):

    unicos = []

    vistos = set()

    for valor in valores:

        valor = limpar_valor(
            valor
        )

        if valor is None:
            continue

        chave = normalizar_texto(
            valor
        )

        if not chave:
            continue

        if chave in vistos:
            continue

        vistos.add(
            chave
        )

        unicos.append(
            str(
                valor
            ).strip()
        )

    if not unicos:
        return None

    if len(
        unicos
    ) == 1:

        return unicos[
            0
        ]

    return " | ".join(
        unicos
    )


def consolidar(
    registros,
):

    por_id = defaultdict(
        list
    )

    for registro in registros:

        acao_id = registro[
            "acao_corretiva_id"
        ]

        if not acao_id:

            raise RuntimeError(
                "Acao elegivel sem ID."
            )

        por_id[
            acao_id
        ].append(
            registro
        )

    consolidados = []

    campos_agregaveis = {
        "oficial_aprovacao",
    }

    campos_controle = {
        "registros_origem",
        "duplicado_origem",
    }

    for (
        acao_id,
        grupo,
    ) in por_id.items():

        base = dict(
            grupo[
                0
            ]
        )

        for coluna in COLUNAS_SAIDA:

            if (
                coluna
                in campos_agregaveis
                or coluna
                in campos_controle
            ):
                continue

            valores = {
                normalizar_texto(
                    registro.get(
                        coluna
                    )
                )
                for registro in grupo
                if limpar_valor(
                    registro.get(
                        coluna
                    )
                )
                is not None
            }

            if len(
                valores
            ) > 1:

                raise RuntimeError(
                    "Conflito ao consolidar "
                    f"acao {acao_id}, "
                    f"campo {coluna}: "
                    + str(
                        [
                            registro.get(
                                coluna
                            )
                            for registro in grupo
                        ]
                    )
                )

        base[
            "oficial_aprovacao"
        ] = juntar_textos(
            [
                registro.get(
                    "oficial_aprovacao"
                )
                for registro in grupo
            ]
        )

        base[
            "registros_origem"
        ] = len(
            grupo
        )

        base[
            "duplicado_origem"
        ] = (
            1
            if len(
                grupo
            ) > 1
            else 0
        )

        consolidados.append(
            base
        )

    return consolidados


# =========================================================
# QUALITY GATE
# =========================================================

def validar_prata(
    registros,
):

    if not registros:

        raise RuntimeError(
            "Prata vazia."
        )

    ids = [
        registro[
            "acao_corretiva_id"
        ]
        for registro in registros
    ]

    if len(
        ids
    ) != len(
        set(
            ids
        )
    ):

        raise RuntimeError(
            "Existem IDs duplicados "
            "apos consolidacao."
        )

    grupos_validos = {
        "SEGURANCA",
        "OUTRAS_AREAS",
    }

    grupos_invalidos = {
        str(
            registro.get(
                "grupo_sso"
            )
            or
            ""
        ).strip()
        for registro in registros
        if str(
            registro.get(
                "grupo_sso"
            )
            or
            ""
        ).strip()
        not in grupos_validos
    }

    if grupos_invalidos:

        raise RuntimeError(
            "grupo_sso invalido na Prata: "
            + str(
                sorted(
                    grupos_invalidos
                )
            )
        )

    sem_plano = [
        registro[
            "acao_corretiva_id"
        ]
        for registro in registros
        if not registro.get(
            "plano_id"
        )
    ]

    if sem_plano:

        raise RuntimeError(
            "Acoes sem plano_id: "
            + str(
                sem_plano[:20]
            )
        )


# =========================================================
# GRAVACAO
# =========================================================

def salvar_prata(
    registros,
):

    PASTA_PRATA.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = (
        NOME_ABA
    )

    worksheet.append(
        COLUNAS_SAIDA
    )

    for registro in registros:

        worksheet.append(
            [
                registro.get(
                    coluna
                )
                for coluna
                in COLUNAS_SAIDA
            ]
        )

    workbook.save(
        ARQUIVO_PRATA
    )


# =========================================================
# EXECUCAO
# =========================================================

def main():

    print(
        "=" * 70
    )

    print(
        "TRATAMENTO ACOES CORRETIVAS - PRATA"
    )

    print(
        "=" * 70
    )

    planos_elegiveis = (
        carregar_planos_elegiveis()
    )

    bronze = ler_excel(
        ARQUIVO_BRONZE,
        "id da acao corretiva",
    )

    print()
    print(
        "Linhas recebidas da Bronze:",
        len(
            bronze
        ),
    )

    elegiveis = []

    excluidos = Counter()

    for origem in bronze:

        plano_id = normalizar_id(
            origem.get(
                "plano de identificacao"
            )
        )

        if not plano_id:

            excluidos[
                "SEM_PLANO"
            ] += 1

            continue

        plano = (
            planos_elegiveis.get(
                plano_id
            )
        )

        if plano is None:

            excluidos[
                "PLANO_NAO_ELEGIVEL"
            ] += 1

            continue

        elegiveis.append(
            transformar_registro(
                origem,
                plano,
            )
        )

    print(
        "Linhas elegiveis:",
        len(
            elegiveis
        ),
    )

    print(
        "Linhas excluidas:",
        sum(
            excluidos.values()
        ),
    )

    for (
        motivo,
        quantidade,
    ) in excluidos.items():

        print(
            f"  {motivo}: "
            f"{quantidade}"
        )

    consolidados = consolidar(
        elegiveis
    )

    validar_prata(
        consolidados
    )

    grupos = Counter(
        registro[
            "grupo_sso"
        ]
        for registro
        in consolidados
    )

    estados = Counter(
        str(
            registro.get(
                "estado_acao"
            )
            or
            ""
        ).strip()
        for registro
        in consolidados
    )

    duplicados = [
        registro
        for registro
        in consolidados
        if registro[
            "duplicado_origem"
        ] == 1
    ]

    salvar_prata(
        consolidados
    )

    print()
    print(
        "Acoes unicas na Prata:",
        len(
            consolidados
        ),
    )

    print()
    print(
        "Grupo SSO:"
    )

    for (
        grupo,
        quantidade,
    ) in grupos.items():

        print(
            f"  {grupo}: "
            f"{quantidade}"
        )

    print()
    print(
        "Estados:"
    )

    for (
        estado,
        quantidade,
    ) in sorted(
        estados.items()
    ):

        print(
            f"  {estado}: "
            f"{quantidade}"
        )

    print()
    print(
        "Acoes consolidadas "
        "por duplicidade:",
        len(
            duplicados
        ),
    )

    for registro in duplicados:

        print(
            "  ID:",
            registro[
                "acao_corretiva_id"
            ],
            "| registros:",
            registro[
                "registros_origem"
            ],
            "| oficial:",
            registro[
                "oficial_aprovacao"
            ],
        )

    print()
    print(
        "QUALITY GATE: APROVADO"
    )

    print()
    print(
        "Prata gerada:"
    )

    print(
        ARQUIVO_PRATA
    )

    print()
    print(
        "=" * 70
    )

    print(
        "TRATAMENTO DE ACOES CORRETIVAS "
        "CONCLUIDO COM SUCESSO"
    )

    print(
        "=" * 70
    )


if __name__ == "__main__":
    main()
