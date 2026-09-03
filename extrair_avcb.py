import webbrowser
import time
import shutil
import os

from pathlib import Path

import pyautogui
import pygetwindow as gw

from openpyxl import load_workbook
from dotenv import load_dotenv


# =========================================================
# CONFIGURAÇÕES DO PROJETO
# =========================================================

PASTA_PROJETO = Path(__file__).resolve().parent


# =========================================================
# CONFIGURAÇÕES DO AVCB
# =========================================================

NOME_ABA_AVCB = "Em Criação"

load_dotenv(
    PASTA_PROJETO / ".env"
)

URL_AVCB = os.getenv(
    "AVCB_SHAREPOINT_URL"
)

if not URL_AVCB:
    raise RuntimeError(
        "AVCB_SHAREPOINT_URL não encontrada no arquivo .env."
    )


# =========================================================
# PASTAS
# =========================================================

PASTA_BRONZE = (
    PASTA_PROJETO
    / "bronze"
    / "avcb"
)

PASTA_DOWNLOADS = (
    Path.home()
    / "Downloads"
)

ARQUIVO_FINAL = (
    PASTA_BRONZE
    / "Controle de AVCB.xlsx"
)


# =========================================================
# PREPARA A JANELA DO NAVEGADOR
# =========================================================

def preparar_janela_avcb(timeout=30):

    print(
        "Localizando janela do Controle de AVCB..."
    )

    limite = (
        time.time()
        + timeout
    )

    while time.time() < limite:

        janelas = [
            janela
            for janela in gw.getAllWindows()
            if janela.title
            and "Controle de AVCB" in janela.title
        ]

        if janelas:

            janela = janelas[0]

            print(
                f"Janela encontrada: {janela.title}"
            )

            # Restaura caso esteja minimizada
            if janela.isMinimized:

                janela.restore()

                time.sleep(1)

            # Se estiver maximizada,
            # restaura primeiro para permitir
            # mover entre os monitores
            if janela.isMaximized:

                janela.restore()

                time.sleep(1)

            print(
                "Movendo janela para a tela principal..."
            )

            # 0,0 = canto superior esquerdo
            # da tela principal do Windows
            janela.moveTo(
                0,
                0
            )

            time.sleep(1)

            print(
                "Maximizando janela..."
            )

            janela.maximize()

            time.sleep(2)

            # Tenta colocar a janela
            # em primeiro plano
            try:

                janela.activate()

            except Exception:

                pass

            time.sleep(1)

            print(
                "Janela preparada na tela principal."
            )

            return

        time.sleep(1)

    raise RuntimeError(
        "Não foi possível localizar a janela "
        "do Controle de AVCB."
    )


# =========================================================
# LOCALIZA O ARQUIVO RECÉM-BAIXADO
# =========================================================

def encontrar_download_novo(
    inicio_execucao,
    timeout=60
):

    print(
        "Aguardando download..."
    )

    limite = (
        time.time()
        + timeout
    )

    while time.time() < limite:

        arquivos = list(
            PASTA_DOWNLOADS.glob(
                "Controle de AVCB*.xlsx"
            )
        )

        arquivos_novos = [
            arquivo
            for arquivo in arquivos
            if arquivo.stat().st_mtime
            >= inicio_execucao
        ]

        if arquivos_novos:

            mais_recente = max(
                arquivos_novos,
                key=lambda arquivo:
                    arquivo.stat().st_mtime
            )

            # Espera o navegador terminar
            # completamente o download
            time.sleep(3)

            return mais_recente

        time.sleep(1)

    raise TimeoutError(
        "O download não foi encontrado "
        "dentro do tempo esperado."
    )


# =========================================================
# VALIDA O ARQUIVO BAIXADO
# =========================================================

def validar_excel(
    arquivo
):

    print(
        f"Validando arquivo: {arquivo.name}"
    )

    workbook = load_workbook(
        arquivo,
        read_only=True,
        data_only=True
    )

    abas = (
        workbook.sheetnames
    )

    workbook.close()

    if NOME_ABA_AVCB not in abas:

        raise ValueError(
            f"A aba '{NOME_ABA_AVCB}' "
            "não foi encontrada."
        )

    print(
        "Arquivo Excel válido."
    )

    print(
        f"Aba '{NOME_ABA_AVCB}' encontrada."
    )


# =========================================================
# ATUALIZA A CAMADA BRONZE
# =========================================================

def atualizar_bronze(
    arquivo_baixado
):

    PASTA_BRONZE.mkdir(
        parents=True,
        exist_ok=True
    )

    arquivo_temporario = (
        PASTA_BRONZE
        / "Controle de AVCB_novo.xlsx"
    )

    print(
        "Copiando nova versão "
        "para a camada Bronze..."
    )

    shutil.copy2(
        arquivo_baixado,
        arquivo_temporario
    )

    # Só substitui a versão anterior
    # depois que o novo arquivo
    # já foi validado
    os.replace(
        arquivo_temporario,
        ARQUIVO_FINAL
    )

    # Remove a cópia de Downloads
    arquivo_baixado.unlink()

    print(
        "Camada Bronze atualizada "
        "com sucesso:"
    )

    print(
        ARQUIVO_FINAL
    )


# =========================================================
# EXTRAÇÃO DO AVCB
# =========================================================

def extrair_avcb():

    inicio_execucao = (
        time.time()
    )

    print(
        "Abrindo Controle de AVCB "
        "no SharePoint..."
    )

    webbrowser.open(
        URL_AVCB
    )

    # Tempo para o Excel Online carregar
    time.sleep(12)

    # Garante que a janela esteja
    # na tela principal antes dos cliques
    preparar_janela_avcb()

    print(
        "Abrindo menu Arquivo..."
    )

    pyautogui.click(
        77,
        262
    )

    time.sleep(2)

    print(
        "Abrindo submenu "
        "Criar uma Cópia..."
    )

    pyautogui.moveTo(
        305,
        511,
        duration=0.5
    )

    time.sleep(2)

    print(
        "Clicando em "
        "Baixar uma Cópia..."
    )

    pyautogui.click(
        471,
        540
    )

    arquivo = encontrar_download_novo(
        inicio_execucao
    )

    validar_excel(
        arquivo
    )

    atualizar_bronze(
        arquivo
    )


# =========================================================
# EXECUÇÃO
# =========================================================

if __name__ == "__main__":
    extrair_avcb()